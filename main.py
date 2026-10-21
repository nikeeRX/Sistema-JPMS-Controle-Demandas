from flask import Flask, request, redirect, url_for, session, flash, render_template_string, send_file, jsonify
from config import setup_db, conectar_db
import os
import PyPDF2
import re
from datetime import datetime
import base64
from io import BytesIO
import pandas as pd
import unicodedata
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages

app = Flask(__name__)
app.secret_key = 'chave_super_secreta_da_postal_saude'

setup_db()

def get_ph(conn):
    return "%s" if "psycopg2" in str(type(conn)) else "?"

# ==========================================
# INICIALIZAÇÃO E AUTO-REPARAÇÃO DO BANCO
# ==========================================
def inicializar_tabela_credenciamento():
    try:
        conn = conectar_db()
        if not conn: return
        c = conn.cursor()
        is_postgres = "psycopg2" in str(type(conn))
        id_col = "SERIAL PRIMARY KEY" if is_postgres else "INTEGER PRIMARY KEY AUTOINCREMENT"
        c.execute(f"""
            CREATE TABLE IF NOT EXISTS credenciamento_fases (
                id {id_col},
                filial TEXT,
                cnpj_cpf TEXT,
                nome TEXT,
                fantasia TEXT,
                tipo_prestador TEXT,
                categoria TEXT,
                tipo_processo TEXT,
                processo TEXT,
                data_abertura TEXT,
                data_final_proc TEXT,
                usuario_resp TEXT,
                fase TEXT,
                ultima_fase TEXT,
                data_inicial TEXT,
                data_final TEXT,
                proxima_fase TEXT,
                dias_percorridos TEXT,
                dias_percorridos_pf TEXT,
                justificativa_indef TEXT,
                parecer TEXT,
                observacao TEXT,
                origem_parecer TEXT,
                data_aprovacao TEXT,
                data_rejeicao TEXT,
                usuario_aprj TEXT,
                justificativa_corgen TEXT,
                usuario_notificacao TEXT,
                grupo_prestadores TEXT,
                data_atualizacao TEXT
            )
        """)
        conn.commit()
        conn.close()
    except Exception as e:
        print(f"Aviso ao criar tabela credenciamento_fases: {e}")

def sincronizar_sequencias():
    try:
        conn = conectar_db()
        if conn and "psycopg2" in str(type(conn)):
            c = conn.cursor()
            for tabela in ['demands', 'demandas_avulsas', 'users', 'credenciamento_fases']:
                try:
                    c.execute(f"SELECT setval(pg_get_serial_sequence('{tabela}', 'id'), COALESCE((SELECT MAX(id) FROM {tabela}), 0) + 1, false);")
                    conn.commit()
                except Exception:
                    conn.rollback()
            conn.close()
    except Exception as e:
        print(f"Aviso ao sincronizar sequências: {e}")

inicializar_tabela_credenciamento()
sincronizar_sequencias()

# ==========================================
# MOTORES DE CÁLCULO DE SLA (COM TRAVA DE PARADA)
# ==========================================
STATUS_ENCERRADOS_DEMANDAS = {
    'CONCLUÍDO', 'CONCLUIDO', 'FINALIZADA', 'FINALIZADO',
    'DEVOLVIDO', 'DEVOLVIDA', 'CANCELADA', 'CANCELADO',
    'DEFERIDO', 'INDEFERIDO'
}

def parse_data_br(data_str):
    if not data_str: return None
    ds = str(data_str).strip()
    if not ds or ds.upper() in ['NONE', 'NAN', '']: return None
    try:
        part = ds.split(' ')[0]
        if '-' in part:
            return datetime.strptime(part, "%Y-%m-%d")
        return datetime.strptime(part, "%d/%m/%Y")
    except Exception:
        return None

def calcular_sla(data_str, status, data_fim=None):
    st_upper = str(status or '').strip().upper()
    if not data_str or st_upper in STATUS_ENCERRADOS_DEMANDAS:
        return -1 
    dt = parse_data_br(data_str)
    if not dt: return -1
    dias = (datetime.now() - dt).days
    return dias if dias >= 0 else 0

def calcular_dias_fechados(data_ini_str, data_fim_str):
    dt_ini = parse_data_br(data_ini_str)
    dt_fim = parse_data_br(data_fim_str)
    if dt_ini and dt_fim:
        d = (dt_fim - dt_ini).days
        return d if d >= 0 else 0
    return None

def calcular_sla_credenciamento(row):
    parecer = str(row[20] or '').strip().upper()
    data_abertura = row[9]
    data_final_proc = str(row[10] or '').strip()
    data_inicial_fase = row[14]
    data_final_fase = str(row[15] or '').strip()

    parado = False
    if parecer in STATUS_ENCERRADOS_DEMANDAS:
        parado = True
    elif (data_final_proc != '' or data_final_fase != '') and parecer not in ['EM ANÁLISE', 'EM ANALISE', 'PENDENTE', 'AGUARDANDO 90 DIAS', 'AGUARDANDO RETORNO']:
        parado = True

    dt_ini = parse_data_br(data_inicial_fase) or parse_data_br(data_abertura)
    dt_abert = parse_data_br(data_abertura)
    dt_fim = parse_data_br(data_final_fase) or parse_data_br(data_final_proc)

    hoje = datetime.now()
    if not parado:
        dias_fase = max(0, (hoje - dt_ini).days) if dt_ini else 0
        dias_geral = max(0, (hoje - dt_abert).days) if dt_abert else dias_fase
    else:
        if dt_ini and dt_fim:
            dias_fase = max(0, (dt_fim - dt_ini).days)
        elif row[17] and str(row[17]).strip().isdigit():
            dias_fase = int(str(row[17]).strip())
        else:
            dias_fase = 0

        if dt_abert and dt_fim:
            dias_geral = max(0, (dt_fim - dt_abert).days)
        elif row[18] and str(row[18]).strip().isdigit():
            dias_geral = int(str(row[18]).strip())
        else:
            dias_geral = dias_fase

    cor_linha = ""
    faixa_sla = "Parado"
    if not parado:
        if dias_fase >= 6:
            cor_linha = "bg-critico"
            faixa_sla = "Crítico"
        elif dias_fase >= 4:
            cor_linha = "bg-urgente"
            faixa_sla = "Urgente"
        elif dias_fase >= 2:
            cor_linha = "bg-atencao"
            faixa_sla = "Atenção"
        else:
            faixa_sla = "No Prazo"

    return {
        "ativo": not parado,
        "dias_fase": dias_fase,
        "dias_geral": dias_geral,
        "cor_linha": cor_linha,
        "faixa_sla": faixa_sla,
        "estourou_90d": (not parado and dias_geral >= 90)
    }

def corrigir_encoding_postal(texto):
    if pd.isna(texto) or texto is None:
        return ""
    s = str(texto).strip()
    substituicoes = {
        'Ă': 'Ã', 'ă': 'ã',
        'ő': 'õ', 'Ő': 'Õ',
        'ę': 'ê', 'Ę': 'Ê',
        'ŕ': 'à', 'Ŕ': 'À',
        'ş': 'º', 'Ş': 'ª',
        'ć': 'ç', 'Ć': 'Ç'
    }
    for k, v in substituicoes.items():
        s = s.replace(k, v)
    return s

# ==========================================
# LISTAS GLOBAIS E CONSTANTES
# ==========================================
TIPOS_ROTINA = ["Reajuste", "Novo Contrato", "Inclusão", "Exclusão", "Ajuste", "Extensão", "Descredenciamento", "Não Identificada"]
STATUS_ROTINA = ['Pendente', 'Em Análise', 'Aguardando 90 dias', 'Aguardando Retorno', 'Devolvido', 'Finalizada', 'Cancelada']

FASES_PADRAO_CRED = [
    "1-DEMANDA",
    "3-PARECER REGIONAL-NOVO CREDEN (ALÇADA SEDE)",
    "3-PARECER SEDE - NOVO CREDENCIADO",
    "3-REANÁLISE REGIONAL-NOVO CREDEN (ALÇADA SEDE)",
    "3-REANÁLISE SEDE - NOVO CREDENCIADO",
    "4-ANÁLISE DOC. PROPONENTE- (ALÇADA SEDE)",
    "4-PARECER SEDE - DOC. PROPONENTE",
    "4-REANÁLISE DOC. PROPONENTE (ALÇADA SEDE)",
    "4-REANÁLISE SEDE - DOC. PROPONENTE",
    "5-ANÁLISE CONTRATO CREDENCIAMENTO",
    "5-PARECER SEDE - CONTRATO CREDENCIAMENTO",
    "6-PARECER FINAL (SEDE)",
    "8-GARANTIA DE ATENDIMENTO",
    "9-FORNECEDOR"
]

STATUS_CRED = [
    "Pendente",
    "Em Análise",
    "EM ANÁLISE",
    "Aguardando 90 dias",
    "Aguardando Retorno",
    "Devolvido",
    "DEVOLVIDO",
    "Finalizada",
    "DEFERIDO",
    "INDEFERIDO",
    "Cancelada"
]

TIPOS_AVULSA = ['Credenciamento', 'Descredenciamento', 'Análise de Vulnerabilidade', 'Parametrização', 'Resposta de E-mail', 'RPS', 'Lote de Erros', 'Atendimento de Chamado', 'Atualização Cadastral', 'Outro']
STATUS_AVULSA = ['Pendente', 'Recebida', 'Em Análise', 'Aguardando 90 dias', 'Aguardando Retorno', 'Aguardando Área Demandante', 'Em andamento', 'Devolvido', 'Concluído', 'Cancelada']

TODOS_STATUS_ORDEM = [
    'Pendente',
    'Recebida',
    'Em Análise',
    'Em andamento',
    'Aguardando 90 dias',
    'Aguardando Retorno',
    'Aguardando Área Demandante',
    'Devolvido',
    'Finalizada',
    'Concluído',
    'Cancelada'
]

SOLICITANTES = ['GEREL/COATE', 'GEREL/COCAD', 'UAR-AM', 'NAR-AC', 'NAR-RR', 'UAR-BA', 'NAR-AL', 'NAR-SE', 'UAR-CE', 'NAR-MA', 'NAR-PI', 'UAR-DF', 'NAR-GO', 'NAR-TO', 'UAR-MG', 'UAR-MS', 'NAR-MT', 'NAR-RO', 'UAR-PA', 'NAR-AP', 'UAR-PE', 'NAR-PB', 'NAR-RN', 'UAR-RJ', 'NAR-ES', 'UAR-RS', 'NAR-SC', 'UAR-SPI', 'UAR-SPM', 'NAR-PR', 'GEREM/COANS', 'GERIN', 'GEASA/COICS', 'GERED/COCAP']

IGNORADAS = {"DIARIAS", "PACOTE", "TAXAS E GASES", "MATERIAIS E OPME", "MEDICAMENTOS", "FORNECEDOR DE MEDICAMENTO ONCOLOGICO", "INSTRUMENTADOR CIRURGICO", "HOME CARE - ATENDIMENTO DOMICILIAR", "REMOCAO", "MEDICINA LEGAL E PERICIA MEDICA", "MEDICO HIPERBARISTA"}
DE_PARA = {"MEDICO HEMOTERAPEUTA": "HEMATOLOGIA E HEMOTERAPIA", "GENETICA MEDICA": "MEDICO GENETICISTA", "MEDICO CANCEROLOGISTA CIRURGICO": "CIRURGIA ONCOLOGICA", "MEDICO CANCEROLOGISTA PEDIATRICO": "CANCEROLOGIA", "MEDICO DE FAMILIA E COMUNIDADE": "MEDICINA DE FAMILIA E COMUNIDADE", "MEDICO GENERALISTA": "CLINICA MEDICA", "MEDICO PATOLOGISTA": "PATOLOGIA CLINICA/MEDICINA LABORATORIAL", "PATOLOGIA": "PATOLOGIA CLINICA/MEDICINA LABORATORIAL", "MEDICO NEUROFISIOLOGISTA": "NEUROFISIOLOGIA CLINICA", "NEUROPSICOLOGO": "PSICOLOGIA", "PSICOPEDAGOGO": "PSICOLOGIA", "PSICOMOTRICISTA": "FISIOTERAPIA", "MUSICOTERAPEUTA": "TERAPIA OCUPACIONAL", "ORTOPEDISTA": "ORTOPEDIA E TRAUMATOLOGIA", "BIOMEDICO": "PATOLOGIA CLINICA/MEDICINA LABORATORIAL", "FONOAUDIOLOGO EDUCACIONAL": "FONOAUDIOLOGIA", "FONOAUDIOLOGO EM AUDIOLOGIA": "FONOAUDIOLOGIA", "FONOAUDIOLOGO EM DISFAGIA": "FONOAUDIOLOGIA", "FONOAUDIOLOGO EM LINGUAGEM": "FONOAUDIOLOGIA", "FONOAUDIOLOGO EM MOTRICIDADE OROFACIAL": "FONOAUDIOLOGIA", "FONOAUDIOLOGO EM VOZ": "FONOAUDIOLOGIA", "FISIOTERAPEUTA NEUROFUNCIONAL": "FISIOTERAPIA", "FISIOTERAPEUTA OSTEOPATA": "FISIOTERAPIA", "FISIOTERAPEUTA RESPIRATORIA": "FISIOTERAPIA", "FISIOTERAPEUTA TRAUMATO-ORTOPEDICA FUNCIONAL": "FISIOTERAPIA", "ENFERMEIRO DA ESTRATEGIA DE SAUDE DA FAMILIA": "ENFERMEIRO", "CIRURGIAO DENTISTA - TRAUMATOLOGISTA BUCOMAXILOFAC": "CIRURGIA E TRAUMATOLOGIA BUCO-MAXILO-FACIAL", "CLINICA GERAL - ODONTOLOGIA": "ODONTOLOGIA", "CIRURGIAO DENTISTA - DENTISTICA": "ODONTOLOGIA", "CIRURGIAO DENTISTA - DISFUNCAO TEMPOROMANDIBULAR E": "ODONTOLOGIA", "DENTISTICA RESTAURADORA": "ODONTOLOGIA", "DISFUNCAO TEMPOROMANDIBULAR E DOR OROFACIAL": "ODONTOLOGIA", "ENDODONTIA": "ODONTOLOGIA", "ESTOMATOLOGIA": "ODONTOLOGIA", "ODONTOLOGIA - AUDITORIA INICIAL E FINAL": "ODONTOLOGIA", "ODONTOLOGIA DO TRABALHO": "ODONTOLOGIA", "ODONTOLOGIA P/ PACIENTE COM NECESSIDADE ESPECIAL": "ODONTOLOGIA", "ODONTOPEDIATRIA": "ODONTOLOGIA", "PERIODONTIA": "ODONTOLOGIA", "PROTESE DENTARIA": "ODONTOLOGIA", "RADIOLOGIA ODONTOLOGICA E IMAGINOLOGIA - RX ODONTO": "RADIOLOGIA E DIAGNOSTICO POR IMAGEM", "NEUROLIGIA": "NEUROLOGIA"}

# ==========================================
# CÓDIGOS HTML EMBUTIDOS
# ==========================================

TELA_LOGIN = """<!DOCTYPE html><html lang="pt-BR"><head><meta charset="UTF-8"><title>Login - COCAP</title><link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet"></head>
<body style="background-color: #f4f7f6; display: flex; align-items: center; justify-content: center; height: 100vh;">
<div style="border-radius: 15px; box-shadow: 0 10px 30px rgba(0,0,0,0.1); padding: 40px; background: white; width: 100%; max-width: 400px;">
    <div class="text-center mb-4"><h3 style="color: #004b87; font-weight: 800;">POSTAL SAÚDE</h3><p class="text-muted">Central de Operações (COCAP)</p></div>
    {% with messages = get_flashed_messages(with_categories=true) %}{% if messages %}{% for category, message in messages %}<div class="alert alert-{{ category }}">{{ message }}</div>{% endfor %}{% endif %}{% endwith %}
    <form method="POST" action="/">
        <div class="mb-3"><label class="form-label fw-bold">Usuário</label><input type="text" name="username" class="form-control" required></div>
        <div class="mb-4"><label class="form-label fw-bold">Senha</label><input type="password" name="password" class="form-control" required></div>
        <button type="submit" class="btn w-100 py-2" style="background-color: #004b87; color: white; font-weight: bold;">Entrar no Sistema</button>
    </form>
</div></body></html>"""

TELA_DASHBOARD = """<!DOCTYPE html><html lang="pt-BR"><head><meta charset="UTF-8"><title>Portal COCAP</title><link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet"></head>
<body style="background-color: #f0f2f5;">
<nav class="navbar navbar-expand-lg navbar-dark" style="background-color: #004b87;">
  <div class="container-fluid"><a class="navbar-brand fw-bold" href="#">COCAP Web</a>
    <div class="d-flex text-white align-items-center"><span class="me-3">Olá, <strong>{{ user }}</strong> ({{ role }})</span><a href="/logout" class="btn btn-danger btn-sm">Sair</a></div>
  </div>
</nav>
<div class="container mt-5 text-center">
    <h1 style="color: #004b87; font-weight: 800; margin-bottom: 2rem;">PORTAL CENTRAL DE OPERAÇÕES</h1>
    <div class="row justify-content-center g-4">
        <div class="col-md-4">
            <div class="card shadow-sm h-100 p-4 border-0">
                <h1 class="text-primary mb-3" style="font-size: 3rem;">📂</h1>
                <h4 class="fw-bold">Controle de Demandas</h4>
                <p class="text-muted mb-4">Acesse Rotinas (PDFs), Demandas Avulsas e Credenciamento - Fases.</p>
                <a href="/demandas" class="btn btn-primary text-white w-100 fw-bold py-2 mt-auto">Acessar Demandas</a>
            </div>
        </div>
        <div class="col-md-4">
            <div class="card shadow-sm h-100 p-4 border-0">
                <h1 class="text-info mb-3" style="font-size: 3rem;">📊</h1>
                <h4 class="fw-bold">Dashboard Gerencial</h4>
                <p class="text-muted mb-4">Painel executivo consolidado de Status, SLAs, Credenciamento e Relatório PDF.</p>
                <a href="/demandas/dashboard_graficos" class="btn btn-info text-white w-100 fw-bold py-2 mt-auto">Acessar Dashboard</a>
            </div>
        </div>
        <div class="col-md-4">
            <div class="card shadow-sm h-100 p-4 border-0">
                <h1 class="text-warning mb-3" style="font-size: 3rem;">🏥</h1>
                <h4 class="fw-bold">Pente Fino (RN 665)</h4>
                <p class="text-muted mb-4">Substituição de Prestadores (Sniper 2 Etapas) e Cruzamento Regional em Massa.</p>
                <a href="/pentefino" class="btn btn-warning text-white w-100 fw-bold py-2 mt-auto">Acessar Módulo</a>
            </div>
        </div>
    </div>
    {% if role == 'admin' %}<div class="row justify-content-center mt-4"><div class="col-md-6"><a href="/admin" class="btn btn-secondary w-100 fw-bold py-3 shadow-sm border-0">⚙️ Painel Administrativo (BD & Usuários)</a></div></div>{% endif %}
</div></body></html>"""

TELA_HUB_DEMANDAS = """<!DOCTYPE html><html lang="pt-BR"><head><title>Módulo de Demandas</title><link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet"></head>
<body style="background-color: #f0f2f5;">
<nav class="navbar navbar-dark bg-dark">
  <div class="container-fluid">
    <a href="/dashboard" class="btn btn-outline-light btn-sm">⬅ Voltar ao Portal</a>
    <span class="navbar-text text-white fw-bold">Central de Controle de Demandas</span>
    <a href="/demandas/relatorio_pdf" class="btn btn-danger btn-sm fw-bold">📄 Baixar Relatório Geral (PDF)</a>
  </div>
</nav>
<div class="container mt-5 text-center">
    <h2 class="fw-bold mb-4" style="color: #004b87;">SELECIONE A FRENTE OPERACIONAL</h2>
    <div class="row justify-content-center g-4">
        <div class="col-md-3">
            <div class="card shadow-sm border-0 h-100 p-4">
                <h1 class="mb-3">📄</h1>
                <h5 class="fw-bold text-secondary">Rotinas (PDFs)</h5>
                <p class="text-muted small mb-4">Importação automática de lotes PDF e controle de SLA duplo (COCAP/FOP).</p>
                <a href="/demandas/rotinas" class="btn btn-primary w-100 py-3 fw-bold mt-auto">Acessar PDFs</a>
            </div>
        </div>
        <div class="col-md-3">
            <div class="card shadow-sm border-0 h-100 p-4">
                <h1 class="mb-3">📝</h1>
                <h5 class="fw-bold text-secondary">Demandas Avulsas</h5>
                <p class="text-muted small mb-4">Registro manual de solicitações, prazos e acompanhamento por área.</p>
                <a href="/demandas/avulsas" class="btn btn-warning text-white w-100 py-3 fw-bold mt-auto">Acessar Avulsas</a>
            </div>
        </div>
        <div class="col-md-3">
            <div class="card shadow-sm border-0 h-100 p-4">
                <h1 class="mb-3">📋</h1>
                <h5 class="fw-bold text-secondary">Credenciamento - Fases</h5>
                <p class="text-muted small mb-4">Relatório diário POSTAL191, edição de fase/status e linha do tempo.</p>
                <a href="/credenciamento" class="btn btn-success text-white w-100 py-3 fw-bold mt-auto">Acessar Fases</a>
            </div>
        </div>
        <div class="col-md-3">
            <div class="card shadow-sm border-0 h-100 p-4">
                <h1 class="mb-3">📊</h1>
                <h5 class="fw-bold text-secondary">Dashboard Gerencial</h5>
                <p class="text-muted small mb-4">Visão consolidada das 3 frentes, indicadores de SLA, gráficos e PDF.</p>
                <a href="/demandas/dashboard_graficos" class="btn btn-info text-white w-100 py-3 fw-bold mt-auto">Ver Dashboard</a>
            </div>
        </div>
    </div>
</div></body></html>"""

TELA_CREDENCIAMENTO_FASES = """<!DOCTYPE html><html lang="pt-BR"><head><meta charset="UTF-8"><title>Credenciamento - Fases</title>
<link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
<style>
    .bg-atencao { background-color: #fff3cd !important; }
    .bg-urgente { background-color: #ffe8cc !important; }
    .bg-critico { background-color: #f8d7da !important; }
    .fase-card { transition: transform 0.15s ease-in-out; }
    .fase-card:hover { transform: translateY(-2px); }
</style>
</head>
<body style="background-color: #f0f2f5;">
<nav class="navbar navbar-dark bg-dark">
  <div class="container-fluid">
    <div class="d-flex gap-2">
        <a href="/demandas" class="btn btn-outline-light btn-sm">⬅ Voltar às Demandas</a>
        <a href="/demandas/dashboard_graficos" class="btn btn-outline-info btn-sm">📊 Ir ao Dashboard</a>
    </div>
    <span class="navbar-text text-white fw-bold">📋 Demandas: Credenciamento - Fases {% if ultima_att %}<small class="text-warning ms-2">(Atualizado em: {{ ultima_att }})</small>{% endif %}</span>
    <div class="d-flex gap-2">
        <button class="btn btn-info btn-sm fw-bold text-white" type="button" data-bs-toggle="collapse" data-bs-target="#painelGraficosCred">📊 Mostrar/Ocultar Gráficos</button>
        <a href="/demandas/relatorio_pdf" class="btn btn-danger btn-sm fw-bold">📄 Relatório Geral (PDF)</a>
        <a href="/credenciamento/exportar?{{ request.query_string.decode('utf-8') }}" class="btn btn-success btn-sm fw-bold">⬇ Exportar Visão (Excel)</a>
    </div>
  </div>
</nav>

<div class="container-fluid mt-3">
    {% with messages = get_flashed_messages(with_categories=true) %}{% if messages %}{% for category, message in messages %}<div class="alert alert-{{ category }}">{{ message }}</div>{% endfor %}{% endif %}{% endwith %}

    <!-- Upload Diário do Relatório POSTAL191 -->
    <div class="card shadow-sm mb-3 border-0">
        <div class="card-body py-2 bg-white rounded">
            <form method="POST" action="/credenciamento/upload" enctype="multipart/form-data" class="row g-2 align-items-center">
                <div class="col-md-4">
                    <span class="fw-bold text-success">🔄 Atualizar Relatório Diário (TXT / CSV / Excel):</span>
                </div>
                <div class="col-md-5">
                    <input type="file" name="arquivo_cred" class="form-control form-control-sm" accept=".txt,.csv,.xlsx,.xls" required>
                </div>
                <div class="col-md-3">
                    <button type="submit" class="btn btn-success btn-sm w-100 fw-bold" onclick="this.innerHTML='A importar relatório... Aguarde!'; this.style.opacity='0.7';">⬆ Carregar e Atualizar Base</button>
                </div>
            </form>
        </div>
    </div>

    <!-- Monitoramento de SLA do Credenciamento -->
    <div class="d-flex justify-content-between align-items-center mb-2">
        <h6 class="mb-0 text-secondary fw-bold">⏱️ Monitoramento de SLA por Fase (Ativos na Visão: {{ sla_resumo.ativos }} | Parados/Concluídos/Devolvidos: {{ registros|length - sla_resumo.ativos }})</h6>
        <small class="text-muted">Nota: Status Finalizada, Devolvido, Deferido ou Indeferido param a contagem do SLA automaticamente.</small>
    </div>
    <div class="row text-center g-2 mb-3">
        <div class="col-md-2 col-6">
            <div class="card shadow-sm border-0 p-2 h-100" style="background-color: #e2e3e5; color: #383d41;">
                <small class="fw-bold">No Prazo (0 a 1d)</small>
                <h3 class="fw-bold mb-0">{{ sla_resumo.no_prazo }}</h3>
            </div>
        </div>
        <div class="col-md-2 col-6">
            <div class="card shadow-sm border-0 p-2 h-100" style="background-color: #fff3cd; color: #856404;">
                <small class="fw-bold">Atenção (2 a 3d)</small>
                <h3 class="fw-bold mb-0">{{ sla_resumo.atencao }}</h3>
            </div>
        </div>
        <div class="col-md-2 col-6">
            <div class="card shadow-sm border-0 p-2 h-100" style="background-color: #ffe8cc; color: #d35400;">
                <small class="fw-bold">Urgente (4 a 5d)</small>
                <h3 class="fw-bold mb-0">{{ sla_resumo.urgente }}</h3>
            </div>
        </div>
        <div class="col-md-2 col-6">
            <div class="card shadow-sm border-0 p-2 h-100" style="background-color: #f8d7da; color: #721c24;">
                <small class="fw-bold">Crítico (6+ dias)</small>
                <h3 class="fw-bold mb-0">{{ sla_resumo.critico }}</h3>
            </div>
        </div>
        <div class="col-md-2 col-6">
            <div class="card shadow-sm border-0 p-2 h-100 text-white" style="background-color: #6f42c1;">
                <small class="fw-bold">⚠️ Alerta 90+ Dias (Ativos)</small>
                <h3 class="fw-bold mb-0">{{ sla_resumo.alerta_90d }}</h3>
            </div>
        </div>
        <div class="col-md-2 col-6">
            <div class="card shadow-sm border-0 bg-dark text-white p-2 h-100">
                <small class="fw-bold">Processos Únicos</small>
                <h3 class="fw-bold mb-0">{{ total_proc_unicos }}</h3>
            </div>
        </div>
    </div>

    <!-- Cards Resumo por Status / Parecer -->
    <div class="row text-center g-2 mb-3">
        <div class="col-md-2 col-6">
            <div class="card shadow-sm border-0 bg-primary text-white p-2 h-100">
                <small class="fw-bold">Total na Visão</small>
                <h4 class="fw-bold mb-0">{{ registros|length }}</h4>
            </div>
        </div>
        <div class="col-md-3 col-6">
            <div class="card shadow-sm border-0 bg-warning text-dark p-2 h-100">
                <small class="fw-bold">Em Análise / Pendentes (SLA Ativo)</small>
                <h4 class="fw-bold mb-0">{{ kpi_parecer.get('EM ANÁLISE', 0) + kpi_parecer.get('EM ANALISE', 0) + kpi_parecer.get('PENDENTE', 0) + kpi_parecer.get('AGUARDANDO 90 DIAS', 0) + kpi_parecer.get('AGUARDANDO RETORNO', 0) }}</h4>
            </div>
        </div>
        <div class="col-md-3 col-6">
            <div class="card shadow-sm border-0 bg-success text-white p-2 h-100">
                <small class="fw-bold">Deferidos / Finalizados (SLA Parado)</small>
                <h4 class="fw-bold mb-0">{{ kpi_parecer.get('DEFERIDO', 0) + kpi_parecer.get('FINALIZADA', 0) + kpi_parecer.get('CONCLUÍDO', 0) }}</h4>
            </div>
        </div>
        <div class="col-md-2 col-6">
            <div class="card shadow-sm border-0 p-2 h-100 text-white" style="background-color: #fd7e14;">
                <small class="fw-bold">Devolvidos (SLA Parado)</small>
                <h4 class="fw-bold mb-0">{{ kpi_parecer.get('DEVOLVIDO', 0) }}</h4>
            </div>
        </div>
        <div class="col-md-2 col-6">
            <div class="card shadow-sm border-0 bg-danger text-white p-2 h-100">
                <small class="fw-bold">Indeferidos / Cancel.</small>
                <h4 class="fw-bold mb-0">{{ kpi_parecer.get('INDEFERIDO', 0) + kpi_parecer.get('CANCELADA', 0) }}</h4>
            </div>
        </div>
    </div>

    <!-- Gráficos Gerenciais (Colapsável) -->
    <div class="collapse mb-3" id="painelGraficosCred">
        <div class="row g-3">
            <div class="col-md-6">
                <div class="card shadow-sm border-0 h-100">
                    <div class="card-body text-center">
                        <h6 class="fw-bold text-secondary mb-2">📊 Distribuição por Fase (Visão Atual)</h6>
                        {% if chart_fases %}<img src="data:image/png;base64,{{ chart_fases }}" class="img-fluid rounded">{% else %}<p class="text-muted py-4">Sem dados para gráfico.</p>{% endif %}
                    </div>
                </div>
            </div>
            <div class="col-md-6">
                <div class="card shadow-sm border-0 h-100">
                    <div class="card-body text-center">
                        <h6 class="fw-bold text-secondary mb-2">🗺️ Processos por UF (Visão Atual)</h6>
                        {% if chart_ufs %}<img src="data:image/png;base64,{{ chart_ufs }}" class="img-fluid rounded">{% else %}<p class="text-muted py-4">Sem dados para gráfico.</p>{% endif %}
                    </div>
                </div>
            </div>
        </div>
    </div>

    <!-- Contagem Rápida por Fase -->
    {% if kpi_fases %}
    <div class="row g-2 mb-3">
        {% for f_nome, f_qtd in kpi_fases %}
        <div class="col-md-3 col-sm-6">
            <div class="card fase-card shadow-sm border-0 p-2 h-100" style="border-left: 4px solid #004b87 !important;">
                <div class="d-flex justify-content-between align-items-center">
                    <span class="small fw-bold text-secondary text-truncate" title="{{ f_nome }}">{{ f_nome }}</span>
                    <span class="badge bg-primary fs-6 ms-2">{{ f_qtd }}</span>
                </div>
            </div>
        </div>
        {% endfor %}
    </div>
    {% endif %}

    <!-- Barra de Filtros -->
    <div class="card shadow-sm mb-3 border-0 bg-light">
        <div class="card-body py-2">
            <form method="GET" action="/credenciamento" class="row g-2 align-items-end">
                <div class="col-md-2">
                    <label class="form-label small fw-bold mb-1">Última Fase (Atual)</label>
                    <select name="ultima_fase" class="form-select form-select-sm border-primary">
                        <option value="SIM" {% if f_ultima == 'SIM' %}selected{% endif %}>Apenas Fase Atual (SIM)</option>
                        <option value="Todos" {% if f_ultima == 'Todos' %}selected{% endif %}>Todo o Histórico (SIM + NÃO)</option>
                        <option value="NÃO" {% if f_ultima == 'NÃO' %}selected{% endif %}>Fases Anteriores (NÃO)</option>
                    </select>
                </div>
                <div class="col-md-2">
                    <label class="form-label small fw-bold mb-1 text-danger">Filtro de SLA</label>
                    <select name="sla_filtro" class="form-select form-select-sm border-danger">
                        <option value="Todos" {% if f_sla == 'Todos' %}selected{% endif %}>Todos os SLAs</option>
                        <option value="No Prazo" {% if f_sla == 'No Prazo' %}selected{% endif %}>No Prazo (0 a 1d)</option>
                        <option value="Atenção" {% if f_sla == 'Atenção' %}selected{% endif %}>Atenção (2 a 3d)</option>
                        <option value="Urgente" {% if f_sla == 'Urgente' %}selected{% endif %}>Urgente (4 a 5d)</option>
                        <option value="Crítico" {% if f_sla == 'Crítico' %}selected{% endif %}>Crítico (6+ dias)</option>
                        <option value="Parado" {% if f_sla == 'Parado' %}selected{% endif %}>SLA Parado (Deferido/Devolvido)</option>
                        <option value="90d" {% if f_sla == '90d' %}selected{% endif %}>⚠️ Alerta 90+ Dias (Ativos)</option>
                    </select>
                </div>
                <div class="col-md-1">
                    <label class="form-label small fw-bold mb-1">UF</label>
                    <select name="uf" class="form-select form-select-sm">
                        <option value="Todos">Todas</option>
                        {% for u in lista_ufs %}<option value="{{u}}" {% if f_uf == u %}selected{% endif %}>{{u}}</option>{% endfor %}
                    </select>
                </div>
                <div class="col-md-2">
                    <label class="form-label small fw-bold mb-1">Fase</label>
                    <select name="fase" class="form-select form-select-sm">
                        <option value="Todos">Todas as Fases</option>
                        {% for fa in lista_fases %}<option value="{{fa}}" {% if f_fase == fa %}selected{% endif %}>{{fa}}</option>{% endfor %}
                    </select>
                </div>
                <div class="col-md-1">
                    <label class="form-label small fw-bold mb-1">Status/Parecer</label>
                    <select name="parecer" class="form-select form-select-sm">
                        <option value="Todos">Todos</option>
                        {% for pa in lista_pareceres %}<option value="{{pa}}" {% if f_parecer == pa %}selected{% endif %}>{{pa}}</option>{% endfor %}
                    </select>
                </div>
                <div class="col-md-2">
                    <label class="form-label small fw-bold mb-1">Responsável (COCAP)</label>
                    <select name="resp" class="form-select form-select-sm">
                        <option value="Todos">Todos</option>
                        {% for re_ in lista_resps %}<option value="{{re_}}" {% if f_resp == re_ %}selected{% endif %}>{{re_}}</option>{% endfor %}
                    </select>
                </div>
                <div class="col-md-2">
                    <label class="form-label small fw-bold mb-1">Busca (CNPJ / Nome / Proc.)</label>
                    <div class="input-group input-group-sm">
                        <input type="text" name="busca" class="form-control" placeholder="Digite..." value="{{ f_busca }}">
                        <button type="submit" class="btn btn-primary fw-bold">🔍</button>
                        <a href="/credenciamento" class="btn btn-outline-secondary" title="Limpar Filtros">✖</a>
                    </div>
                </div>
            </form>
        </div>
    </div>

    <!-- Tabela de Processos de Credenciamento -->
    <div class="card shadow-sm border-0 mb-5">
        <div class="card-body">
            <div class="d-flex justify-content-between align-items-center mb-2">
                <h5 class="mb-0 text-secondary fw-bold">Processos de Credenciamento <small class="text-muted">(Duplo clique na linha para Editar Fase, Status, Responsável e Observação)</small></h5>
                <span class="badge bg-primary fs-6">{{ registros|length }} Registro(s)</span>
            </div>
            <div class="table-responsive">
                <table class="table table-hover table-bordered align-middle text-center" style="font-size: 0.82em; cursor: pointer;">
                    <thead class="table-dark">
                        <tr>
                            <th>Processo</th>
                            <th>SLA (Fase / Geral)</th>
                            <th>UF</th>
                            <th>CNPJ / CPF</th>
                            <th class="text-start">Prestador / Fantasia</th>
                            <th>Categoria / Tipo</th>
                            <th>Abertura</th>
                            <th class="text-start">Fase</th>
                            <th>Últ. Fase</th>
                            <th>Entrada Fase</th>
                            <th>Status / Parecer</th>
                            <th>Origem</th>
                            <th>Responsável</th>
                            <th>Ações</th>
                        </tr>
                    </thead>
                    <tbody>
                        {% for item in registros %}
                        {% set r = item.row %}
                        {% set s = item.sla %}
                        {% set st_up = (r[20] or '')|upper %}
                        <tr class="{{ s.cor_linha }}" ondblclick="document.getElementById('btn-edit-cred-{{r[0]}}').click();">
                            <td class="fw-bold text-primary">{{ r[8] }}</td>
                            <td>
                                {% if s.ativo %}
                                    <span class="badge bg-{{ 'danger' if s.dias_fase >= 6 else 'warning text-dark' if s.dias_fase >= 2 else 'secondary' }} mb-1">Fase: {{ s.dias_fase }}d</span><br>
                                    <span class="badge" style="background-color: {{ '#6f42c1' if s.estourou_90d else '#6c757d' }};">Geral: {{ s.dias_geral }}d {{ '⚠️' if s.estourou_90d else '' }}</span>
                                {% else %}
                                    <span class="badge bg-{{ 'secondary' if st_up=='DEVOLVIDO' else 'danger' if st_up in ('INDEFERIDO', 'CANCELADA') else 'success' }} mb-1">SLA Parado ({{ s.dias_fase }}d)</span><br>
                                    <small class="text-muted">Geral: {{ s.dias_geral }}d</small>
                                {% endif %}
                            </td>
                            <td><span class="badge bg-secondary">{{ r[1] }}</span></td>
                            <td>{{ r[2] }}</td>
                            <td class="text-start">
                                <div class="fw-bold">{{ r[3] }}</div>
                                {% if r[4] and r[4] != r[3] %}<small class="text-muted">{{ r[4] }}</small>{% endif %}
                                {% if r[28] %}<br><span class="badge bg-light text-dark border">{{ r[28] }}</span>{% endif %}
                            </td>
                            <td>
                                <span class="badge bg-info text-dark">{{ r[6] }}</span><br>
                                <small class="text-muted">{{ r[5] }}</small>
                            </td>
                            <td>{{ r[9] }}</td>
                            <td class="text-start fw-bold">{{ r[12] }}</td>
                            <td>
                                <span class="badge bg-{{ 'primary' if r[13] == 'SIM' else 'light text-muted border' }}">{{ r[13] }}</span>
                            </td>
                            <td>{{ r[14] }}</td>
                            <td>
                                <span class="badge bg-{{ 'success' if st_up in ('DEFERIDO', 'FINALIZADA', 'CONCLUÍDO') else 'dark' if st_up=='PENDENTE' else 'warning text-dark' if st_up in ('EM ANÁLISE', 'EM ANALISE', 'AGUARDANDO 90 DIAS', 'AGUARDANDO RETORNO') else 'danger' if st_up in ('INDEFERIDO', 'DEVOLVIDO') else 'secondary' }}">{{ r[20] }}</span>
                            </td>
                            <td><strong>{{ r[22] }}</strong></td>
                            <td>{{ r[11] }}</td>
                            <td class="text-nowrap">
                                <form method="POST" action="/credenciamento/acao" style="display:inline;">
                                    <input type="hidden" name="id" value="{{ r[0] }}">
                                    <button type="submit" name="acao" value="assumir" class="btn btn-warning btn-sm fw-bold" title="Assumir Processo">✋</button>
                                </form>
                                <button id="btn-edit-cred-{{r[0]}}" type="button" class="btn btn-primary btn-sm fw-bold" title="Editar Fase, Status, Responsável e Observação"
                                    data-id="{{ r[0] }}"
                                    data-proc="{{ r[8] }}"
                                    data-nome="{{ r[3] }}"
                                    data-cnpj="{{ r[2] }}"
                                    data-fase="{{ r[12] if r[12] else '' }}"
                                    data-status="{{ r[20] if r[20] else 'Em Análise' }}"
                                    data-resp="{{ r[11] if r[11] else 'Nenhum' }}"
                                    data-origem="{{ r[22] if r[22] else 'SEDE' }}"
                                    data-obs="{{ r[21] if r[21] else '' }}"
                                    onclick="abrirModalEditarCred(this)">✏️</button>
                                <button type="button" class="btn btn-outline-dark btn-sm fw-bold" onclick="verHistoricoProcesso('{{ r[8] }}')" title="Ver Linha do Tempo do Processo">🔍</button>
                            </td>
                        </tr>
                        {% else %}
                        <tr><td colspan="14" class="text-center text-muted py-4 fs-6">Nenhum processo encontrado. Carregue o relatório diário acima ou ajuste os filtros.</td></tr>
                        {% endfor %}
                    </tbody>
                </table>
            </div>
        </div>
    </div>
</div>

<!-- Modal de Edição (Duplo Clique na Linha de Credenciamento - Fases) -->
<div class="modal fade" id="modalEditarCred" tabindex="-1">
  <div class="modal-dialog modal-lg">
    <div class="modal-content">
      <div class="modal-header bg-dark text-white">
        <h5 class="modal-title fw-bold">✏️ Analisar / Editar Processo de Credenciamento</h5>
        <button type="button" class="btn-close btn-close-white" data-bs-dismiss="modal"></button>
      </div>
      <form method="POST" action="/credenciamento/editar">
      <div class="modal-body">
        <input type="hidden" name="id" id="edit_cred_id">
        <div class="alert alert-secondary py-2 mb-3">
            <strong>Processo Nº:</strong> <span id="lbl_cred_proc" class="text-primary fw-bold"></span> |
            <strong>CNPJ/CPF:</strong> <span id="lbl_cred_cnpj"></span><br>
            <strong>Prestador:</strong> <span id="lbl_cred_nome" class="fw-bold"></span>
        </div>
        <div class="row g-3">
            <div class="col-md-12">
                <label class="form-label fw-bold text-primary">Fase do Credenciamento</label>
                <select name="fase" id="edit_cred_fase" class="form-select border-primary" required>
                    {% for fa in opcoes_fases_modal %}
                    <option value="{{ fa }}">{{ fa }}</option>
                    {% endfor %}
                </select>
            </div>
            <div class="col-md-4">
                <label class="form-label fw-bold text-success">Status / Parecer</label>
                <select name="status" id="edit_cred_status" class="form-select border-success" required>
                    {% for st in status_cred_modal %}
                    <option value="{{ st }}">{{ st }}</option>
                    {% endfor %}
                </select>
            </div>
            <div class="col-md-4">
                <label class="form-label fw-bold text-primary">Responsável (Usuários COCAP)</label>
                <select name="responsavel" id="edit_cred_resp" class="form-select border-primary">
                    {% for u in usuarios_login %}
                    <option value="{{ u }}">{{ u }}</option>
                    {% endfor %}
                </select>
            </div>
            <div class="col-md-4">
                <label class="form-label fw-bold">Origem / Alçada</label>
                <select name="origem_parecer" id="edit_cred_origem" class="form-select">
                    <option value="SEDE">SEDE</option>
                    <option value="REGIONAL">REGIONAL</option>
                    <option value="URR/NAR">URR/NAR</option>
                </select>
            </div>
            <div class="col-md-12">
                <label class="form-label fw-bold">Observação / Parecer Técnico</label>
                <textarea name="observacao" id="edit_cred_obs" class="form-control" rows="5" placeholder="Digite as observações da análise..."></textarea>
            </div>
        </div>
      </div>
      <div class="modal-footer">
        <button type="button" class="btn btn-secondary" data-bs-dismiss="modal">Cancelar</button>
        <button type="submit" class="btn btn-success fw-bold">💾 Salvar Alterações</button>
      </div>
      </form>
    </div>
  </div>
</div>

<!-- Modal Linha do Tempo e Detalhes do Processo -->
<div class="modal fade" id="modalHistoricoCred" tabindex="-1">
  <div class="modal-dialog modal-xl modal-dialog-scrollable">
    <div class="modal-content">
      <div class="modal-header bg-dark text-white">
        <h5 class="modal-title fw-bold" id="tituloModalCred">📋 Histórico Completo do Processo</h5>
        <button type="button" class="btn-close btn-close-white" data-bs-dismiss="modal"></button>
      </div>
      <div class="modal-body" id="corpoModalCred">
        <div class="text-center py-4"><div class="spinner-border text-primary"></div><p class="mt-2">A carregar linha do tempo...</p></div>
      </div>
      <div class="modal-footer">
        <button type="button" class="btn btn-secondary fw-bold" data-bs-dismiss="modal">Fechar</button>
      </div>
    </div>
  </div>
</div>

<script src="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/js/bootstrap.bundle.min.js"></script>
<script>
function selecionarOuCriarOpcao(selectEl, valor) {
    if (!valor) return;
    let existe = Array.from(selectEl.options).some(opt => opt.value === valor);
    if (!existe) {
        selectEl.options.add(new Option(valor, valor));
    }
    selectEl.value = valor;
}

function selecionarApenasCadastrados(selectEl, valor) {
    let existe = Array.from(selectEl.options).some(opt => opt.value === valor);
    selectEl.value = existe ? valor : 'Nenhum';
}

function abrirModalEditarCred(btn) {
    document.getElementById('edit_cred_id').value = btn.getAttribute('data-id');
    document.getElementById('lbl_cred_proc').innerText = btn.getAttribute('data-proc') || '-';
    document.getElementById('lbl_cred_cnpj').innerText = btn.getAttribute('data-cnpj') || '-';
    document.getElementById('lbl_cred_nome').innerText = btn.getAttribute('data-nome') || '-';

    selecionarOuCriarOpcao(document.getElementById('edit_cred_fase'), btn.getAttribute('data-fase'));
    selecionarOuCriarOpcao(document.getElementById('edit_cred_status'), btn.getAttribute('data-status'));
    // Mantém estritamente os usuários cadastrados na tabela users
    selecionarApenasCadastrados(document.getElementById('edit_cred_resp'), btn.getAttribute('data-resp'));
    selecionarOuCriarOpcao(document.getElementById('edit_cred_origem'), btn.getAttribute('data-origem'));

    let obsBruta = btn.getAttribute('data-obs') || '';
    document.getElementById('edit_cred_obs').value = obsBruta.replace(/--/g, '\\n');

    new bootstrap.Modal(document.getElementById('modalEditarCred')).show();
}

function verHistoricoProcesso(numProcesso) {
    let modalEl = document.getElementById('modalHistoricoCred');
    let modal = new bootstrap.Modal(modalEl);
    document.getElementById('tituloModalCred').innerText = '📋 Linha do Tempo - Processo Nº ' + numProcesso;
    let corpo = document.getElementById('corpoModalCred');
    corpo.innerHTML = '<div class="text-center py-4"><div class="spinner-border text-primary"></div><p class="mt-2">A procurar histórico do processo...</p></div>';
    modal.show();

    fetch('/api/credenciamento/processo/' + encodeURIComponent(numProcesso))
    .then(r => r.json())
    .then(data => {
        if (data.erro || !data.etapas || data.etapas.length === 0) {
            corpo.innerHTML = '<div class="alert alert-warning">Não foi possível carregar os detalhes deste processo.</div>';
            return;
        }
        let info = data.etapas[0];
        let html = `
            <div class="card bg-light border-0 mb-3">
                <div class="card-body">
                    <div class="row">
                        <div class="col-md-6">
                            <h5 class="fw-bold text-primary mb-1">${info.nome}</h5>
                            <p class="mb-1 text-muted small"><strong>Fantasia:</strong> ${info.fantasia || '-'} | <strong>CNPJ/CPF:</strong> ${info.cnpj_cpf}</p>
                            <p class="mb-0 small"><strong>Tipo Prestador:</strong> ${info.tipo_prestador} | <strong>Categoria:</strong> ${info.categoria}</p>
                        </div>
                        <div class="col-md-6 text-md-end">
                            <span class="badge bg-dark fs-6">Processo: ${info.processo}</span>
                            <span class="badge bg-secondary fs-6">UF: ${info.filial}</span>
                            <p class="mb-0 small mt-2"><strong>Data de Abertura:</strong> ${info.data_abertura} | <strong>Grupo:</strong> ${info.grupo_prestadores || '-'}</p>
                        </div>
                    </div>
                </div>
            </div>
            <h6 class="fw-bold text-secondary mb-3">⏳ Etapas e Pareceres Registados (${data.etapas.length}):</h6>
        `;

        data.etapas.forEach((et, idx) => {
            let stUp = (et.parecer || '').toUpperCase();
            let badgeCor = (stUp === 'DEFERIDO' || stUp === 'FINALIZADA') ? 'success' : (stUp === 'EM ANÁLISE' || stUp === 'EM ANALISE' || stUp === 'PENDENTE' ? 'warning text-dark' : (stUp === 'INDEFERIDO' || stUp === 'DEVOLVIDO' ? 'danger' : 'secondary'));
            let destaqueAtual = et.ultima_fase === 'SIM' ? 'border-primary border-2' : '';
            let obsFormatada = (et.observacao || 'Sem observações registadas.').replace(/--/g, '\\n');
            html += `
                <div class="card shadow-sm mb-3 ${destaqueAtual}">
                    <div class="card-header d-flex justify-content-between align-items-center ${et.ultima_fase === 'SIM' ? 'bg-primary text-white' : 'bg-white'}">
                        <div>
                            <strong>#${idx + 1} - ${et.fase}</strong>
                            ${et.ultima_fase === 'SIM' ? '<span class="badge bg-warning text-dark ms-2">⭐ FASE ATUAL</span>' : ''}
                        </div>
                        <div>
                            <span class="badge bg-${badgeCor}">${et.parecer || '-'}</span>
                            <span class="badge bg-dark ms-1">${et.origem_parecer || '-'}</span>
                        </div>
                    </div>
                    <div class="card-body">
                        <div class="row small text-muted mb-2">
                            <div class="col-md-3"><strong>Responsável:</strong> ${et.usuario_resp || '-'}</div>
                            <div class="col-md-3"><strong>Início / Fim:</strong> ${et.data_inicial || '-'} a ${et.data_final || '-'}</div>
                            <div class="col-md-3"><strong>Dias (Relatório):</strong> ${et.dias_percorridos || '0'}d / ${et.dias_percorridos_pf || '0'}d</div>
                            <div class="col-md-3"><strong>Próx. Fase:</strong> ${et.proxima_fase || '-'}</div>
                        </div>
                        ${et.justificativa_indef ? `<p class="mb-2 small"><strong>Motivo / Enquadramento:</strong> <span class="text-danger fw-bold">${et.justificativa_indef}</span></p>` : ''}
                        <div class="p-3 bg-light rounded border small" style="white-space: pre-wrap;">${obsFormatada}</div>
                        ${(et.data_aprovacao || et.usuario_aprj || et.justificativa_corgen) ? `
                        <div class="mt-2 pt-2 border-top small text-secondary">
                            <strong>Aprovação Sede:</strong> ${et.data_aprovacao || '-'} | <strong>Aprovador:</strong> ${et.usuario_aprj || '-'} | <strong> Despacho:</strong> ${et.justificativa_corgen || '-'}
                        </div>` : ''}
                    </div>
                </div>
            `;
        });
        corpo.innerHTML = html;
    })
    .catch(e => {
        corpo.innerHTML = '<div class="alert alert-danger">Erro ao comunicar com o servidor.</div>';
    });
}
</script>
</body></html>"""

TELA_ADMIN = """<!DOCTYPE html><html lang="pt-BR"><head><meta charset="UTF-8"><title>Painel Administrativo</title><link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet"></head>
<body style="background-color: #f0f2f5;">
<nav class="navbar navbar-dark" style="background-color: #333;"><div class="container-fluid"><a href="/dashboard" class="btn btn-outline-light btn-sm">⬅ Voltar ao Portal</a><span class="navbar-text text-white fw-bold">Painel Administrativo</span></div></nav>
<div class="container mt-4">
    {% with messages = get_flashed_messages(with_categories=true) %}{% if messages %}{% for category, message in messages %}<div class="alert alert-{{ category }}">{{ message }}</div>{% endfor %}{% endif %}{% endwith %}
    <div class="row">
        <div class="col-md-12 mb-4">
            <div class="card shadow-sm border-0">
                <div class="card-header bg-dark text-white fw-bold">💾 Gestão de Banco de Dados</div>
                <div class="card-body">
                    <div class="row text-center">
                        <div class="col-md-6 border-end">
                            <h5 class="text-success">Exportar Backup</h5><a href="/admin/exportar" class="btn btn-success fw-bold w-75 mt-2">⬇ Baixar Backup (Excel)</a>
                        </div>
                        <div class="col-md-6">
                            <h5 class="text-danger">Restaurar ou Migrar Banco de Dados</h5>
                            <form action="/admin/importar" method="POST" enctype="multipart/form-data" class="d-flex flex-column align-items-center mt-2">
                                <input type="file" name="file_backup" class="form-control form-control-sm w-75 mb-2" accept=".xlsx,.db" required>
                                <button type="submit" class="btn btn-danger fw-bold w-75" onclick="return confirm('ATENÇÃO: Todas as informações da Nuvem serão APAGADAS. Deseja continuar?');">⬆ Importar Arquivo</button>
                            </form>
                        </div>
                    </div>
                </div>
            </div>
        </div>
        <div class="col-md-12">
            <div class="card shadow-sm mb-4 border-0">
                <div class="card-body"><h5 class="card-title mb-3">Criar Novo Usuário</h5>
                    <form method="POST" action="/admin"><input type="hidden" name="acao" value="adicionar">
                        <div class="row g-2 align-items-end">
                            <div class="col-md-4"><label class="form-label small fw-bold">Usuário</label><input type="text" name="username" class="form-control" required></div>
                            <div class="col-md-3"><label class="form-label small fw-bold">Senha</label><input type="password" name="password" class="form-control" required></div>
                            <div class="col-md-3"><label class="form-label small fw-bold">Nível</label><select name="role" class="form-select"><option value="user">User</option><option value="admin">Admin</option></select></div>
                            <div class="col-md-2"><button type="submit" class="btn btn-primary w-100 fw-bold">Salvar</button></div>
                        </div>
                    </form>
                </div>
            </div>
            <div class="card shadow-sm border-0">
                <div class="card-body">
                    <table class="table table-hover align-middle"><thead class="table-light"><tr><th>ID</th><th>Usuário</th><th>Nível</th><th>Ações</th></tr></thead>
                        <tbody>{% for u in usuarios %}<tr><td>{{ u[0] }}</td><td>{{ u[1] }}</td><td><span class="badge bg-{{ 'dark' if u[2] == 'admin' else 'secondary' }}">{{ u[2] | upper }}</span></td><td>
                            {% if u[1] != 'admin' and u[1] != user %}<form method="POST" action="/admin" style="display:inline;"><input type="hidden" name="acao" value="eliminar"><input type="hidden" name="user_id" value="{{ u[0] }}"><button type="submit" class="btn btn-outline-danger btn-sm" onclick="return confirm('Excluir o usuário {{ u[1] }}?');">Excluir</button></form>
                            {% else %}<span class="text-muted small">Sistema</span>{% endif %}</td></tr>{% endfor %}
                        </tbody>
                    </table>
                </div>
            </div>
        </div>
    </div>
</div></body></html>"""

TELA_AVULSAS = """<!DOCTYPE html><html lang="pt-BR"><head><title>Demandas Avulsas</title><link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
<style>
    .bg-atencao { background-color: #fff3cd !important; }
    .bg-urgente { background-color: #ffe8cc !important; }
    .bg-critico { background-color: #f8d7da !important; }
</style>
</head>
<body style="background-color: #f0f2f5;">
<nav class="navbar navbar-dark bg-dark"><div class="container-fluid"><a href="/demandas" class="btn btn-outline-light btn-sm">⬅ Voltar às Demandas</a><span class="text-white fw-bold">Demandas Avulsas</span><a href="/demandas/dashboard_graficos" class="btn btn-outline-info btn-sm">📊 Dashboard</a></div></nav>
<div class="container-fluid mt-4">
    {% with messages = get_flashed_messages(with_categories=true) %}{% if messages %}{% for category, message in messages %}<div class="alert alert-{{ category }}">{{ message }}</div>{% endfor %}{% endif %}{% endwith %}
    
    <div class="card shadow-sm mb-3 border-0 bg-light">
        <div class="card-body py-2">
            <form method="GET" action="/demandas/avulsas" class="row g-2 align-items-center">
                <div class="col-md-3"><select name="status" class="form-select form-select-sm"><option value="Todos">Status (Todos)</option>{% for s in status_list %}<option value="{{s}}" {% if s == atual_status %}selected{% endif %}>{{s}}</option>{% endfor %}</select></div>
                <div class="col-md-3"><select name="tipo" class="form-select form-select-sm"><option value="Todos">Tipo (Todos)</option>{% for t in tipos %}<option value="{{t}}" {% if t == atual_tipo %}selected{% endif %}>{{t}}</option>{% endfor %}</select></div>
                <div class="col-md-3"><select name="resp" class="form-select form-select-sm"><option value="Todos">Responsável (Todos)</option>{% for u in usuarios %}<option value="{{u}}" {% if u == atual_resp %}selected{% endif %}>{{u}}</option>{% endfor %}</select></div>
                <div class="col-md-3"><button type="submit" class="btn btn-secondary btn-sm w-100 fw-bold">🔍 Filtrar Resultados</button></div>
            </form>
        </div>
    </div>

    <div class="card shadow-sm mb-4 border-0">
        <div class="card-body">
            <form method="POST" action="/demandas/avulsas"><input type="hidden" name="acao" value="nova">
                <div class="row g-2">
                    <div class="col-md-2"><select name="tipo" class="form-select" required><option value="">Tipo...</option>{% for t in tipos %}<option value="{{t}}">{{t}}</option>{% endfor %}</select></div>
                    <div class="col-md-3"><input type="text" name="assunto" class="form-control" placeholder="Assunto/Demanda" required></div>
                    <div class="col-md-2"><select name="solicitante" class="form-select" required><option value="">Solicitante...</option>{% for s in sols %}<option value="{{s}}">{{s}}</option>{% endfor %}</select></div>
                    <div class="col-md-2"><input type="date" name="prazo" class="form-control" required></div>
                    <div class="col-md-2"><button type="submit" class="btn btn-success w-100 fw-bold">+ Nova Demanda</button></div>
                </div>
            </form>
        </div>
    </div>
    
    <div class="card shadow-sm border-0"><div class="card-body">
        <div class="d-flex justify-content-between align-items-center mb-3">
            <h5 class="mb-0 text-secondary fw-bold">Lista de Demandas <small class="text-muted">(Duplo clique p/ Editar)</small></h5>
            <span class="badge bg-primary fs-6">{{ demandas|length }} Registro(s) Encontrado(s)</span>
        </div>
        <div class="table-responsive">
        <table class="table table-hover table-bordered align-middle text-center" style="font-size: 0.9em; cursor: pointer;">
            <thead class="table-dark"><tr><th>ID</th><th>SLA (Dias)</th><th>Tipo</th><th>Assunto</th><th>Solicitante</th><th>Entrada</th><th>Prazo</th><th>Status</th><th>Resp.</th><th>Conclusão</th><th>Ações</th></tr></thead>
            <tbody>{% for d in demandas %}
            {% set sla = calc_sla(d[4], d[6]) %}
            {% set dias_fech = calc_fechados(d[4], d[8]) %}
            {% set row_class = '' %}
            {% if sla >= 6 %}{% set row_class = 'bg-critico' %}
            {% elif sla >= 4 %}{% set row_class = 'bg-urgente' %}
            {% elif sla >= 2 %}{% set row_class = 'bg-atencao' %}
            {% endif %}
            <tr class="{{ row_class }}" ondblclick="document.getElementById('btn-edit-avulsa-{{d[0]}}').click();">
                <td class="fw-bold">{{ d[0] }}</td>
                <td>
                    {% if sla >= 0 %}
                        <span class="badge bg-{{ 'danger' if sla >= 6 else 'warning text-dark' if sla >= 4 else 'secondary' }}">{{ sla }} dia(s)</span>
                    {% else %}
                        <span class="badge bg-{{ 'secondary' if d[6]=='Devolvido' else 'success' }}">Parado {{ '(' ~ dias_fech ~ 'd)' if dias_fech is not none else '' }}</span>
                    {% endif %}
                </td>
                <td>{{ d[1] }}</td><td class="text-start">{{ d[2] }}</td><td>{{ d[3] }}</td><td>{{ d[4] }}</td><td>{{ d[5] }}</td>
                <td><span class="badge bg-{{ 'dark' if d[6]=='Pendente' else 'success' if d[6] in ('Concluído', 'Finalizada') else 'danger' if d[6]=='Devolvido' else 'warning text-dark' if d[6]=='Aguardando 90 dias' else 'info' }}">{{ d[6] }}</span></td>
                <td>{{ d[7] }}</td><td>{{ d[8] }}</td>
                <td>
                    <form method="POST" action="/demandas/avulsas" style="display:inline;"><input type="hidden" name="id" value="{{ d[0] }}">
                        <button type="submit" name="acao" value="assumir" class="btn btn-warning btn-sm fw-bold" {% if d[6] in ('Concluído', 'Devolvido') %}disabled{% endif %} title="Assumir">✋</button>
                        <button type="submit" name="acao" value="concluir" class="btn btn-success btn-sm fw-bold" {% if d[6] in ('Concluído', 'Devolvido') %}disabled{% endif %} title="Concluir">✔</button>
                    </form>
                    <button id="btn-edit-avulsa-{{d[0]}}" type="button" class="btn btn-primary btn-sm fw-bold" title="Editar" 
                        data-id="{{ d[0] }}" data-tipo="{{ d[1] }}" data-assunto="{{ d[2] }}" data-sol="{{ d[3] }}" 
                        data-prazo="{{ d[5] }}" data-status="{{ d[6] }}" data-resp="{{ d[7] }}" data-obs="{{ d[9] if d[9] else '' }}"
                        onclick="abrirModalAvulsa(this)">✏️</button>
                </td>
            </tr>{% endfor %}</tbody>
        </table>
    </div></div></div>
</div>

<!-- Modal Edição Avulsa -->
<div class="modal fade" id="modalEditAvulsa" tabindex="-1">
  <div class="modal-dialog modal-lg">
    <div class="modal-content">
      <div class="modal-header bg-dark text-white">
        <h5 class="modal-title fw-bold">✏️ Editar Demanda Avulsa</h5>
        <button type="button" class="btn-close btn-close-white" data-bs-dismiss="modal"></button>
      </div>
      <form method="POST" action="/demandas/avulsas/editar">
      <div class="modal-body">
        <input type="hidden" name="id" id="edit_av_id">
        <div class="row g-3">
            <div class="col-md-6"><label class="form-label fw-bold">Tipo</label>
                <select name="tipo" id="edit_av_tipo" class="form-select">{% for t in tipos %}<option value="{{t}}">{{t}}</option>{% endfor %}</select>
            </div>
            <div class="col-md-6"><label class="form-label fw-bold">Status</label>
                <select name="status" id="edit_av_status" class="form-select">{% for s in status_list %}<option value="{{s}}">{{s}}</option>{% endfor %}</select>
            </div>
            <div class="col-md-12"><label class="form-label fw-bold">Assunto</label>
                <input type="text" name="assunto" id="edit_av_assunto" class="form-control">
            </div>
            <div class="col-md-4"><label class="form-label fw-bold">Solicitante</label>
                <select name="solicitante" id="edit_av_sol" class="form-select">{% for s in sols %}<option value="{{s}}">{{s}}</option>{% endfor %}</select>
            </div>
            <div class="col-md-4"><label class="form-label fw-bold">Responsável</label>
                <select name="responsavel" id="edit_av_resp" class="form-select">{% for u in usuarios if u != 'Todos' %}<option value="{{u}}">{{u}}</option>{% endfor %}</select>
            </div>
            <div class="col-md-4"><label class="form-label fw-bold">Prazo</label>
                <input type="date" name="prazo" id="edit_av_prazo" class="form-control">
            </div>
            <div class="col-md-12"><label class="form-label fw-bold">Observações</label>
                <textarea name="observacao" id="edit_av_obs" class="form-control" rows="3"></textarea>
            </div>
        </div>
      </div>
      <div class="modal-footer">
        <button type="button" class="btn btn-secondary" data-bs-dismiss="modal">Cancelar</button>
        <button type="submit" class="btn btn-success fw-bold">Salvar Alterações</button>
      </div>
      </form>
    </div>
  </div>
</div>
<script src="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/js/bootstrap.bundle.min.js"></script>
<script>
function abrirModalAvulsa(btn) {
    document.getElementById('edit_av_id').value = btn.getAttribute('data-id');
    document.getElementById('edit_av_tipo').value = btn.getAttribute('data-tipo');
    document.getElementById('edit_av_assunto').value = btn.getAttribute('data-assunto');
    document.getElementById('edit_av_sol').value = btn.getAttribute('data-sol');
    document.getElementById('edit_av_prazo').value = btn.getAttribute('data-prazo');
    document.getElementById('edit_av_status').value = btn.getAttribute('data-status');
    document.getElementById('edit_av_resp').value = btn.getAttribute('data-resp');
    document.getElementById('edit_av_obs').value = btn.getAttribute('data-obs');
    new bootstrap.Modal(document.getElementById('modalEditAvulsa')).show();
}
</script>
</body></html>"""

TELA_ROTINAS = """<!DOCTYPE html><html lang="pt-BR"><head><title>Rotinas (PDFs)</title><link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
<style>
    .bg-atencao { background-color: #fff3cd !important; }
    .bg-urgente { background-color: #ffe8cc !important; }
    .bg-critico { background-color: #f8d7da !important; }
</style>
</head>
<body style="background-color: #f0f2f5;">
<nav class="navbar navbar-dark bg-dark"><div class="container-fluid"><a href="/demandas" class="btn btn-outline-light btn-sm">⬅ Voltar às Demandas</a><span class="text-white fw-bold">Processamento de PDFs</span><a href="/demandas/dashboard_graficos" class="btn btn-outline-info btn-sm">📊 Dashboard</a></div></nav>
<div class="container-fluid mt-4">
    {% with messages = get_flashed_messages(with_categories=true) %}{% if messages %}{% for category, message in messages %}<div class="alert alert-{{ category }}">{{ message }}</div>{% endfor %}{% endif %}{% endwith %}

    <div class="card shadow-sm mb-3 border-0 bg-light">
        <div class="card-body py-2">
            <form method="GET" action="/demandas/rotinas" class="row g-2 align-items-center">
                <div class="col-md-3"><select name="status" class="form-select form-select-sm"><option value="Todos">Status (Todos)</option>{% for s in status_list %}<option value="{{s}}" {% if atual_status == s %}selected{% endif %}>{{s}}</option>{% endfor %}</select></div>
                <div class="col-md-3"><select name="tipo" class="form-select form-select-sm"><option value="Todos">Tipo (Todos)</option>{% for t in tipos %}<option value="{{t}}" {% if t == atual_tipo %}selected{% endif %}>{{t}}</option>{% endfor %}</select></div>
                <div class="col-md-3"><select name="resp" class="form-select form-select-sm"><option value="Todos">Responsável (Todos)</option>{% for u in usuarios %}<option value="{{u}}" {% if u == atual_resp %}selected{% endif %}>{{u}}</option>{% endfor %}</select></div>
                <div class="col-md-3"><button type="submit" class="btn btn-secondary btn-sm w-100 fw-bold">🔍 Filtrar Resultados</button></div>
            </form>
        </div>
    </div>

    <div class="card shadow-sm mb-4 border-0">
        <div class="card-body text-center bg-light rounded">
            <form method="POST" action="/demandas/rotinas/upload" enctype="multipart/form-data" class="d-flex justify-content-center align-items-center gap-3">
                <span class="text-primary fw-bold">Importar PDFs:</span>
                <input type="file" name="pdfs" class="form-control w-25" multiple accept=".pdf" required>
                <button type="submit" class="btn btn-primary fw-bold px-4">Processar Arquivos</button>
            </form>
        </div>
    </div>
    
    <div class="card shadow-sm border-0"><div class="card-body">
        <div class="d-flex justify-content-between align-items-center mb-3">
            <h5 class="mb-0 text-secondary fw-bold">Lista de Processamentos <small class="text-muted">(Duplo clique p/ Editar)</small></h5>
            <span class="badge bg-primary fs-6">{{ rotinas|length }} Registro(s) Encontrado(s)</span>
        </div>
        <div class="table-responsive">
        <table class="table table-hover table-bordered align-middle text-center" style="font-size: 0.85em; cursor: pointer;">
            <thead class="table-dark"><tr><th>ID</th><th>SLA (COCAP / FOP)</th><th>Data Entrada</th><th>Prestador</th><th>CNPJ</th><th>Mun/UF</th><th>Demanda</th><th>Status</th><th>Resp.</th><th>Ações</th></tr></thead>
            <tbody>{% for d in rotinas %}
            {% set sla_cocap = calc_sla(d[9], d[3]) %}
            {% set sla_geral = calc_sla(d[18], d[3]) %}
            {% set dias_fech_cocap = calc_fechados(d[9], d[19]) %}
            {% set dias_fech_geral = calc_fechados(d[18], d[19]) %}
            {% set row_class = '' %}
            {% if sla_cocap >= 6 %}{% set row_class = 'bg-critico' %}
            {% elif sla_cocap >= 4 %}{% set row_class = 'bg-urgente' %}
            {% elif sla_cocap >= 2 %}{% set row_class = 'bg-atencao' %}
            {% endif %}
            <tr class="{{ row_class }}" ondblclick="document.getElementById('btn-edit-rot-{{d[0]}}').click();">
                <td class="fw-bold">{{ d[0] }}</td>
                <td>
                    {% if sla_cocap >= 0 %}
                        <span class="badge bg-{{ 'danger' if sla_cocap >= 6 else 'warning text-dark' if sla_cocap >= 4 else 'secondary' }} mb-1">COCAP: {{ sla_cocap }}d</span><br>
                    {% else %}
                        <span class="badge bg-{{ 'secondary' if d[3]=='Devolvido' else 'success' }} mb-1">COCAP: Parado {{ '(' ~ dias_fech_cocap ~ 'd)' if dias_fech_cocap is not none else '' }}</span><br>
                    {% endif %}
                    
                    {% if sla_geral >= 0 %}
                        <span class="badge bg-{{ 'danger' if sla_geral >= 15 else 'warning text-dark' if sla_geral >= 10 else 'secondary' }}">Geral: {{ sla_geral }}d</span>
                    {% elif d[18] and d[3] in ['Finalizada', 'Concluído', 'Devolvido', 'Cancelada'] %}
                        <span class="badge bg-{{ 'secondary' if d[3]=='Devolvido' else 'success' }}">Geral: Parado {{ '(' ~ dias_fech_geral ~ 'd)' if dias_fech_geral is not none else '' }}</span>
                    {% else %}
                        <span class="badge bg-light text-muted border">Geral: -</span>
                    {% endif %}
                </td>
                <td>{{ d[9] }}</td><td class="text-start fw-bold">{{ d[5] }}</td><td>{{ d[6] }}</td><td>{{ d[7] }}/{{ d[8] }}</td><td>{{ d[2] }}</td>
                <td><span class="badge bg-{{ 'dark' if d[3]=='Pendente' else 'success' if d[3]=='Finalizada' else 'warning text-dark' if d[3]=='Aguardando 90 dias' else 'danger' if d[3]=='Devolvido' else 'info' }}">{{ d[3] }}</span></td>
                <td>{{ d[4] }}</td>
                <td>
                    <form method="POST" action="/demandas/rotinas/acao" style="display:inline;"><input type="hidden" name="id" value="{{ d[0] }}">
                        <button type="submit" name="acao" value="assumir" class="btn btn-warning btn-sm fw-bold" {% if d[3] in ('Finalizada', 'Devolvido') %}disabled{% endif %} title="Assumir">✋</button>
                        <button type="submit" name="acao" value="finalizar" class="btn btn-success btn-sm fw-bold" {% if d[3] in ('Finalizada', 'Devolvido') %}disabled{% endif %} title="Finalizar">✔</button>
                    </form>
                    <button id="btn-edit-rot-{{d[0]}}" type="button" class="btn btn-primary btn-sm fw-bold" title="Editar Informações"
                        data-id="{{ d[0] }}" data-status="{{ d[3] }}" data-tipo="{{ d[11] if d[11] else '' }}" data-alcada="{{ d[12] if d[12] else '' }}" 
                        data-perc="{{ d[13] if d[13] else '' }}" data-rs="{{ d[14] if d[14] else '' }}" data-nup="{{ d[15] if d[15] else '' }}"
                        data-proc="{{ d[16] if d[16] else '' }}" data-fop="{{ d[17] if d[17] else '' }}" 
                        data-fop-date="{{ d[18] if d[18] else '' }}" data-obs="{{ d[22] if d[22] else '' }}"
                        data-resp="{{ d[4] }}"
                        onclick="abrirModalRotina(this)">✏️</button>
                </td>
            </tr>{% endfor %}</tbody>
        </table>
    </div></div></div>
</div>

<!-- Modal Edição Rotina -->
<div class="modal fade" id="modalEditRotina" tabindex="-1">
  <div class="modal-dialog modal-lg">
    <div class="modal-content">
      <div class="modal-header bg-dark text-white">
        <h5 class="modal-title fw-bold">✏️ Completar Informações da Demanda</h5>
        <button type="button" class="btn-close btn-close-white" data-bs-dismiss="modal"></button>
      </div>
      <form method="POST" action="/demandas/rotinas/editar">
      <div class="modal-body">
        <input type="hidden" name="id" id="edit_rot_id">
        <div class="row g-3">
            <div class="col-md-4"><label class="form-label fw-bold text-primary">Status da Demanda</label>
                <select name="status" id="edit_rot_status" class="form-select border-primary">{% for s in status_list %}<option value="{{s}}">{{s}}</option>{% endfor %}</select>
            </div>
            <div class="col-md-4"><label class="form-label fw-bold text-primary">Responsável Atual</label>
                <select name="responsavel" id="edit_rot_resp" class="form-select border-primary">{% for u in usuarios if u != 'Todos' %}<option value="{{u}}">{{u}}</option>{% endfor %}</select>
            </div>
            <div class="col-md-4"><label class="form-label fw-bold text-danger">Data Assin. FOP (SLA Geral)</label>
                <input type="date" name="data_aprovacao" id="edit_rot_data_fop" class="form-control border-danger">
            </div>
            <div class="col-md-4"><label class="form-label fw-bold">NUP / e-Doc</label>
                <input type="text" name="nup_edoc" id="edit_rot_nup" class="form-control">
            </div>
            <div class="col-md-4"><label class="form-label fw-bold">Tipo Negociação</label>
                <select name="tipo_negociacao" id="edit_rot_tipo" class="form-select">
                    <option value="">Selecione...</option><option value="DIFERENCIADA">DIFERENCIADA</option><option value="MISTA">MISTA</option><option value="ESTADO">ESTADO</option><option value="NÃO IDENTIFICADA">NÃO IDENTIFICADA</option>
                </select>
            </div>
            <div class="col-md-4"><label class="form-label fw-bold">Alçada</label>
                <select name="alcada" id="edit_rot_alcada" class="form-select">
                    <option value="">Selecione...</option><option value="GERED">GERED</option><option value="REGIONAL">REGIONAL</option><option value="CONEC">CONEC</option><option value="DIOPE">DIOPE</option><option value="CODER">CODER</option>
                </select>
            </div>
            <div class="col-md-4"><label class="form-label fw-bold">Tipo Processo</label>
                <select name="tipo_processo" id="edit_rot_proc" class="form-select">
                    <option value="">Selecione...</option><option value="PARAMETRIZADO">PARAMETRIZADO</option><option value="AGUARDANDO 90 DIAS">AGUARDANDO 90 DIAS</option><option value="AGUARDANDO RETORNO">AGUARDANDO RETORNO</option><option value="DEVOLVIDO">DEVOLVIDO</option><option value="NÃO LOCALIZADO">NÃO LOCALIZADO</option>
                </select>
            </div>
            <div class="col-md-4"><label class="form-label fw-bold">FOP Autorização (Nº)</label>
                <input type="text" name="fop_autorizacao" id="edit_rot_fop" class="form-control">
            </div>
            <div class="col-md-2"><label class="form-label fw-bold">Impacto (%)</label>
                <input type="text" name="impacto_perc" id="edit_rot_perc" class="form-control" placeholder="Ex: 5,4%">
            </div>
            <div class="col-md-2"><label class="form-label fw-bold">Impacto (R$)</label>
                <input type="text" name="impacto_rs" id="edit_rot_rs" class="form-control" placeholder="Ex: 15000,00">
            </div>
            <div class="col-md-12"><label class="form-label fw-bold">Observações</label>
                <textarea name="observacao" id="edit_rot_obs" class="form-control" rows="3"></textarea>
            </div>
        </div>
      </div>
      <div class="modal-footer">
        <button type="button" class="btn btn-secondary" data-bs-dismiss="modal">Cancelar</button>
        <button type="submit" class="btn btn-success fw-bold">Salvar Informações</button>
      </div>
      </form>
    </div>
  </div>
</div>
<script src="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/js/bootstrap.bundle.min.js"></script>
<script>
function abrirModalRotina(btn) {
    document.getElementById('edit_rot_id').value = btn.getAttribute('data-id');
    document.getElementById('edit_rot_status').value = btn.getAttribute('data-status') || 'Pendente';
    document.getElementById('edit_rot_tipo').value = btn.getAttribute('data-tipo');
    document.getElementById('edit_rot_alcada').value = btn.getAttribute('data-alcada');
    document.getElementById('edit_rot_perc').value = btn.getAttribute('data-perc');
    document.getElementById('edit_rot_rs').value = btn.getAttribute('data-rs');
    document.getElementById('edit_rot_nup').value = btn.getAttribute('data-nup');
    document.getElementById('edit_rot_proc').value = btn.getAttribute('data-proc');
    document.getElementById('edit_rot_fop').value = btn.getAttribute('data-fop');
    
    let dStr = btn.getAttribute('data-fop-date');
    if (dStr && dStr.includes('/')) {
        let parts = dStr.split(' ')[0].split('/');
        if (parts.length === 3) dStr = parts[2] + '-' + parts[1] + '-' + parts[0];
    }
    document.getElementById('edit_rot_data_fop').value = dStr;
    document.getElementById('edit_rot_obs').value = btn.getAttribute('data-obs');
    document.getElementById('edit_rot_resp').value = btn.getAttribute('data-resp');
    new bootstrap.Modal(document.getElementById('modalEditRotina')).show();
}
</script>
</body></html>"""

TELA_DASHBOARD_GRAFICOS = """<!DOCTYPE html><html lang="pt-BR"><head><title>Dashboard Gerencial Unificado</title><link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet"></head>
<body style="background-color: #f0f2f5;">
<nav class="navbar navbar-dark bg-dark">
  <div class="container-fluid">
    <div class="d-flex gap-2">
        <a href="/dashboard" class="btn btn-outline-light btn-sm">⬅ Portal Principal</a>
        <a href="/demandas" class="btn btn-outline-info btn-sm">📂 Central de Demandas</a>
    </div>
    <span class="text-white fw-bold">📊 Dashboard Gerencial Unificado (PDFs + Avulsas + Credenciamento)</span>
    <div class="d-flex gap-2">
        <a href="/demandas/relatorio_pdf" class="btn btn-danger btn-sm fw-bold">📄 Baixar Relatório Geral do Cenário (PDF)</a>
        <a href="/credenciamento" class="btn btn-success btn-sm fw-bold">📋 Ir para Credenciamento - Fases</a>
    </div>
  </div>
</nav>
<div class="container mt-4">
    <!-- Resumo Geral Unificado (As 3 Frentes) -->
    <div class="row text-center mb-4 g-3">
        <div class="col-md-3">
            <div class="card shadow-sm border-0 bg-primary text-white p-3 h-100">
                <h6 class="mb-1 fw-bold">Total Geral Consolidado</h6>
                <h1 class="fw-bold mb-0">{{ d_total_geral }}</h1>
                <small class="text-light opacity-75">Soma das 3 Frentes Operacionais</small>
            </div>
        </div>
        <div class="col-md-3">
            <div class="card shadow-sm border-0 bg-dark text-white p-3 h-100">
                <h6 class="mb-1 fw-bold">1. Rotinas (PDFs)</h6>
                <h1 class="fw-bold mb-0">{{ t_rotina }}</h1>
                <small class="text-light opacity-75">Processos via Upload PDF</small>
            </div>
        </div>
        <div class="col-md-3">
            <div class="card shadow-sm border-0 bg-secondary text-white p-3 h-100">
                <h6 class="mb-1 fw-bold">2. Demandas Avulsas</h6>
                <h1 class="fw-bold mb-0">{{ t_avulsa }}</h1>
                <small class="text-light opacity-75">Registros manuais</small>
            </div>
        </div>
        <div class="col-md-3">
            <div class="card shadow-sm border-0 bg-success text-white p-3 h-100">
                <h6 class="mb-1 fw-bold">3. Credenciamento (Fase Atual)</h6>
                <h1 class="fw-bold mb-0">{{ t_cred_atual }}</h1>
                <small class="text-light opacity-75">Total Histórico: {{ t_cred_total }} linhas</small>
            </div>
        </div>
    </div>

    <!-- Contagem Detalhada de TODOS os Status (Rotinas + Avulsas) -->
    <h5 class="mb-3 text-secondary fw-bold">📊 Status: Rotinas (PDFs) & Demandas Avulsas</h5>
    <div class="row text-center g-3 mb-4">
        {% for item in status_cards %}
        <div class="col-md-3 col-sm-6">
            <div class="card shadow-sm border-0 h-100 p-3" style="border-left: 5px solid {{ item.cor }} !important;">
                <h6 class="text-secondary fw-bold mb-1">{{ item.nome }}</h6>
                <h2 class="fw-bold mb-1" style="color: {{ item.cor }};">{{ item.total }}</h2>
                <small class="text-muted">PDFs: <strong>{{ item.rotina }}</strong> | Avulsas: <strong>{{ item.avulsa }}</strong></small>
            </div>
        </div>
        {% endfor %}
    </div>

    <!-- Contagem Detalhada do Módulo Credenciamento - Fases -->
    <h5 class="mb-3 text-secondary fw-bold">📋 Status/Parecer: Credenciamento - Fases (Fase Atual = SIM)</h5>
    <div class="row text-center g-3 mb-4">
        <div class="col-md-3 col-6">
            <div class="card shadow-sm border-0 p-3 h-100" style="border-left: 5px solid #f0ad4e !important;">
                <h6 class="text-secondary fw-bold mb-1">EM ANÁLISE / PENDENTES (SLA Ativo)</h6>
                <h2 class="fw-bold mb-0 text-warning">{{ cred_parecer.get('EM ANÁLISE', 0) + cred_parecer.get('EM ANALISE', 0) + cred_parecer.get('PENDENTE', 0) + cred_parecer.get('AGUARDANDO 90 DIAS', 0) + cred_parecer.get('AGUARDANDO RETORNO', 0) }}</h2>
            </div>
        </div>
        <div class="col-md-3 col-6">
            <div class="card shadow-sm border-0 p-3 h-100" style="border-left: 5px solid #198754 !important;">
                <h6 class="text-secondary fw-bold mb-1">DEFERIDOS / FINALIZADOS (SLA Parado)</h6>
                <h2 class="fw-bold mb-0 text-success">{{ cred_parecer.get('DEFERIDO', 0) + cred_parecer.get('FINALIZADA', 0) + cred_parecer.get('CONCLUÍDO', 0) }}</h2>
            </div>
        </div>
        <div class="col-md-3 col-6">
            <div class="card shadow-sm border-0 p-3 h-100" style="border-left: 5px solid #fd7e14 !important;">
                <h6 class="text-secondary fw-bold mb-1">DEVOLVIDOS (SLA Parado)</h6>
                <h2 class="fw-bold mb-0" style="color: #fd7e14;">{{ cred_parecer.get('DEVOLVIDO', 0) }}</h2>
            </div>
        </div>
        <div class="col-md-3 col-6">
            <div class="card shadow-sm border-0 p-3 h-100" style="border-left: 5px solid #dc3545 !important;">
                <h6 class="text-secondary fw-bold mb-1">INDEFERIDOS / CANCEL. (SLA Parado)</h6>
                <h2 class="fw-bold mb-0 text-danger">{{ cred_parecer.get('INDEFERIDO', 0) + cred_parecer.get('CANCELADA', 0) }}</h2>
            </div>
        </div>
    </div>
    
    <!-- SLA UNIFICADO (PDFs + Avulsas + Credenciamento) -->
    <h5 class="mb-0 text-secondary fw-bold mt-4">⏱️ Monitoramento Unificado de SLA (Total em Aberto nas 3 Frentes: {{ total_ativas_unificado }})</h5>
    <small class="text-muted">Processos Finalizados, Concluídos, Deferidos, Indeferidos ou Devolvidos têm o SLA congelado automaticamente.</small>
    <div class="row text-center mt-2 mb-4 g-3">
        <div class="col-md-3">
            <div class="card shadow-sm border-0 p-3 h-100" style="background-color: #e2e3e5; color: #383d41;">
                <h6 class="mb-1 fw-bold">No Prazo (0 a 1 dia)</h6>
                <h2 class="fw-bold mb-1">{{ sla_em_dia + cred_sla.no_prazo }}</h2>
                <small>Demandas: {{ sla_em_dia }} | Credenc.: {{ cred_sla.no_prazo }}</small>
            </div>
        </div>
        <div class="col-md-3">
            <div class="card shadow-sm border-0 p-3 h-100" style="background-color: #fff3cd; color: #856404;">
                <h6 class="mb-1 fw-bold">Atenção (2 a 3 dias)</h6>
                <h2 class="fw-bold mb-1">{{ sla_atencao + cred_sla.atencao }}</h2>
                <small>Demandas: {{ sla_atencao }} | Credenc.: {{ cred_sla.atencao }}</small>
            </div>
        </div>
        <div class="col-md-3">
            <div class="card shadow-sm border-0 p-3 h-100" style="background-color: #ffe8cc; color: #d35400;">
                <h6 class="mb-1 fw-bold">Urgente (4 a 5 dias)</h6>
                <h2 class="fw-bold mb-1">{{ sla_urgente + cred_sla.urgente }}</h2>
                <small>Demandas: {{ sla_urgente }} | Credenc.: {{ cred_sla.urgente }}</small>
            </div>
        </div>
        <div class="col-md-3">
            <div class="card shadow-sm border-0 p-3 h-100" style="background-color: #f8d7da; color: #721c24;">
                <h6 class="mb-1 fw-bold">Crítico (6+ dias)</h6>
                <h2 class="fw-bold mb-1">{{ sla_critico + cred_sla.critico }}</h2>
                <small>Demandas: {{ sla_critico }} | Credenc.: {{ cred_sla.critico }}</small>
            </div>
        </div>
    </div>

    <div class="row">
        <div class="col-md-6 mb-4"><div class="card shadow-sm border-0 h-100"><div class="card-body text-center"><h5 class="fw-bold text-secondary mb-3">Produtividade por Colaborador (Rotinas)</h5><img src="data:image/png;base64,{{ chart_prod }}" class="img-fluid rounded"></div></div></div>
        <div class="col-md-6 mb-4"><div class="card shadow-sm border-0 h-100"><div class="card-body text-center"><h5 class="fw-bold text-secondary mb-3">Distribuição por Tipo de Demanda</h5><img src="data:image/png;base64,{{ chart_tipo }}" class="img-fluid rounded"></div></div></div>
    </div>
</div></body></html>"""

TELA_HUB_PENTEFINO = """<!DOCTYPE html><html lang="pt-BR"><head><title>Pente Fino RN 665</title><link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet"></head>
<body style="background-color: #f0f2f5;">
<nav class="navbar navbar-dark bg-dark"><div class="container-fluid"><a href="/dashboard" class="btn btn-outline-light btn-sm">⬅ Voltar ao Portal</a><span class="navbar-text text-white fw-bold">Pente Fino (RN 665)</span></div></nav>
<div class="container mt-5 text-center">
    <div class="row justify-content-center">
        <div class="col-md-5 mb-4">
            <div class="card shadow-sm border-0 h-100 p-4"><h1 class="mb-3 text-warning" style="font-size: 3rem;">🌍</h1><h4 class="fw-bold text-secondary">Cruzamento Regional em Massa</h4><p class="text-muted">Análise de cobertura completa comparando a Postal com a Operadora Intermediária.</p><a href="/pentefino/regional" class="btn btn-warning text-white w-100 py-3 fw-bold fs-6">Acessar</a></div>
        </div>
        <div class="col-md-5 mb-4">
            <div class="card shadow-sm border-0 h-100 p-4"><h1 class="mb-3 text-danger" style="font-size: 3rem;">🎯</h1><h4 class="fw-bold text-secondary">Busca Individual (Sniper)</h4><p class="text-muted">Procure substitutos internamente e avalie o Pool da Região.</p><a href="/pentefino/individual" class="btn btn-danger text-white w-100 py-3 fw-bold fs-6">Acessar</a></div>
        </div>
    </div>
</div></body></html>"""

TELA_PENTEFINO_REGIONAL = """<!DOCTYPE html><html lang="pt-BR"><head><title>Cruzamento Regional</title><link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet"></head>
<body style="background-color: #f0f2f5;">
<nav class="navbar navbar-dark bg-dark"><div class="container-fluid"><a href="/pentefino" class="btn btn-outline-light btn-sm">⬅ Voltar</a><span class="text-white fw-bold">Cruzamento Regional (Em Massa)</span></div></nav>
<div class="container mt-4">
    {% with messages = get_flashed_messages(with_categories=true) %}{% if messages %}{% for category, message in messages %}<div class="alert alert-{{ category }}">{{ message }}</div>{% endfor %}{% endif %}{% endwith %}
    <div class="card shadow-sm border-0">
        <div class="card-header bg-warning text-white fw-bold">🌍 Mapeamento e Upload das Bases</div>
        <div class="card-body">
            <form method="POST" action="/pentefino/regional" enctype="multipart/form-data">
                <div class="row mb-3 bg-light p-3 rounded">
                    <div class="col-md-6 border-end">
                        <label class="form-label fw-bold text-primary">1. Base Postal Saúde</label>
                        <input type="file" id="f_postal" name="f_postal" class="form-control mb-2" accept=".xlsx,.xls,.csv" required onchange="lerCabecalho(this, 'col_tipo_pos', 'col_esp_pos')">
                        <label class="form-label small">Coluna: Tipo Prestador</label>
                        <select id="col_tipo_pos" name="col_tipo_pos" class="form-select form-select-sm mb-2"><option value="">Aguardando arquivo...</option></select>
                        <label class="form-label small">Coluna: Especialidade</label>
                        <select id="col_esp_pos" name="col_esp_pos" class="form-select form-select-sm"><option value="">Aguardando arquivo...</option></select>
                    </div>
                    <div class="col-md-6">
                        <label class="form-label fw-bold text-success">2. Base Operadora Intermediária</label>
                        <input type="file" id="f_op" name="f_op" class="form-control mb-2" accept=".xlsx,.xls,.csv" required onchange="lerCabecalho(this, 'col_tipo_op', 'col_esp_op')">
                        <label class="form-label small">Coluna: Tipo Prestador</label>
                        <select id="col_tipo_op" name="col_tipo_op" class="form-select form-select-sm mb-2"><option value="">Aguardando arquivo...</option></select>
                        <label class="form-label small">Coluna: Especialidade</label>
                        <select id="col_esp_op" name="col_esp_op" class="form-select form-select-sm"><option value="">Aguardando arquivo...</option></select>
                    </div>
                </div>
                <div class="mb-4 text-center">
                    <label class="form-label fw-bold">3. Base Geográfica (IBGE)</label>
                    <input type="file" name="f_ibge" class="form-control w-50 mx-auto" accept=".xlsx,.xls,.csv" required>
                </div>
                
                <div id="progresso-container" class="mt-3 mb-3" style="display: none;">
                    <p id="progresso-texto" class="text-center fw-bold text-warning mb-2">A processar...</p>
                    <div class="progress" style="height: 20px;">
                        <div class="progress-bar progress-bar-striped progress-bar-animated bg-warning w-100"></div>
                    </div>
                </div>

                <button type="submit" class="btn btn-success w-100 py-3 fw-bold fs-5" onclick="mostrarProgresso('A processar cruzamento regional... Pode demorar até 2 minutos.')">Processar e Gerar Relatório Excel</button>
            </form>
        </div>
    </div>
</div>
<script>
function mostrarProgresso(texto) {
    document.getElementById('progresso-container').style.display = 'block';
    document.getElementById('progresso-texto').innerText = texto;
}

function lerCabecalho(input, selectId1, selectId2) {
    if (!input.files || input.files.length === 0) return;
    mostrarProgresso('A ler as colunas da planilha no servidor...');
    let s1 = document.getElementById(selectId1); let s2 = document.getElementById(selectId2);
    s1.innerHTML = '<option value="">A ler colunas no servidor...</option>'; s2.innerHTML = '<option value="">A ler colunas no servidor...</option>';
    let formData = new FormData(); formData.append("file", input.files[0]);
    fetch('/api/ler_cabecalhos', { method: 'POST', body: formData })
    .then(r => r.json())
    .then(data => {
        document.getElementById('progresso-container').style.display = 'none';
        if(data.erro) { s1.innerHTML = '<option value="">Erro ao ler</option>'; s2.innerHTML = '<option value="">Erro ao ler</option>'; return; }
        s1.innerHTML = ''; s2.innerHTML = '';
        data.colunas.forEach(c => { s1.options.add(new Option(c, c)); s2.options.add(new Option(c, c)); });
    }).catch(e => { 
        document.getElementById('progresso-container').style.display = 'none';
        s1.innerHTML = '<option value="">Falha na leitura</option>'; s2.innerHTML = '<option value="">Falha na leitura</option>'; 
    });
}
</script>
</body></html>"""

TELA_PENTEFINO_INDIVIDUAL = """<!DOCTYPE html><html lang="pt-BR"><head><title>Busca Sniper Individual</title><link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet"></head>
<body style="background-color: #f0f2f5;">
<nav class="navbar navbar-dark bg-dark"><div class="container-fluid"><a href="/pentefino" class="btn btn-outline-light btn-sm">⬅ Voltar</a><span class="text-white fw-bold">Busca Individual (Sniper)</span></div></nav>
<div class="container mt-4">
    {% with messages = get_flashed_messages(with_categories=true) %}{% if messages %}{% for category, message in messages %}<div class="alert alert-{{ category }}">{{ message }}</div>{% endfor %}{% endif %}{% endwith %}
    
    <div class="card shadow-sm border-0">
        <div class="card-header bg-danger text-white fw-bold">🎯 Parâmetros de Substituição</div>
        <div class="card-body">
            <form method="POST" action="/pentefino/individual" enctype="multipart/form-data">
                <div class="row mb-3 bg-light p-3 rounded">
                    <div class="col-md-6 border-end">
                        <label class="form-label fw-bold text-primary">1. Base Postal Saúde (Upload p/ liberar as buscas)</label>
                        <input type="file" name="f_postal" class="form-control mb-2" accept=".xlsx,.xls,.csv" required onchange="carregarInteligencia(this)">
                    </div>
                    <div class="col-md-6">
                        <label class="form-label fw-bold text-success">2. Base Geográfica (IBGE)</label>
                        <input type="file" name="f_ibge" class="form-control mb-2" accept=".xlsx,.xls,.csv" required>
                    </div>
                </div>
                <hr>
                <div class="row mt-4">
                    <div class="col-md-6 mb-3">
                        <label class="form-label fw-bold text-danger">Pesquisar Alvo (Nome OU CNPJ)</label>
                        <input type="text" name="alvo" list="lista_alvos" class="form-control border-danger" placeholder="Carregue a planilha acima para liberar a pesquisa..." required autocomplete="off" disabled>
                        <datalist id="lista_alvos"></datalist>
                    </div>
                    <div class="col-md-6 mb-3">
                        <label class="form-label fw-bold">Filtro de Tipo Prestador</label>
                        <select name="filtro_tipo" id="filtro_tipo" class="form-select">
                            <option value="">Todos (Sem filtro)</option>
                        </select>
                    </div>
                </div>
                
                <div id="progresso-container" class="mt-3 mb-3" style="display: none;">
                    <p id="progresso-texto" class="text-center fw-bold text-danger mb-2">A processar...</p>
                    <div class="progress" style="height: 20px;">
                        <div class="progress-bar progress-bar-striped progress-bar-animated bg-danger w-100"></div>
                    </div>
                </div>

                <button id="btn-submit" type="submit" class="btn btn-danger w-100 py-3 fw-bold fs-5 mt-3" onclick="mostrarProgresso('A gerar Relatório Sniper... Aguarde alguns segundos.')">🔍 Procurar Substitutos e Gerar Análise</button>
            </form>
        </div>
    </div>
</div>
<script>
function mostrarProgresso(texto) {
    document.getElementById('progresso-container').style.display = 'block';
    document.getElementById('progresso-texto').innerText = texto;
}

function carregarInteligencia(input) {
    if (!input.files || input.files.length === 0) return;
    mostrarProgresso('A extrair inteligência da planilha (CNPJs e Tipos)...');
    
    let inputAlvo = document.querySelector('input[name="alvo"]');
    let selectTipo = document.getElementById('filtro_tipo');
    let datalist = document.getElementById('lista_alvos');
    
    inputAlvo.disabled = true;
    inputAlvo.placeholder = "A analisar a planilha no servidor... Aguarde...";
    selectTipo.innerHTML = '<option value="">A carregar tipos...</option>';
    
    let formData = new FormData(); formData.append("f_postal", input.files[0]);
    
    fetch('/api/carregar_alvos', { method: 'POST', body: formData })
    .then(r => r.json())
    .then(data => {
        document.getElementById('progresso-container').style.display = 'none';
        if(data.erro) { 
            inputAlvo.placeholder = "Erro: " + data.erro; 
            selectTipo.innerHTML = '<option value="">Todos (Sem filtro)</option>'; 
            return; 
        }
        
        datalist.innerHTML = '';
        data.alvos.forEach(a => { let opt = document.createElement('option'); opt.value = a; datalist.appendChild(opt); });
        
        selectTipo.innerHTML = '<option value="">Todos (Sem filtro)</option><option value="MESMO TIPO DO ALVO">Mesmo Tipo do Alvo</option>';
        data.tipos.forEach(t => { let opt = document.createElement('option'); opt.value = t; opt.text = t; selectTipo.appendChild(opt); });
        
        inputAlvo.placeholder = "Lista pronta! Digite Nome ou CNPJ...";
        inputAlvo.disabled = false;
    }).catch(e => { 
        document.getElementById('progresso-container').style.display = 'none';
        inputAlvo.placeholder = "Falha ao processar. Digite manualmente."; 
        selectTipo.innerHTML = '<option value="">Todos (Sem filtro)</option>'; 
        inputAlvo.disabled = false; 
    });
}
</script>
</body></html>"""

TELA_PENTEFINO_INDIVIDUAL_RESULTADO = """<!DOCTYPE html><html lang="pt-BR"><head><title>Resultados Pente Fino</title><link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet"></head>
<body style="background-color: #f0f2f5;">
<nav class="navbar navbar-dark bg-dark"><div class="container-fluid"><a href="/pentefino/individual" class="btn btn-outline-light btn-sm">⬅ Voltar</a><span class="text-white fw-bold">Resultados Sniper</span></div></nav>
<div class="container mt-4">
    <div class="alert alert-success d-flex justify-content-between align-items-center shadow-sm">
        <div>
            <h4 class="fw-bold mb-1">✅ Análise Concluída!</h4>
            <p class="mb-0 text-dark">Alvo: <strong>{{ alvo_nome }}</strong> ({{ alvo_cnpj }})</p>
        </div>
        <a href="data:application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;base64,{{ excel_b64 }}" download="Relatorio_Substituicao_{{ alvo_cnpj }}.xlsx" class="btn btn-success btn-lg fw-bold shadow-sm">⬇ Baixar Planilha Completa (Excel)</a>
    </div>
    
    <!-- Linha POOL DO MUNICÍPIO -->
    <div class="alert alert-info shadow-sm border-info mt-3">
        <h5 class="fw-bold text-dark mb-1">📍 Avaliação do Pool do Município ({{ mun_alvo }}/{{ uf_alvo }})</h5>
        <p class="mb-0 text-dark"><strong>Status:</strong> {{ pool_mun_status }}</p>
        {% if pool_mun_faltantes and pool_mun_faltantes != 'A soma de todos os prestadores do município cobre as especialidades.' %}
        <hr><p class="mb-0 text-danger fw-bold small" style="white-space: pre-wrap;">{{ pool_mun_faltantes }}</p>
        {% endif %}
    </div>

    <!-- Linha POOL DA REGIÃO -->
    <div class="alert alert-warning shadow-sm border-warning mt-3">
        <h5 class="fw-bold text-dark mb-1">⭐ Avaliação do Pool da Região de Saúde ({{ regiao_alvo }})</h5>
        <p class="mb-0 text-dark"><strong>Status:</strong> {{ pool_reg_status }}</p>
        {% if pool_reg_faltantes and pool_reg_faltantes != 'A soma de todos os prestadores da região cobre as especialidades.' %}
        <hr><p class="mb-0 text-danger fw-bold small" style="white-space: pre-wrap;">{{ pool_reg_faltantes }}</p>
        {% endif %}
    </div>

    <h4 class="fw-bold text-secondary mt-5 mb-3 border-bottom pb-2">1. Análise Municipal (1ª Etapa)</h4>
    
    <!-- Tabela 1: 100% Município -->
    <div class="card shadow-sm border-0 mb-4">
        <div class="card-header bg-primary text-white fw-bold fs-5">🏆 100% de Cobertura no MUNICÍPIO</div>
        <div class="card-body p-0">
            <div class="table-responsive">
                <table class="table table-hover align-middle mb-0 text-center">
                    <thead class="table-light"><tr><th>Proximidade</th><th>CNPJ</th><th class="text-start">Prestador</th><th>Município/UF</th><th>Tipo</th></tr></thead>
                    <tbody>
                        {% for r in res_100_mun %}
                        <tr>
                            <td><span class="badge bg-info text-dark">{{ r[2] }}</span></td>
                            <td>{{ r[3] }}</td><td class="fw-bold text-start">{{ r[4] }}</td><td>{{ r[6] }}</td><td>{{ r[5] }}</td>
                        </tr>
                        {% else %}
                        <tr><td colspan="5" class="text-center text-muted py-4">Nenhum prestador 100% equivalente no município.</td></tr>
                        {% endfor %}
                    </tbody>
                </table>
            </div>
        </div>
    </div>
    
    <!-- Tabela 2: Vulneráveis Município -->
    <div class="card shadow-sm border-0 mb-5">
        <div class="card-header bg-danger text-white fw-bold fs-5">⚠️ Cobertura Parcial no MUNICÍPIO (Vulneráveis)</div>
        <div class="card-body p-0">
            <div class="table-responsive">
                <table class="table table-hover align-middle mb-0 text-center" style="font-size: 0.9em;">
                    <thead class="table-light"><tr><th>Proximidade</th><th>CNPJ</th><th class="text-start">Prestador</th><th>Município/UF</th><th class="text-start">Especialidades Faltantes</th></tr></thead>
                    <tbody>
                        {% for r in res_vul_mun %}
                        <tr>
                            <td><span class="badge bg-secondary">{{ r[2] }}</span></td>
                            <td>{{ r[3] }}</td><td class="fw-bold text-start">{{ r[4] }}</td><td>{{ r[6] }}</td>
                            <td class="text-start text-danger small" style="white-space: pre-wrap;">{{ r[8] }}</td>
                        </tr>
                        {% else %}
                        <tr><td colspan="5" class="text-center text-muted py-4">Nenhum prestador vulnerável no município.</td></tr>
                        {% endfor %}
                    </tbody>
                </table>
            </div>
        </div>
    </div>

    <h4 class="fw-bold text-secondary mt-5 mb-3 border-bottom pb-2">2. Análise Regional (2ª Etapa - Demais Municípios da Região)</h4>

    <!-- Tabela 3: 100% Região -->
    <div class="card shadow-sm border-0 mb-4">
        <div class="card-header text-white fw-bold fs-5" style="background-color: #0dcaf0;">🌟 100% de Cobertura na REGIÃO DE SAÚDE</div>
        <div class="card-body p-0">
            <div class="table-responsive">
                <table class="table table-hover align-middle mb-0 text-center">
                    <thead class="table-light"><tr><th>Proximidade</th><th>CNPJ</th><th class="text-start">Prestador</th><th>Município/UF</th><th>Tipo</th></tr></thead>
                    <tbody>
                        {% for r in res_100_reg %}
                        <tr>
                            <td><span class="badge bg-info text-dark">{{ r[2] }}</span></td>
                            <td>{{ r[3] }}</td><td class="fw-bold text-start">{{ r[4] }}</td><td>{{ r[6] }}</td><td>{{ r[5] }}</td>
                        </tr>
                        {% else %}
                        <tr><td colspan="5" class="text-center text-muted py-4">Nenhum prestador 100% equivalente na região de saúde.</td></tr>
                        {% endfor %}
                    </tbody>
                </table>
            </div>
        </div>
    </div>
    
    <!-- Tabela 4: Vulneráveis Região -->
    <div class="card shadow-sm border-0 mb-5">
        <div class="card-header text-white fw-bold fs-5" style="background-color: #fd7e14;">⚡ Cobertura Parcial na REGIÃO DE SAÚDE (Vulneráveis)</div>
        <div class="card-body p-0">
            <div class="table-responsive">
                <table class="table table-hover align-middle mb-0 text-center" style="font-size: 0.9em;">
                    <thead class="table-light"><tr><th>Proximidade</th><th>CNPJ</th><th class="text-start">Prestador</th><th>Município/UF</th><th class="text-start">Especialidades Faltantes</th></tr></thead>
                    <tbody>
                        {% for r in res_vul_reg %}
                        <tr>
                            <td><span class="badge bg-secondary">{{ r[2] }}</span></td>
                            <td>{{ r[3] }}</td><td class="fw-bold text-start">{{ r[4] }}</td><td>{{ r[6] }}</td>
                            <td class="text-start text-danger small" style="white-space: pre-wrap;">{{ r[8] }}</td>
                        </tr>
                        {% else %}
                        <tr><td colspan="5" class="text-center text-muted py-4">Nenhum prestador regional vulnerável encontrado.</td></tr>
                        {% endfor %}
                    </tbody>
                </table>
            </div>
        </div>
    </div>
    
    <div class="text-center mt-4 mb-5"><a href="/pentefino/individual" class="btn btn-secondary px-4 fw-bold">Fazer Nova Busca</a></div>
</div></body></html>"""


# ==========================================
# FUNÇÕES DE APOIO PENTE FINO E ARQUIVOS
# ==========================================
def normalizar_texto(texto):
    if pd.isna(texto) or str(texto).lower() in ['nan', 'none', '']: return ""
    return ''.join(c for c in unicodedata.normalize('NFD', str(texto)) if unicodedata.category(c) != 'Mn').upper().strip()

def cacador_de_colunas(df, palavras_chave):
    for k in palavras_chave:
        for col in df.columns:
            if k == str(col).strip().upper(): return col
    for k in palavras_chave:
        for col in df.columns:
            if k in str(col).strip().upper(): return col
    return None

def limpar_ibge(val):
    if pd.isna(val): return ""
    v = ''.join(filter(str.isdigit, str(val).split('.')[0].strip()))
    return v[:6] if len(v) >= 6 else v

def extrair_especialidades(series):
    specs = set()
    for val in series.dropna():
        for s in str(val).replace(';', ',').split(','):
            v = normalizar_texto(s)
            if v:
                if v in IGNORADAS: continue
                specs.add(DE_PARA.get(v, v))
    return specs

def ler_arquivo(file_obj, nrows=None):
    filename = file_obj.filename.lower()
    if filename.endswith('.csv') or filename.endswith('.txt'):
        raw_bytes = file_obj.read()
        file_obj.seek(0)
        
        texto_decodificado = None
        for enc in ['utf-8-sig', 'utf-8', 'cp1250', 'iso-8859-2', 'cp1252', 'latin1']:
            try:
                t = raw_bytes.decode(enc)
                texto_decodificado = t
                break
            except Exception:
                continue
        if texto_decodificado is None:
            texto_decodificado = raw_bytes.decode('latin1', errors='replace')
            
        linhas = [l for l in texto_decodificado.splitlines() if l.strip()]
        if not linhas:
            return pd.DataFrame()
            
        primeira_linha = linhas[0]
        
        m_sep = re.search(r'FILIALPADRAO(.+?)CNPJCPF', primeira_linha, re.IGNORECASE)
        if m_sep:
            sep = m_sep.group(1)
        elif 'Ź' in primeira_linha:
            sep = 'Ź'
        elif '\x8f' in primeira_linha:
            sep = '\x8f'
        elif ';' in primeira_linha:
            sep = ';'
        elif '|' in primeira_linha:
            sep = '|'
        elif '\t' in primeira_linha:
            sep = '\t'
        else:
            sep = ','

        if m_sep or sep in ['Ź', '\x8f'] or filename.endswith('.txt'):
            cabecalhos = [c.strip() for c in primeira_linha.split(sep)]
            while cabecalhos and cabecalhos[-1] == '':
                cabecalhos.pop()
            num_cols = len(cabecalhos)
            dados = []
            limite = (nrows + 1) if nrows is not None else len(linhas)
            for l in linhas[1:limite]:
                partes = l.split(sep)
                if len(partes) >= num_cols:
                    partes = partes[:num_cols]
                else:
                    partes.extend([''] * (num_cols - len(partes)))
                dados.append(partes)
            return pd.DataFrame(dados, columns=cabecalhos, dtype=str)

        import csv
        from io import StringIO
        try:
            return pd.read_csv(StringIO(texto_decodificado), sep=sep, dtype=str, nrows=nrows, engine='python', on_bad_lines='skip')
        except Exception:
            return pd.read_csv(StringIO(texto_decodificado), sep=sep, dtype=str, nrows=nrows, engine='python', quoting=csv.QUOTE_NONE, on_bad_lines='skip')
    else:
        return pd.read_excel(file_obj, dtype=str, nrows=nrows)


# ==========================================
# ROTAS DO SERVIDOR WEB E APIs INVISÍVEIS
# ==========================================

@app.route('/api/ler_cabecalhos', methods=['POST'])
def api_ler_cabecalhos():
    if 'user' not in session: return jsonify({"erro": "Não autorizado"}), 401
    file = request.files.get('file')
    if not file: return jsonify({"erro": "Arquivo não enviado"}), 400
    try:
        df = ler_arquivo(file, nrows=0)
        return jsonify({"colunas": [str(c).strip().upper() for c in df.columns]})
    except Exception as e: return jsonify({"erro": str(e)}), 500

@app.route('/api/carregar_alvos', methods=['POST'])
def api_carregar_alvos():
    if 'user' not in session: return jsonify({"erro": "Não autorizado"}), 401
    f_postal = request.files.get('f_postal')
    if not f_postal: return jsonify({"erro": "Arquivo não enviado"}), 400
    try:
        df = ler_arquivo(f_postal)
        df.columns = [str(c).strip().upper() for c in df.columns]
        
        col_cnpj = cacador_de_colunas(df, ['CNPJ', 'CPFCNPJ'])
        col_nome = cacador_de_colunas(df, ['NOME', 'RAZAO', 'PRESTADOR', 'FANTASIA'])
        col_tipo = cacador_de_colunas(df, ['TIPOPRESTADOR', 'TIPO_PRESTADOR', 'TIPO PRESTADOR', 'TIPO', 'CATEGORIA', 'TIPO DE PRESTADOR'])
        
        if not col_cnpj or not col_nome: return jsonify({"erro": "Colunas de Nome ou CNPJ não encontradas"})
        
        df[col_cnpj] = df[col_cnpj].astype(str).str.replace(r'\D', '', regex=True)
        df_clean = df.dropna(subset=[col_cnpj, col_nome]).drop_duplicates(subset=[col_cnpj])
        
        alvos = []
        tipos_set = set()
        for _, row in df_clean.iterrows():
            cnpj = str(row[col_cnpj]).strip(); nome = str(row[col_nome]).strip()
            if cnpj and cnpj != 'NAN' and nome: alvos.append(f"{cnpj} - {nome}")
            if col_tipo:
                tipo = str(row[col_tipo]).strip().upper()
                if tipo and tipo != 'NAN' and tipo != 'NONE' and tipo != '': tipos_set.add(tipo)
                
        return jsonify({"alvos": alvos, "tipos": sorted(list(tipos_set))})
    except Exception as e: return jsonify({"erro": str(e)}), 500

@app.route('/', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        user = request.form.get('username')
        password = request.form.get('password')
        conn = conectar_db()
        if not conn:
            flash("Erro de ligação com a base de dados.", "danger")
            return render_template_string(TELA_LOGIN)
            
        c = conn.cursor(); ph = get_ph(conn)
        c.execute(f"SELECT id, role FROM users WHERE username={ph} AND password={ph}", (user, password))
        res = c.fetchone()
        conn.close()
        
        if res:
            session['user'] = user; session['role'] = res[1]
            return redirect(url_for('dashboard'))
        else: flash("Usuário ou senha incorretos!", "danger")
    return render_template_string(TELA_LOGIN)

@app.route('/dashboard')
def dashboard():
    if 'user' not in session: return redirect(url_for('login'))
    return render_template_string(TELA_DASHBOARD, user=session['user'], role=session['role'])

# -------- ROTAS DO MÓDULO CREDENCIAMENTO - FASES --------
@app.route('/credenciamento')
def credenciamento_fases():
    if 'user' not in session: return redirect(url_for('login'))
    inicializar_tabela_credenciamento()
    conn = conectar_db(); c = conn.cursor(); ph = get_ph(conn)

    f_ultima = request.args.get('ultima_fase', 'SIM')
    f_sla = request.args.get('sla_filtro', 'Todos')
    f_uf = request.args.get('uf', 'Todos')
    f_fase = request.args.get('fase', 'Todos')
    f_parecer = request.args.get('parecer', 'Todos')
    f_origem = request.args.get('origem', 'Todos')
    f_resp = request.args.get('resp', 'Todos')
    f_busca = request.args.get('busca', '').strip()

    q = "SELECT * FROM credenciamento_fases WHERE 1=1"
    p = []
    if f_ultima != 'Todos':
        q += f" AND ultima_fase={ph}"; p.append(f_ultima)
    if f_uf != 'Todos':
        q += f" AND filial={ph}"; p.append(f_uf)
    if f_fase != 'Todos':
        q += f" AND fase={ph}"; p.append(f_fase)
    if f_parecer != 'Todos':
        q += f" AND parecer={ph}"; p.append(f_parecer)
    if f_origem != 'Todos':
        q += f" AND origem_parecer={ph}"; p.append(f_origem)
    if f_resp != 'Todos':
        q += f" AND usuario_resp={ph}"; p.append(f_resp)
    if f_busca:
        termo = f"%{f_busca.upper()}%"
        q += f" AND (UPPER(cnpj_cpf) LIKE {ph} OR UPPER(nome) LIKE {ph} OR UPPER(fantasia) LIKE {ph} OR UPPER(processo) LIKE {ph})"
        p.extend([termo, termo, termo, termo])

    q += " ORDER BY id ASC"
    c.execute(q, tuple(p))
    rows_brutas = c.fetchall()

    registros = []
    sla_resumo = {"ativos": 0, "no_prazo": 0, "atencao": 0, "urgente": 0, "critico": 0, "alerta_90d": 0}
    kpi_parecer = {}
    kpi_fases_dict = {}
    kpi_ufs_dict = {}
    proc_set = set()

    for r in rows_brutas:
        info_sla = calcular_sla_credenciamento(r)
        if f_sla != 'Todos':
            if f_sla == '90d' and not info_sla['estourou_90d']:
                continue
            elif f_sla != '90d' and info_sla['faixa_sla'] != f_sla:
                continue

        registros.append({"row": r, "sla": info_sla})

        if info_sla["ativo"]:
            sla_resumo["ativos"] += 1
            if info_sla["faixa_sla"] == "Crítico": sla_resumo["critico"] += 1
            elif info_sla["faixa_sla"] == "Urgente": sla_resumo["urgente"] += 1
            elif info_sla["faixa_sla"] == "Atenção": sla_resumo["atencao"] += 1
            else: sla_resumo["no_prazo"] += 1
            if info_sla["estourou_90d"]: sla_resumo["alerta_90d"] += 1

        par = str(r[20] or 'SEM PARECER').strip().upper()
        kpi_parecer[par] = kpi_parecer.get(par, 0) + 1
        fa = r[12] if r[12] else 'SEM FASE'
        kpi_fases_dict[fa] = kpi_fases_dict.get(fa, 0) + 1
        uf_val = r[1] if r[1] else 'S/UF'
        kpi_ufs_dict[uf_val] = kpi_ufs_dict.get(uf_val, 0) + 1
        if r[8]: proc_set.add(r[8])

    kpi_fases = sorted(kpi_fases_dict.items(), key=lambda x: x[0])

    chart_fases = ""
    if kpi_fases:
        fig1, ax1 = plt.subplots(figsize=(6.5, 3.8))
        nomes_f = [k[0][:28] for k in kpi_fases]
        vals_f = [k[1] for k in kpi_fases]
        ax1.barh(nomes_f, vals_f, color='#004b87')
        ax1.invert_yaxis()
        ax1.tick_params(axis='y', labelsize=8)
        buf1 = BytesIO(); fig1.tight_layout(); fig1.savefig(buf1, format="png"); buf1.seek(0)
        chart_fases = base64.b64encode(buf1.read()).decode('utf-8'); plt.close(fig1)

    chart_ufs = ""
    if kpi_ufs_dict:
        ufs_ord = sorted(kpi_ufs_dict.items(), key=lambda x: x[1], reverse=True)
        fig2, ax2 = plt.subplots(figsize=(6.5, 3.8))
        ax2.bar([u[0] for u in ufs_ord], [u[1] for u in ufs_ord], color='#198754')
        buf2 = BytesIO(); fig2.tight_layout(); fig2.savefig(buf2, format="png"); buf2.seek(0)
        chart_ufs = base64.b64encode(buf2.read()).decode('utf-8'); plt.close(fig2)

    c.execute("SELECT DISTINCT filial FROM credenciamento_fases WHERE filial IS NOT NULL AND filial != '' ORDER BY filial")
    lista_ufs = [r[0] for r in c.fetchall()]
    c.execute("SELECT DISTINCT fase FROM credenciamento_fases WHERE fase IS NOT NULL AND fase != '' ORDER BY fase")
    fases_banco = [r[0] for r in c.fetchall()]
    lista_fases = sorted(list(set(FASES_PADRAO_CRED + fases_banco)))

    c.execute("SELECT DISTINCT parecer FROM credenciamento_fases WHERE parecer IS NOT NULL AND parecer != '' ORDER BY parecer")
    pareceres_banco = [r[0] for r in c.fetchall()]
    lista_pareceres = list(dict.fromkeys(STATUS_CRED + pareceres_banco))

    c.execute("SELECT DISTINCT origem_parecer FROM credenciamento_fases WHERE origem_parecer IS NOT NULL AND origem_parecer != '' ORDER BY origem_parecer")
    lista_origens = [r[0] for r in c.fetchall()]

    # Puxa EXCLUSIVAMENTE os usuários cadastrados na tabela users (Painel Admin)
    c.execute("SELECT username FROM users ORDER BY id ASC")
    logins_sistema = [r[0] for r in c.fetchall() if r[0]]
    usuarios_login = ['Nenhum'] + logins_sistema
    lista_resps = ['Nenhum'] + logins_sistema

    c.execute("SELECT MAX(data_atualizacao) FROM credenciamento_fases")
    row_att = c.fetchone()
    ultima_att = row_att[0] if row_att and row_att[0] else None

    conn.close()
    return render_template_string(
        TELA_CREDENCIAMENTO_FASES,
        registros=registros,
        sla_resumo=sla_resumo,
        kpi_parecer=kpi_parecer,
        kpi_fases=kpi_fases,
        chart_fases=chart_fases,
        chart_ufs=chart_ufs,
        total_proc_unicos=len(proc_set),
        lista_ufs=lista_ufs,
        lista_fases=lista_fases,
        opcoes_fases_modal=lista_fases,
        status_cred_modal=lista_pareceres,
        usuarios_login=usuarios_login,
        lista_pareceres=lista_pareceres,
        lista_origens=lista_origens,
        lista_resps=lista_resps,
        f_ultima=f_ultima, f_sla=f_sla, f_uf=f_uf, f_fase=f_fase, f_parecer=f_parecer, f_origem=f_origem, f_resp=f_resp, f_busca=f_busca,
        ultima_att=ultima_att
    )

@app.route('/credenciamento/editar', methods=['POST'])
def editar_credenciamento():
    if 'user' not in session: return redirect(url_for('login'))
    c_id = request.form.get('id')
    nova_fase = request.form.get('fase', '').strip()
    novo_status = request.form.get('status', '').strip()
    novo_resp = request.form.get('responsavel', '').strip()
    nova_origem = request.form.get('origem_parecer', '').strip()
    nova_obs = request.form.get('observacao', '').strip()

    conn = conectar_db(); c = conn.cursor(); ph = get_ph(conn)
    try:
        c.execute(f"SELECT fase, data_inicial, data_final FROM credenciamento_fases WHERE id={ph}", (c_id,))
        atual = c.fetchone()
        hoje_str = datetime.now().strftime("%d/%m/%Y")

        data_ini = atual[1] if atual else hoje_str
        data_fim = atual[2] if atual else ""

        if atual and str(atual[0] or '').strip() != nova_fase:
            data_ini = hoje_str

        if novo_status.upper() in STATUS_ENCERRADOS_DEMANDAS:
            if not data_fim:
                data_fim = hoje_str
        else:
            data_fim = ""

        c.execute(f"""
            UPDATE credenciamento_fases
            SET fase={ph}, parecer={ph}, usuario_resp={ph}, origem_parecer={ph}, observacao={ph},
                data_inicial={ph}, data_final={ph}
            WHERE id={ph}
        """, (nova_fase, novo_status, novo_resp, nova_origem, nova_obs, data_ini, data_fim, c_id))
        conn.commit()
        flash("Processo de credenciamento atualizado com sucesso!", "success")
    except Exception as e:
        conn.rollback()
        flash(f"Erro ao atualizar processo de credenciamento: {e}", "danger")
    finally:
        conn.close()
    return redirect(url_for('credenciamento_fases'))

@app.route('/credenciamento/acao', methods=['POST'])
def acao_credenciamento():
    if 'user' not in session: return redirect(url_for('login'))
    c_id = request.form.get('id')
    acao = request.form.get('acao')
    conn = conectar_db(); c = conn.cursor(); ph = get_ph(conn)
    try:
        if acao == 'assumir':
            c.execute(f"UPDATE credenciamento_fases SET usuario_resp={ph}, parecer='Em Análise', data_final='' WHERE id={ph}", (session['user'], c_id))
            flash(f"Processo assumido por {session['user']}!", "info")
        conn.commit()
    except Exception as e:
        conn.rollback()
        flash(f"Erro ao executar ação: {e}", "danger")
    finally:
        conn.close()
    return redirect(url_for('credenciamento_fases'))

@app.route('/credenciamento/upload', methods=['POST'])
def upload_credenciamento():
    if 'user' not in session: return redirect(url_for('login'))
    file = request.files.get('arquivo_cred')
    if not file or file.filename == '':
        flash("Selecione um arquivo válido.", "warning")
        return redirect(url_for('credenciamento_fases'))

    try:
        inicializar_tabela_credenciamento()
        df = ler_arquivo(file)
        df.columns = [corrigir_encoding_postal(c).strip().upper() for c in df.columns]
        df = df.loc[:, ~df.columns.str.contains('^UNNAMED')]

        def get_val(row, col_names):
            for cn in col_names:
                if cn in row and pd.notna(row[cn]):
                    return corrigir_encoding_postal(row[cn])
            return ""

        dt_att = datetime.now().strftime("%d/%m/%Y %H:%M")
        linhas_para_inserir = []
        for _, row in df.iterrows():
            filial = get_val(row, ['FILIALPADRAO', 'FILIAL', 'UF'])
            cnpj_cpf = get_val(row, ['CNPJCPF', 'CNPJ', 'CPF'])
            if not filial and not cnpj_cpf:
                continue
            ultima_f = get_val(row, ['ULTIMAFASE', 'ULTIMA_FASE']).upper()
            if ultima_f in ['NAO', 'NĂO', 'NÃO', 'N']:
                ultima_f = 'NÃO'
            elif ultima_f in ['SIM', 'S']:
                ultima_f = 'SIM'

            linhas_para_inserir.append((
                filial,
                cnpj_cpf,
                get_val(row, ['NOME', 'RAZAO SOCIAL']),
                get_val(row, ['FANTASIA', 'NOME FANTASIA']),
                get_val(row, ['TIPOPRESTADOR', 'TIPO PRESTADOR']),
                get_val(row, ['CATEGORIA']),
                get_val(row, ['TIPOPROCESSO', 'TIPO PROCESSO']),
                get_val(row, ['PROCESSO', 'NUM PROCESSO']),
                get_val(row, ['DATADATAABERTURAPROCESSO', 'DATAABERTURAPROCESSO', 'DATA ABERTURA']),
                get_val(row, ['DATAFINALPROCESSO', 'DATA FINAL PROCESSO']),
                get_val(row, ['USUARIORESPONSAVEL', 'RESPONSAVEL']),
                get_val(row, ['FASE']),
                ultima_f,
                get_val(row, ['DATAINICIAL', 'DATA INICIAL']),
                get_val(row, ['DATAFINAL', 'DATA FINAL']),
                get_val(row, ['K_PROXIMAFASE', 'PROXIMAFASE']),
                get_val(row, ['K_DIASPERCORRIDOS', 'DIASPERCORRIDOS']),
                get_val(row, ['K_DIASPERCORRIDOSPF', 'DIASPERCORRIDOSPF']),
                get_val(row, ['JUSTIFICATIVAINDEFERIMENTO']),
                get_val(row, ['PARECER']).upper(),
                get_val(row, ['OBSERVACAO']),
                get_val(row, ['ORIGEMPARECER']).upper(),
                get_val(row, ['DATAAPROVACAO']),
                get_val(row, ['DATAREJEICAO']),
                get_val(row, ['USUARIOAPRJ']),
                get_val(row, ['K_JUSTIFICATIVACORGEN']),
                get_val(row, ['USUARIONOTIFICACAO']),
                get_val(row, ['GRUPO_PRESTADORES']),
                dt_att
            ))

        conn = conectar_db(); c = conn.cursor(); ph = get_ph(conn)
        c.execute("DELETE FROM credenciamento_fases")
        sincronizar_sequencias()
        placeholders = ", ".join([ph] * 29)
        q_ins = f"""
            INSERT INTO credenciamento_fases (
                filial, cnpj_cpf, nome, fantasia, tipo_prestador, categoria, tipo_processo,
                processo, data_abertura, data_final_proc, usuario_resp, fase, ultima_fase,
                data_inicial, data_final, proxima_fase, dias_percorridos, dias_percorridos_pf,
                justificativa_indef, parecer, observacao, origem_parecer, data_aprovacao,
                data_rejeicao, usuario_aprj, justificativa_corgen, usuario_notificacao,
                grupo_prestadores, data_atualizacao
            ) VALUES ({placeholders})
        """
        c.executemany(q_ins, linhas_para_inserir)
        conn.commit(); conn.close()
        flash(f"Relatório atualizado com sucesso! {len(linhas_para_inserir)} registros de fases importados.", "success")
    except Exception as e:
        flash(f"Erro ao processar relatório de Credenciamento: {e}", "danger")

    return redirect(url_for('credenciamento_fases'))

@app.route('/api/credenciamento/processo/<num_proc>')
def api_historico_credenciamento(num_proc):
    if 'user' not in session: return jsonify({"erro": "Não autorizado"}), 401
    conn = conectar_db(); c = conn.cursor(); ph = get_ph(conn)
    c.execute(f"SELECT * FROM credenciamento_fases WHERE processo={ph} ORDER BY id ASC", (num_proc,))
    rows = c.fetchall()
    conn.close()
    etapas = []
    for r in rows:
        etapas.append({
            "id": r[0], "filial": r[1], "cnpj_cpf": r[2], "nome": r[3], "fantasia": r[4],
            "tipo_prestador": r[5], "categoria": r[6], "tipo_processo": r[7], "processo": r[8],
            "data_abertura": r[9], "data_final_proc": r[10], "usuario_resp": r[11], "fase": r[12],
            "ultima_fase": r[13], "data_inicial": r[14], "data_final": r[15], "proxima_fase": r[16],
            "dias_percorridos": r[17], "dias_percorridos_pf": r[18], "justificativa_indef": r[19],
            "parecer": r[20], "observacao": r[21], "origem_parecer": r[22], "data_aprovacao": r[23],
            "data_rejeicao": r[24], "usuario_aprj": r[25], "justificativa_corgen": r[26],
            "usuario_notificacao": r[27], "grupo_prestadores": r[28]
        })
    return jsonify({"etapas": etapas})

@app.route('/credenciamento/exportar')
def exportar_credenciamento():
    if 'user' not in session: return redirect(url_for('login'))
    conn = conectar_db(); c = conn.cursor(); ph = get_ph(conn)
    f_ultima = request.args.get('ultima_fase', 'SIM')
    f_sla = request.args.get('sla_filtro', 'Todos')
    f_uf = request.args.get('uf', 'Todos')
    f_fase = request.args.get('fase', 'Todos')
    f_parecer = request.args.get('parecer', 'Todos')
    f_origem = request.args.get('origem', 'Todos')
    f_resp = request.args.get('resp', 'Todos')
    f_busca = request.args.get('busca', '').strip()

    q = "SELECT * FROM credenciamento_fases WHERE 1=1"
    p = []
    if f_ultima != 'Todos': q += f" AND ultima_fase={ph}"; p.append(f_ultima)
    if f_uf != 'Todos': q += f" AND filial={ph}"; p.append(f_uf)
    if f_fase != 'Todos': q += f" AND fase={ph}"; p.append(f_fase)
    if f_parecer != 'Todos': q += f" AND parecer={ph}"; p.append(f_parecer)
    if f_origem != 'Todos': q += f" AND origem_parecer={ph}"; p.append(f_origem)
    if f_resp != 'Todos': q += f" AND usuario_resp={ph}"; p.append(f_resp)
    if f_busca:
        termo = f"%{f_busca.upper()}%"
        q += f" AND (UPPER(cnpj_cpf) LIKE {ph} OR UPPER(nome) LIKE {ph} OR UPPER(fantasia) LIKE {ph} OR UPPER(processo) LIKE {ph})"
        p.extend([termo, termo, termo, termo])
    q += " ORDER BY id ASC"
    c.execute(q, tuple(p))
    rows_brutas = c.fetchall()
    conn.close()

    dados_exp = []
    for r in rows_brutas:
        info_sla = calcular_sla_credenciamento(r)
        if f_sla != 'Todos':
            if f_sla == '90d' and not info_sla['estourou_90d']:
                continue
            elif f_sla != '90d' and info_sla['faixa_sla'] != f_sla:
                continue
        dados_exp.append([
            r[8], info_sla['faixa_sla'], info_sla['dias_fase'], info_sla['dias_geral'],
            r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[9], r[11], r[12], r[13],
            r[14], r[15], r[16], r[19], r[20], r[21], r[22], r[23], r[25], r[28]
        ])

    cols = ["Processo", "Status_SLA", "SLA_Dias_Fase", "SLA_Dias_Geral", "UF", "CNPJ_CPF", "Nome", "Fantasia", "Tipo_Prestador", "Categoria", "Tipo_Processo", "Data_Abertura", "Responsavel", "Fase", "Ultima_Fase", "Data_Inicial", "Data_Final", "Proxima_Fase", "Motivo", "Parecer", "Observacao", "Origem", "Data_Aprovacao", "Aprovador", "Grupo"]
    df_exp = pd.DataFrame(dados_exp, columns=cols)
    output = BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_exp.to_excel(writer, sheet_name='Credenciamento_Fases', index=False)
    output.seek(0)
    return send_file(output, download_name=f"Credenciamento_Fases_{datetime.now().strftime('%Y_%m_%d_%Hh%Mm')}.xlsx", as_attachment=True)

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/admin', methods=['GET', 'POST'])
def admin():
    if 'user' not in session or session.get('role') != 'admin': return redirect(url_for('dashboard'))
    conn = conectar_db(); c = conn.cursor(); ph = get_ph(conn)
    if request.method == 'POST':
        acao = request.form.get('acao')
        if acao == 'adicionar':
            u = request.form.get('username'); p = request.form.get('password'); r = request.form.get('role')
            try:
                sincronizar_sequencias()
                c.execute(f"INSERT INTO users (username, password, role, first_login) VALUES ({ph}, {ph}, {ph}, 1)", (u, p, r))
                conn.commit(); flash(f"Usuário '{u}' criado com sucesso!", "success")
            except Exception: conn.rollback(); flash("Erro: Nome de usuário já existe!", "danger")
        elif acao == 'eliminar':
            uid = request.form.get('user_id')
            c.execute(f"DELETE FROM users WHERE id={ph}", (uid,))
            conn.commit(); flash("Usuário excluído com sucesso!", "warning")
    c.execute("SELECT id, username, role, first_login FROM users ORDER BY id ASC")
    usuarios = c.fetchall(); conn.close()
    return render_template_string(TELA_ADMIN, user=session['user'], role=session['role'], usuarios=usuarios)

@app.route('/admin/exportar')
def exportar_backup():
    if 'user' not in session or session.get('role') != 'admin': return redirect(url_for('dashboard'))
    try:
        conn = conectar_db()
        df_demands = pd.read_sql_query("SELECT * FROM demands", conn)
        df_avulsas = pd.read_sql_query("SELECT * FROM demandas_avulsas", conn)
        df_users = pd.read_sql_query("SELECT * FROM users", conn)
        try:
            df_cred = pd.read_sql_query("SELECT * FROM credenciamento_fases", conn)
        except Exception:
            df_cred = pd.DataFrame()
        conn.close()
        output = BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df_demands.to_excel(writer, sheet_name='Rotina_PDFs', index=False)
            df_avulsas.to_excel(writer, sheet_name='Demandas_Avulsas', index=False)
            df_users.to_excel(writer, sheet_name='Usuarios', index=False)
            if not df_cred.empty:
                df_cred.to_excel(writer, sheet_name='Credenciamento_Fases', index=False)
        output.seek(0)
        return send_file(output, download_name=f"Backup_COCAP_{datetime.now().strftime('%Y_%m_%d_%Hh%Mm')}.xlsx", as_attachment=True)
    except Exception as e: flash(f"Erro ao gerar backup: {e}", "danger"); return redirect(url_for('admin'))

@app.route('/admin/importar', methods=['POST'])
def importar_backup():
    if 'user' not in session or session.get('role') != 'admin': return redirect(url_for('dashboard'))
    file = request.files.get('file_backup')
    if not file: flash("Nenhum arquivo selecionado.", "danger"); return redirect(url_for('admin'))
    filename = file.filename.lower()
    try:
        conn = conectar_db(); c = conn.cursor(); is_postgres = "psycopg2" in str(type(conn))
        if filename.endswith('.xlsx'):
            xls = pd.ExcelFile(file); sheets = xls.sheet_names
            mapeamento = {'Rotina_PDFs': 'demands', 'Demandas_Avulsas': 'demandas_avulsas', 'Usuarios': 'users', 'Credenciamento_Fases': 'credenciamento_fases'}
            for aba, tabela in mapeamento.items():
                if aba in sheets:
                    df = pd.read_excel(xls, sheet_name=aba); c.execute(f"DELETE FROM {tabela}") 
                    if not df.empty:
                        df = df.where(pd.notnull(df), None); cols = ", ".join(df.columns); placeholders = ", ".join(["%s" if is_postgres else "?"] * len(df.columns))
                        c.executemany(f"INSERT INTO {tabela} ({cols}) VALUES ({placeholders})", df.values.tolist())
        elif filename.endswith('.db'):
            import sqlite3, tempfile
            temp_db = os.path.join(tempfile.gettempdir(), "temp_migration.db")
            file.save(temp_db) 
            sqlite_conn = sqlite3.connect(temp_db); sqlite_c = sqlite_conn.cursor()
            for tabela in ['demands', 'demandas_avulsas', 'users', 'credenciamento_fases']:
                try:
                    sqlite_c.execute(f"SELECT * FROM {tabela}"); rows = sqlite_c.fetchall()
                    if rows:
                        col_names = [desc[0] for desc in sqlite_c.description]; cols = ", ".join(col_names); placeholders = ", ".join(["%s" if is_postgres else "?"] * len(col_names))
                        c.execute(f"DELETE FROM {tabela}"); c.executemany(f"INSERT INTO {tabela} ({cols}) VALUES ({placeholders})", rows)
                except Exception: pass
            sqlite_conn.close(); os.remove(temp_db)
        conn.commit(); conn.close()
        sincronizar_sequencias()
        session.clear()
        flash("Base de Dados importada e sincronizada com sucesso! Inicie sessão novamente.", "success")
        return redirect(url_for('login'))
    except Exception as e: flash(f"Falha ao restaurar banco. Erro: {e}", "danger"); return redirect(url_for('admin'))

@app.route('/demandas')
def hub_demandas():
    if 'user' not in session: return redirect(url_for('login'))
    return render_template_string(TELA_HUB_DEMANDAS)

@app.route('/demandas/avulsas', methods=['GET', 'POST'])
def demandas_avulsas():
    if 'user' not in session: return redirect(url_for('login'))
    conn = conectar_db(); c = conn.cursor(); ph = get_ph(conn)
    if request.method == 'POST':
        acao = request.form.get('acao')
        try:
            if acao == 'nova':
                sincronizar_sequencias()
                td = request.form.get('tipo'); assunto = request.form.get('assunto')
                sol = request.form.get('solicitante'); prazo = request.form.get('prazo')
                dt_ent = datetime.now().strftime("%d/%m/%Y %H:%M")
                c.execute(f"INSERT INTO demandas_avulsas (tipo_demanda, assunto, solicitante, data_entrada, prazo, status, responsavel, observacao) VALUES ({ph}, {ph}, {ph}, {ph}, {ph}, 'Pendente', 'Nenhum', '')", (td, assunto, sol, dt_ent, prazo))
                flash("Demanda Avulsa criada com sucesso!", "success")
            elif acao == 'assumir':
                c.execute(f"UPDATE demandas_avulsas SET status='Em Análise', responsavel={ph} WHERE id={ph}", (session['user'], request.form.get('id')))
            elif acao == 'concluir':
                c.execute(f"UPDATE demandas_avulsas SET status='Concluído', data_conclusao={ph} WHERE id={ph}", (datetime.now().strftime("%d/%m/%Y"), request.form.get('id')))
            conn.commit()
        except Exception as e:
            conn.rollback()
            flash(f"Erro ao processar ação na Demanda Avulsa: {e}", "danger")

    s_filtro = request.args.get('status', 'Todos'); t_filtro = request.args.get('tipo', 'Todos'); r_filtro = request.args.get('resp', 'Todos')
    q = "SELECT * FROM demandas_avulsas WHERE 1=1"; p = []
    if s_filtro != 'Todos': q += f" AND status={ph}"; p.append(s_filtro)
    if t_filtro != 'Todos': q += f" AND tipo_demanda={ph}"; p.append(t_filtro)
    if r_filtro != 'Todos': q += f" AND responsavel={ph}"; p.append(r_filtro)
    c.execute(q + " ORDER BY id DESC", tuple(p)); demandas = c.fetchall()
    c.execute("SELECT username FROM users ORDER BY id ASC"); usuarios = ['Todos', 'Nenhum'] + [r[0] for r in c.fetchall()]; conn.close()
    return render_template_string(TELA_AVULSAS, demandas=demandas, tipos=TIPOS_AVULSA, sols=SOLICITANTES, usuarios=usuarios, status_list=STATUS_AVULSA, atual_status=s_filtro, atual_tipo=t_filtro, atual_resp=r_filtro, calc_sla=calcular_sla, calc_fechados=calcular_dias_fechados)

@app.route('/demandas/avulsas/editar', methods=['POST'])
def editar_avulsa():
    if 'user' not in session: return redirect(url_for('login'))
    d_id = request.form.get('id'); tipo = request.form.get('tipo'); assunto = request.form.get('assunto')
    solicitante = request.form.get('solicitante'); prazo = request.form.get('prazo'); status = request.form.get('status')
    resp = request.form.get('responsavel'); obs = request.form.get('observacao')

    conn = conectar_db(); c = conn.cursor(); ph = get_ph(conn)
    try:
        if str(status or '').strip().upper() in STATUS_ENCERRADOS_DEMANDAS:
            dt_fim = datetime.now().strftime("%d/%m/%Y")
            c.execute(f"UPDATE demandas_avulsas SET tipo_demanda={ph}, assunto={ph}, solicitante={ph}, prazo={ph}, status={ph}, responsavel={ph}, observacao={ph}, data_conclusao=COALESCE(NULLIF(data_conclusao, ''), {ph}) WHERE id={ph}",
                      (tipo, assunto, solicitante, prazo, status, resp, obs, dt_fim, d_id))
        else:
            c.execute(f"UPDATE demandas_avulsas SET tipo_demanda={ph}, assunto={ph}, solicitante={ph}, prazo={ph}, status={ph}, responsavel={ph}, observacao={ph} WHERE id={ph}",
                      (tipo, assunto, solicitante, prazo, status, resp, obs, d_id))
        conn.commit()
        flash("Demanda Avulsa atualizada com sucesso!", "success")
    except Exception as e:
        conn.rollback()
        flash(f"Erro ao editar demanda: {e}", "danger")
    finally:
        conn.close()
    return redirect(url_for('demandas_avulsas'))

@app.route('/demandas/rotinas')
def demandas_rotinas():
    if 'user' not in session: return redirect(url_for('login'))
    conn = conectar_db(); c = conn.cursor(); ph = get_ph(conn)
    s_filtro = request.args.get('status', 'Todos'); t_filtro = request.args.get('tipo', 'Todos'); r_filtro = request.args.get('resp', 'Todos')
    q = "SELECT * FROM demands WHERE 1=1"; p = []
    if s_filtro != 'Todos': q += f" AND status={ph}"; p.append(s_filtro)
    if t_filtro != 'Todos': q += f" AND type={ph}"; p.append(t_filtro)
    if r_filtro != 'Todos': q += f" AND assigned_to={ph}"; p.append(r_filtro)
    c.execute(q + " ORDER BY id DESC", tuple(p)); rotinas = c.fetchall()
    c.execute("SELECT username FROM users ORDER BY id ASC"); usuarios = ['Todos', 'Nenhum'] + [r[0] for r in c.fetchall()]; conn.close()
    return render_template_string(TELA_ROTINAS, rotinas=rotinas, tipos=TIPOS_ROTINA, status_list=STATUS_ROTINA, usuarios=usuarios, atual_status=s_filtro, atual_tipo=t_filtro, atual_resp=r_filtro, calc_sla=calcular_sla, calc_fechados=calcular_dias_fechados)

@app.route('/demandas/rotinas/upload', methods=['POST'])
def upload_rotinas():
    if 'user' not in session: return redirect(url_for('login'))
    arquivos = request.files.getlist('pdfs')
    if not arquivos or arquivos[0].filename == '':
        flash("Nenhum arquivo selecionado.", "warning")
        return redirect(url_for('demandas_rotinas'))
        
    sincronizar_sequencias()
    kw = {"REAJUSTE": "Reajuste", "NOVO CONTRATO": "Novo Contrato", "INCLUSÃO": "Inclusão", "INCLUSAO": "Inclusão", "EXCLUSÃO": "Exclusão", "EXCLUSAO": "Exclusão", "AJUSTE": "Ajuste", "EXTENSÃO": "Extensão", "EXTENSAO": "Extensão", "DESCREDENCIAMENTO": "Descredenciamento"}
    valid_ufs = {'AC', 'AL', 'AM', 'AP', 'BA', 'CE', 'DF', 'ES', 'GO', 'MA', 'MT', 'MS', 'MG', 'PA', 'PB', 'PR', 'PE', 'PI', 'RJ', 'RN', 'RS', 'RO', 'RR', 'SC', 'SP', 'SE', 'TO'}
    conn = conectar_db(); c = conn.cursor(); ph = get_ph(conn)
    count = 0; ignorados = 0
    
    for file in arquivos:
        try:
            f_name = file.filename
            c.execute(f"SELECT id FROM demands WHERE filename={ph}", (f_name,))
            if c.fetchone():
                ignorados += 1
                continue
            
            fu = f_name.upper(); dem = "Não Identificada"
            for k, v in kw.items():
                if k in fu: dem = v; break
            f_clean = re.sub(r'\.pdf$', '', fu, flags=re.IGNORECASE)
            docs_brutos = re.findall(r'\d{2}[.\s_-]*\d{3}[.\s_-]*\d{3}[.\s/-]*\d{4}[-\s_]*\d{2}|\d{3}[.\s_-]*\d{3}[.\s_-]*\d{3}[-\s_]*\d{2}|\b\d{14}\b|\b\d{11}\b', f_clean)
            cnpj_p = "Não encontrado"
            for c_ in docs_brutos:
                cn = re.sub(r'\D', '', c_)
                if len(cn) == 14: cnpj_p = f"{cn[:2]}.{cn[2:5]}.{cn[5:8]}/{cn[8:12]}-{cn[12:]}"; f_clean = f_clean.replace(c_, ''); break
                elif len(cn) == 11: cnpj_p = f"{cn[:3]}.{cn[3:6]}.{cn[6:9]}-{cn[9:]}"; f_clean = f_clean.replace(c_, ''); break
            mun, uf, prest = "-", "-", "Não encontrado"
            for k in kw.keys(): f_clean = re.sub(k, '', f_clean, flags=re.IGNORECASE)
            parts = [p.strip() for p in f_clean.split('-') if p.strip()]
            if len(parts) >= 2:
                p1 = parts[-1].upper(); p2 = parts[-2].upper() if len(parts) > 1 else ""
                if p1 in valid_ufs: uf = p1; mun = p2 if len(parts) >= 3 else "-"; prest = " ".join(parts[:-2]).strip() if len(parts) >= 3 else p2
                elif p2 in valid_ufs: uf = p2; mun = p1; prest = " ".join(parts[:-2]).strip()
                else:
                    m_uf_end = re.search(r'\b([A-Z]{2})\b$', p1)
                    if m_uf_end and m_uf_end.group(1) in valid_ufs: uf = m_uf_end.group(1); mun = re.sub(r'\b' + uf + r'\b$', '', p1).strip(); prest = " ".join(parts[:-1]).strip()
                    else: prest = " ".join(parts).strip()
            elif len(parts) == 1:
                m_uf_end = re.search(r'\b([A-Z]{2})\b$', parts[0].upper())
                if m_uf_end and m_uf_end.group(1) in valid_ufs: uf = m_uf_end.group(1); prest = re.sub(r'\b' + uf + r'\b$', '', parts[0]).strip()
                else: prest = parts[0]
            prest = re.sub(r'^[-_\s]+|[-_\s]+$', '', prest); prest = re.sub(r'\s+', ' ', prest).strip()
            if not prest or len(prest) <= 3: prest = "Não encontrado"
            
            if prest == "Não encontrado" or uf == "-" or cnpj_p == "Não encontrado":
                try:
                    pdf_bytes = BytesIO(file.read())
                    reader = PyPDF2.PdfReader(pdf_bytes); full_txt = ""
                    for page in reader.pages: full_txt += (page.extract_text() or "") + " "
                    if cnpj_p == "Não encontrado":
                        docs_text = re.findall(r'\d{2}[.\s]*\d{3}[.\s]*\d{3}[/\s]*\d{4}[-\s]*\d{2}|\b\d{14}\b', full_txt)
                        for c_ in docs_text:
                            cn = re.sub(r'\D', '', c_)
                            if len(cn) == 14: cnpj_p = f"{cn[:2]}.{cn[2:5]}.{cn[5:8]}/{cn[8:12]}-{cn[12:]}"; break
                except Exception:
                    pass
                    
            c.execute(f"INSERT INTO demands (filename, filepath, type, status, assigned_to, prestador, cnpj, municipio, uf, data_entrada) VALUES ({ph}, 'WEB_UPLOAD', {ph}, 'Pendente', 'Nenhum', {ph}, {ph}, {ph}, {ph}, {ph})", 
                         (f_name, dem, prest, cnpj_p, mun, uf, datetime.now().strftime("%d/%m/%Y %H:%M")))
            conn.commit()
            count += 1
        except Exception as e:
            conn.rollback()
            flash(f"Erro ao importar '{file.filename}': {e}", "danger")
            
    conn.close()
    msg = f"{count} PDF(s) processado(s) com sucesso!"
    if ignorados > 0: msg += f" ({ignorados} já existiam na base e foram ignorados)."
    flash(msg, "success" if count > 0 else "warning")
    return redirect(url_for('demandas_rotinas'))

@app.route('/demandas/rotinas/acao', methods=['POST'])
def acao_rotinas():
    d_id = request.form.get('id'); acao = request.form.get('acao')
    conn = conectar_db(); c = conn.cursor(); ph = get_ph(conn)
    try:
        if acao == 'assumir': c.execute(f"UPDATE demands SET status='Em Análise', assigned_to={ph} WHERE id={ph}", (session['user'], d_id))
        elif acao == 'finalizar': c.execute(f"UPDATE demands SET status='Finalizada', data_finalizacao={ph} WHERE id={ph}", (datetime.now().strftime("%d/%m/%Y %H:%M"), d_id))
        conn.commit()
    except Exception as e:
        conn.rollback()
        flash(f"Erro ao atualizar rotina: {e}", "danger")
    finally:
        conn.close()
    return redirect(url_for('demandas_rotinas'))

@app.route('/demandas/rotinas/editar', methods=['POST'])
def editar_rotina():
    if 'user' not in session: return redirect(url_for('login'))
    d_id = request.form.get('id'); status = request.form.get('status')
    tipo_neg = request.form.get('tipo_negociacao'); alcada = request.form.get('alcada')
    perc = request.form.get('impacto_perc'); rs = request.form.get('impacto_rs'); nup = request.form.get('nup_edoc')
    proc = request.form.get('tipo_processo'); fop = request.form.get('fop_autorizacao'); obs = request.form.get('observacao')
    resp = request.form.get('responsavel'); data_fop = request.form.get('data_aprovacao')

    if str(proc or '').strip().upper() == 'DEVOLVIDO' and status not in ['Finalizada', 'Cancelada']:
        status = 'Devolvido'

    conn = conectar_db(); c = conn.cursor(); ph = get_ph(conn)
    try:
        if str(status or '').strip().upper() in STATUS_ENCERRADOS_DEMANDAS:
            dt_fim = datetime.now().strftime("%d/%m/%Y %H:%M")
            c.execute(f"UPDATE demands SET status={ph}, assigned_to={ph}, tipo_negociacao={ph}, alcada={ph}, impacto_perc={ph}, impacto_rs={ph}, nup_edoc={ph}, tipo_processo={ph}, fop_autorizacao={ph}, data_aprovacao={ph}, observacao={ph}, data_finalizacao=COALESCE(NULLIF(data_finalizacao, ''), {ph}) WHERE id={ph}",
                      (status, resp, tipo_neg, alcada, perc, rs, nup, proc, fop, data_fop, obs, dt_fim, d_id))
        else:
            c.execute(f"UPDATE demands SET status={ph}, assigned_to={ph}, tipo_negociacao={ph}, alcada={ph}, impacto_perc={ph}, impacto_rs={ph}, nup_edoc={ph}, tipo_processo={ph}, fop_autorizacao={ph}, data_aprovacao={ph}, observacao={ph} WHERE id={ph}",
                      (status, resp, tipo_neg, alcada, perc, rs, nup, proc, fop, data_fop, obs, d_id))
        conn.commit()
        flash("Informações do processo atualizadas com sucesso!", "success")
    except Exception as e:
        conn.rollback()
        flash(f"Erro ao salvar informações: {e}", "danger")
    finally:
        conn.close()
    return redirect(url_for('demandas_rotinas'))

@app.route('/demandas/dashboard_graficos')
def dashboard_graficos():
    if 'user' not in session: return redirect(url_for('login'))
    inicializar_tabela_credenciamento()
    conn = conectar_db(); c = conn.cursor()
    c.execute("SELECT COUNT(*) FROM demands"); t_rotina = c.fetchone()[0]
    c.execute("SELECT COUNT(*) FROM demandas_avulsas"); t_avulsa = c.fetchone()[0]
    c.execute("SELECT COUNT(*) FROM credenciamento_fases"); t_cred_total = c.fetchone()[0]
    c.execute("SELECT COUNT(*) FROM credenciamento_fases WHERE ultima_fase='SIM'"); t_cred_atual = c.fetchone()[0]
    
    d_total_geral = t_rotina + t_avulsa + t_cred_atual
    
    c.execute("SELECT status, COUNT(*) FROM demands GROUP BY status")
    rot_por_status = {str(r[0]).strip(): r[1] for r in c.fetchall() if r[0]}
    
    c.execute("SELECT status, COUNT(*) FROM demandas_avulsas GROUP BY status")
    av_por_status = {str(r[0]).strip(): r[1] for r in c.fetchall() if r[0]}
    
    todos_status = list(TODOS_STATUS_ORDEM)
    for st in list(rot_por_status.keys()) + list(av_por_status.keys()):
        if st not in todos_status:
            todos_status.append(st)
            
    cores_status = {
        'Pendente': '#f0ad4e',
        'Recebida': '#5bc0de',
        'Em Análise': '#0275d8',
        'Em andamento': '#0dcaf0',
        'Aguardando 90 dias': '#6f42c1',
        'Aguardando Retorno': '#fd7e14',
        'Aguardando Área Demandante': '#d63384',
        'Devolvido': '#dc3545',
        'Finalizada': '#198754',
        'Concluído': '#20c997',
        'Cancelada': '#6c757d'
    }
    
    status_cards = []
    for st in todos_status:
        qtd_r = rot_por_status.get(st, 0)
        qtd_a = av_por_status.get(st, 0)
        soma_st = qtd_r + qtd_a
        if soma_st > 0 or st in ['Pendente', 'Em Análise', 'Aguardando 90 dias', 'Aguardando Retorno', 'Devolvido', 'Finalizada', 'Concluído']:
            status_cards.append({
                'nome': st,
                'rotina': qtd_r,
                'avulsa': qtd_a,
                'total': soma_st,
                'cor': cores_status.get(st, '#004b87')
            })
    
    c.execute("SELECT data_entrada, status FROM demands WHERE status NOT IN ('Finalizada', 'Concluído', 'Devolvido', 'Cancelada')")
    pendentes_rotina = c.fetchall()
    c.execute("SELECT data_entrada, status FROM demandas_avulsas WHERE status NOT IN ('Finalizada', 'Concluído', 'Devolvido', 'Cancelada')")
    pendentes_avulsas = c.fetchall()
    
    todas_pendentes = pendentes_rotina + pendentes_avulsas
    total_ativas = 0
    sla_em_dia = 0; sla_atencao = 0; sla_urgente = 0; sla_critico = 0
    for p in todas_pendentes:
        s = calcular_sla(p[0], p[1])
        if s < 0: continue
        total_ativas += 1
        if s >= 6: sla_critico += 1
        elif s >= 4: sla_urgente += 1
        elif s >= 2: sla_atencao += 1
        else: sla_em_dia += 1

    c.execute("SELECT * FROM credenciamento_fases WHERE ultima_fase='SIM'")
    rows_cred_atual = c.fetchall()
    cred_parecer = {}
    cred_sla = {"ativos": 0, "no_prazo": 0, "atencao": 0, "urgente": 0, "critico": 0}
    for rc in rows_cred_atual:
        par = str(rc[20] or 'SEM PARECER').strip().upper()
        cred_parecer[par] = cred_parecer.get(par, 0) + 1
        info_c = calcular_sla_credenciamento(rc)
        if info_c["ativo"]:
            cred_sla["ativos"] += 1
            if info_c["faixa_sla"] == "Crítico": cred_sla["critico"] += 1
            elif info_c["faixa_sla"] == "Urgente": cred_sla["urgente"] += 1
            elif info_c["faixa_sla"] == "Atenção": cred_sla["atencao"] += 1
            else: cred_sla["no_prazo"] += 1

    total_ativas_unificado = total_ativas + cred_sla["ativos"]

    fig1, ax1 = plt.subplots(figsize=(6,4))
    c.execute("SELECT assigned_to, COUNT(*) FROM demands WHERE status='Finalizada' AND assigned_to != 'Nenhum' GROUP BY assigned_to")
    dados_prod = c.fetchall()
    if dados_prod:
        ax1.bar([r[0] for r in dados_prod], [r[1] for r in dados_prod], color='#5cb85c')
        ax1.set_xticklabels([r[0] for r in dados_prod], rotation=45, ha='right')
    buf1 = BytesIO(); fig1.tight_layout(); fig1.savefig(buf1, format="png"); buf1.seek(0)
    chart_prod = base64.b64encode(buf1.read()).decode('utf-8'); plt.close(fig1)

    fig2, ax2 = plt.subplots(figsize=(6,4))
    c.execute("SELECT type, COUNT(*) FROM demands GROUP BY type")
    dados_tipo = c.fetchall()
    if dados_tipo:
        ax2.bar([r[0] for r in dados_tipo], [r[1] for r in dados_tipo], color='#008CBA')
        ax2.set_xticklabels([r[0] for r in dados_tipo], rotation=45, ha='right')
    buf2 = BytesIO(); fig2.tight_layout(); fig2.savefig(buf2, format="png"); buf2.seek(0)
    chart_tipo = base64.b64encode(buf2.read()).decode('utf-8'); plt.close(fig2)
    conn.close()
    return render_template_string(
        TELA_DASHBOARD_GRAFICOS,
        d_total_geral=d_total_geral, t_rotina=t_rotina, t_avulsa=t_avulsa,
        t_cred_total=t_cred_total, t_cred_atual=t_cred_atual,
        status_cards=status_cards, cred_parecer=cred_parecer, cred_sla=cred_sla,
        chart_prod=chart_prod, chart_tipo=chart_tipo,
        sla_em_dia=sla_em_dia, sla_atencao=sla_atencao, sla_urgente=sla_urgente, sla_critico=sla_critico,
        total_ativas_unificado=total_ativas_unificado
    )

# -------- GERADOR DE RELATÓRIO EXECUTIVO GERAL EM PDF --------
@app.route('/demandas/relatorio_pdf')
def exportar_relatorio_pdf():
    if 'user' not in session: return redirect(url_for('login'))
    inicializar_tabela_credenciamento()
    conn = conectar_db(); c = conn.cursor()

    c.execute("SELECT COUNT(*) FROM demands"); t_rotina = c.fetchone()[0]
    c.execute("SELECT COUNT(*) FROM demandas_avulsas"); t_avulsa = c.fetchone()[0]
    c.execute("SELECT COUNT(*) FROM credenciamento_fases WHERE ultima_fase='SIM'"); t_cred_atual = c.fetchone()[0]
    d_total_geral = t_rotina + t_avulsa + t_cred_atual

    c.execute("SELECT status, COUNT(*) FROM demands GROUP BY status")
    rot_por_status = {str(r[0]).strip(): r[1] for r in c.fetchall() if r[0]}
    c.execute("SELECT status, COUNT(*) FROM demandas_avulsas GROUP BY status")
    av_por_status = {str(r[0]).strip(): r[1] for r in c.fetchall() if r[0]}

    todos_status = list(TODOS_STATUS_ORDEM)
    for st in list(rot_por_status.keys()) + list(av_por_status.keys()):
        if st not in todos_status: todos_status.append(st)

    tabela_status_rows = []
    for st in todos_status:
        qr = rot_por_status.get(st, 0)
        qa = av_por_status.get(st, 0)
        if (qr + qa) > 0 or st in ['Pendente', 'Em Análise', 'Aguardando 90 dias', 'Aguardando Retorno', 'Devolvido', 'Finalizada', 'Concluído']:
            tabela_status_rows.append([st, str(qr), str(qa), str(qr + qa)])

    c.execute("SELECT data_entrada, status FROM demands WHERE status NOT IN ('Finalizada', 'Concluído', 'Devolvido', 'Cancelada')")
    pend_rot = c.fetchall()
    c.execute("SELECT data_entrada, status FROM demandas_avulsas WHERE status NOT IN ('Finalizada', 'Concluído', 'Devolvido', 'Cancelada')")
    pend_av = c.fetchall()

    sla_dem = {"no_prazo": 0, "atencao": 0, "urgente": 0, "critico": 0}
    for p in (pend_rot + pend_av):
        s = calcular_sla(p[0], p[1])
        if s < 0: continue
        if s >= 6: sla_dem["critico"] += 1
        elif s >= 4: sla_dem["urgente"] += 1
        elif s >= 2: sla_dem["atencao"] += 1
        else: sla_dem["no_prazo"] += 1

    c.execute("SELECT * FROM credenciamento_fases WHERE ultima_fase='SIM'")
    rows_cred = c.fetchall()
    cred_parecer = {}
    cred_fases_dict = {}
    sla_cred = {"no_prazo": 0, "atencao": 0, "urgente": 0, "critico": 0, "alerta_90d": 0}
    for rc in rows_cred:
        par = str(rc[20] or 'SEM PARECER').strip().upper()
        cred_parecer[par] = cred_parecer.get(par, 0) + 1
        fa = str(rc[12] or 'SEM FASE').strip()
        cred_fases_dict[fa] = cred_fases_dict.get(fa, 0) + 1
        info_c = calcular_sla_credenciamento(rc)
        if info_c["ativo"]:
            if info_c["faixa_sla"] == "Crítico": sla_cred["critico"] += 1
            elif info_c["faixa_sla"] == "Urgente": sla_cred["urgente"] += 1
            elif info_c["faixa_sla"] == "Atenção": sla_cred["atencao"] += 1
            else: sla_cred["no_prazo"] += 1
            if info_c["estourou_90d"]: sla_cred["alerta_90d"] += 1

    c.execute("SELECT assigned_to, COUNT(*) FROM demands WHERE status='Finalizada' AND assigned_to != 'Nenhum' GROUP BY assigned_to ORDER BY COUNT(*) DESC")
    dados_prod = c.fetchall()

    c.execute("SELECT type, COUNT(*) FROM demands GROUP BY type ORDER BY COUNT(*) DESC")
    dados_tipo = c.fetchall()
    conn.close()

    pdf_buffer = BytesIO()
    dt_emissao = datetime.now().strftime("%d/%m/%Y %H:%M")

    with PdfPages(pdf_buffer) as pdf:
        # ==========================================================
        # PÁGINA 1: RESUMO EXECUTIVO E MONITORAMENTO DE SLA
        # ==========================================================
        fig1 = plt.figure(figsize=(11.69, 8.27))
        fig1.patch.set_facecolor('#ffffff')

        fig1.text(0.04, 0.94, "POSTAL SAÚDE - CENTRAL DE OPERAÇÕES (COCAP)", fontsize=16, fontweight='bold', color='#004b87')
        fig1.text(0.04, 0.91, f"Relatório Executivo Geral do Cenário Operacional | Emitido em: {dt_emissao} por {session.get('user', 'Sistema')}", fontsize=10, color='#555555')

        # 1. Tabela de Indicadores Gerais (KPIs das 3 Frentes)
        ax_kpi = fig1.add_axes([0.04, 0.77, 0.92, 0.11])
        ax_kpi.axis('off')
        ax_kpi.set_title("1. Volume Consolidado das 3 Frentes Operacionais", fontsize=11, fontweight='bold', loc='left', color='#004b87', pad=6)
        kpi_data = [[
            f"{d_total_geral}",
            f"{t_rotina}",
            f"{t_avulsa}",
            f"{t_cred_atual}"
        ]]
        kpi_cols = ["TOTAL GERAL CONSOLIDADO", "1. Rotinas (PDFs)", "2. Demandas Avulsas", "3. Credenciamento (Fase Atual)"]
        t1 = ax_kpi.table(cellText=kpi_data, colLabels=kpi_cols, loc='center', cellLoc='center')
        t1.auto_set_font_size(False); t1.set_fontsize(11); t1.scale(1, 2.0)
        for (row_i, col_i), cell in t1.get_celld().items():
            if row_i == 0:
                cell.set_facecolor('#004b87'); cell.set_text_props(color='white', fontweight='bold', fontsize=9.5)
            else:
                cell.set_facecolor('#f0f4f8'); cell.set_text_props(fontweight='bold', fontsize=13, color='#002060')

        # 2. Tabela Completa por Status (PDFs + Avulsas)
        ax_st = fig1.add_axes([0.04, 0.10, 0.46, 0.60])
        ax_st.axis('off')
        ax_st.set_title("2. Contagem Detalhada por Status (PDFs + Avulsas)", fontsize=11, fontweight='bold', loc='left', color='#004b87', pad=8)
        if tabela_status_rows:
            t_st = ax_st.table(cellText=tabela_status_rows, colLabels=["Status da Demanda", "PDFs", "Avulsas", "Total"], loc='upper center', cellLoc='center', colWidths=[0.46, 0.18, 0.18, 0.18])
            t_st.auto_set_font_size(False); t_st.set_fontsize(9); t_st.scale(1, 1.55)
            for (r_i, c_i), cell in t_st.get_celld().items():
                if r_i == 0:
                    cell.set_facecolor('#333333'); cell.set_text_props(color='white', fontweight='bold')
                elif r_i % 2 == 0:
                    cell.set_facecolor('#f8f9fa')

        # 3. Gráfico de Monitoramento Unificado de SLA
        ax_sla = fig1.add_axes([0.56, 0.43, 0.40, 0.27])
        categorias_sla = ["No Prazo\n(0-1d)", "Atenção\n(2-3d)", "Urgente\n(4-5d)", "Crítico\n(6+d)"]
        vals_dem = [sla_dem["no_prazo"], sla_dem["atencao"], sla_dem["urgente"], sla_dem["critico"]]
        vals_crd = [sla_cred["no_prazo"], sla_cred["atencao"], sla_cred["urgente"], sla_cred["critico"]]
        x_pos = range(len(categorias_sla))
        ax_sla.bar(x_pos, vals_dem, label='Demandas (PDF/Avulsa)', color='#0275d8', width=0.5)
        ax_sla.bar(x_pos, vals_crd, bottom=vals_dem, label='Credenciamento (Ativos)', color='#fd7e14', width=0.5)
        ax_sla.set_xticks(list(x_pos)); ax_sla.set_xticklabels(categorias_sla, fontsize=8.5)
        ax_sla.set_title("3. Fila Ativa por Faixa de SLA (Unificado)", fontsize=11, fontweight='bold', color='#004b87')
        ax_sla.legend(fontsize=8)
        ax_sla.grid(axis='y', linestyle='--', alpha=0.4)

        # 4. Tabela Resumo de SLA e Credenciamento na Página 1
        ax_res_sla = fig1.add_axes([0.56, 0.08, 0.40, 0.27])
        ax_res_sla.axis('off')
        ax_res_sla.set_title("4. Detalhamento da Fila de SLA (Em Aberto)", fontsize=10.5, fontweight='bold', loc='left', color='#004b87', pad=6)
        rows_sla_tb = [
            ["No Prazo (0 a 1 dia)", str(sla_dem["no_prazo"]), str(sla_cred["no_prazo"]), str(sla_dem["no_prazo"] + sla_cred["no_prazo"])],
            ["Atenção (2 a 3 dias)", str(sla_dem["atencao"]), str(sla_cred["atencao"]), str(sla_dem["atencao"] + sla_cred["atencao"])],
            ["Urgente (4 a 5 dias)", str(sla_dem["urgente"]), str(sla_cred["urgente"]), str(sla_dem["urgente"] + sla_cred["urgente"])],
            ["Crítico (6+ dias)", str(sla_dem["critico"]), str(sla_cred["critico"]), str(sla_dem["critico"] + sla_cred["critico"])],
            ["Alerta > 90 Dias (Geral)", "-", str(sla_cred["alerta_90d"]), str(sla_cred["alerta_90d"])]
        ]
        t_sla = ax_res_sla.table(cellText=rows_sla_tb, colLabels=["Faixa de SLA", "Demandas", "Credenc.", "Total"], loc='upper center', cellLoc='center', colWidths=[0.46, 0.18, 0.18, 0.18])
        t_sla.auto_set_font_size(False); t_sla.set_fontsize(8.5); t_sla.scale(1, 1.45)
        cores_linhas_sla = ['#e2e3e5', '#fff3cd', '#ffe8cc', '#f8d7da', '#e8d9f8']
        for (r_i, c_i), cell in t_sla.get_celld().items():
            if r_i == 0:
                cell.set_facecolor('#333333'); cell.set_text_props(color='white', fontweight='bold')
            else:
                cell.set_facecolor(cores_linhas_sla[r_i - 1])

        pdf.savefig(fig1)
        plt.close(fig1)

        # ==========================================================
        # PÁGINA 2: PRODUTIVIDADE E CENÁRIO DO CREDENCIAMENTO - FASES
        # ==========================================================
        fig2 = plt.figure(figsize=(11.69, 8.27))
        fig2.patch.set_facecolor('#ffffff')

        fig2.text(0.04, 0.94, "POSTAL SAÚDE - PRODUTIVIDADE & CREDENCIAMENTO EM FASES", fontsize=15, fontweight='bold', color='#004b87')
        fig2.text(0.04, 0.91, f"Página 2 - Análise de Produtividade, Tipos de Demanda e Pipeline de Credenciamento ({dt_emissao})", fontsize=9.5, color='#555555')

        # Gráfico 1: Produtividade por Colaborador
        ax_p1 = fig2.add_axes([0.06, 0.54, 0.41, 0.31])
        if dados_prod:
            nomes_p = [str(r[0]) for r in dados_prod]
            vals_p = [r[1] for r in dados_prod]
            ax_p1.bar(nomes_p, vals_p, color='#198754')
            ax_p1.set_xticklabels(nomes_p, rotation=30, ha='right', fontsize=8)
        else:
            ax_p1.text(0.5, 0.5, "Sem dados de demandas finalizadas", ha='center', va='center')
        ax_p1.set_title("5. Produtividade por Colaborador (Rotinas Finalizadas)", fontsize=10.5, fontweight='bold', color='#004b87')
        ax_p1.grid(axis='y', linestyle='--', alpha=0.4)

        # Gráfico 2: Distribuição por Tipo de Demanda
        ax_p2 = fig2.add_axes([0.55, 0.54, 0.41, 0.31])
        if dados_tipo:
            nomes_t = [str(r[0]) for r in dados_tipo]
            vals_t = [r[1] for r in dados_tipo]
            ax_p2.bar(nomes_t, vals_t, color='#008CBA')
            ax_p2.set_xticklabels(nomes_t, rotation=30, ha='right', fontsize=8)
        else:
            ax_p2.text(0.5, 0.5, "Sem dados por tipo", ha='center', va='center')
        ax_p2.set_title("6. Volume por Tipo de Demanda (Rotinas)", fontsize=10.5, fontweight='bold', color='#004b87')
        ax_p2.grid(axis='y', linestyle='--', alpha=0.4)

        # Tabela de Fases do Credenciamento (Fase Atual = SIM)
        ax_cf = fig2.add_axes([0.04, 0.06, 0.92, 0.38])
        ax_cf.axis('off')
        em_ana_c = cred_parecer.get('EM ANÁLISE', 0) + cred_parecer.get('EM ANALISE', 0) + cred_parecer.get('PENDENTE', 0)
        def_c = cred_parecer.get('DEFERIDO', 0) + cred_parecer.get('FINALIZADA', 0)
        dev_c = cred_parecer.get('DEVOLVIDO', 0)
        ind_c = cred_parecer.get('INDEFERIDO', 0)
        ax_cf.set_title(
            f"7. Pipeline de Credenciamento - Fases Atuais (Em Análise: {em_ana_c} | Deferidos: {def_c} | Devolvidos: {dev_c} | Indeferidos: {ind_c})",
            fontsize=10.5, fontweight='bold', loc='left', color='#004b87', pad=8
        )
        fases_ordenadas = sorted(cred_fases_dict.items(), key=lambda x: x[0])
        if fases_ordenadas:
            rows_fases = [[f_n, str(f_q)] for f_n, f_q in fases_ordenadas]
            t_cf = ax_cf.table(cellText=rows_fases, colLabels=["Fase Atual do Processo (ULTIMAFASE = SIM)", "Qtd. Processos"], loc='upper center', cellLoc='left', colWidths=[0.80, 0.20])
            t_cf.auto_set_font_size(False); t_cf.set_fontsize(8.5); t_cf.scale(1, 1.35)
            for (r_i, c_i), cell in t_cf.get_celld().items():
                if r_i == 0:
                    cell.set_facecolor('#004b87'); cell.set_text_props(color='white', fontweight='bold')
                elif r_i % 2 == 0:
                    cell.set_facecolor('#f8f9fa')
        else:
            ax_cf.text(0.5, 0.5, "Nenhum registro carregado no módulo Credenciamento - Fases.", ha='center', va='center', fontsize=10, color='#777777')

        pdf.savefig(fig2)
        plt.close(fig2)

    pdf_buffer.seek(0)
    return send_file(
        pdf_buffer,
        download_name=f"Relatorio_Executivo_COCAP_{datetime.now().strftime('%Y_%m_%d_%Hh%Mm')}.pdf",
        as_attachment=True,
        mimetype='application/pdf'
    )

@app.route('/pentefino')
def hub_pentefino():
    if 'user' not in session: return redirect(url_for('login'))
    return render_template_string(TELA_HUB_PENTEFINO)

@app.route('/pentefino/regional', methods=['GET', 'POST'])
def pentefino_regional():
    if 'user' not in session: return redirect(url_for('login'))
    if request.method == 'GET': return render_template_string(TELA_PENTEFINO_REGIONAL)
        
    try:
        f_postal = request.files.get('f_postal'); f_op = request.files.get('f_op'); f_ibge = request.files.get('f_ibge')
        df_postal = ler_arquivo(f_postal); df_operadora = ler_arquivo(f_op); df_ibge = ler_arquivo(f_ibge)
        
        df_postal.columns = [str(c).strip().upper() for c in df_postal.columns]
        df_operadora.columns = [str(c).strip().upper() for c in df_operadora.columns]
        df_ibge.columns = [str(c).strip().upper() for c in df_ibge.columns]

        col_ibge_7 = cacador_de_colunas(df_ibge, ['CD_MUNICIPIO', 'COMPLETO', 'IBGE'])
        col_ibge_6 = cacador_de_colunas(df_ibge, ['AJUSTADO', 'CÓD. MUNIC'])
        c_regiao = cacador_de_colunas(df_ibge, ['NM_REGIAO_SAUDE', 'REGIÃO DE SAÚDE', 'REGIAO DE SAUDE', 'NOME DA REGIÃO DE SAÚDE', 'REGIÃO'])
        col_mun_csv = cacador_de_colunas(df_ibge, ['NM_MUNICIPIO', 'MUNICÍPIO', 'MUNICIPIO', 'CIDADE'])
        col_uf_csv = cacador_de_colunas(df_ibge, ['SG_UF', 'UF', 'ESTADO'])

        map_ibge_regiao = {}
        if col_ibge_7: map_ibge_regiao.update(dict(zip(df_ibge[col_ibge_7].apply(limpar_ibge), df_ibge[c_regiao].apply(lambda x: str(x).strip().upper() if pd.notna(x) else ""))))
        if col_ibge_6: map_ibge_regiao.update(dict(zip(df_ibge[col_ibge_6].apply(limpar_ibge), df_ibge[c_regiao].apply(lambda x: str(x).strip().upper() if pd.notna(x) else ""))))
        df_ibge['CHAVE_GEO'] = df_ibge[col_mun_csv].apply(normalizar_texto) + "_" + df_ibge[col_uf_csv].apply(normalizar_texto)
        map_geo_regiao = dict(zip(df_ibge['CHAVE_GEO'], df_ibge[c_regiao].apply(lambda x: str(x).strip().upper() if pd.notna(x) else "")))

        dics_processados = []
        for df in [df_postal, df_operadora]:
            col_ibge = cacador_de_colunas(df, ['IBGE', 'CÓDIGO IBGE']); col_mun = cacador_de_colunas(df, ['MUNICÍPIO', 'MUNICIPIO', 'CIDADE']); col_uf = cacador_de_colunas(df, ['UF', 'ESTADO'])
            if col_ibge: df['REG_IBGE'] = df[col_ibge].apply(limpar_ibge).map(map_ibge_regiao)
            else: df['REG_IBGE'] = None
            if col_mun and col_uf: df['REG_GEO'] = (df[col_mun].apply(normalizar_texto) + "_" + df[col_uf].apply(normalizar_texto)).map(map_geo_regiao)
            else: df['REG_GEO'] = None
            df['REGIAO_SAUDE'] = df['REG_IBGE'].fillna(df['REG_GEO']).fillna("REGIÃO NÃO IDENTIFICADA")
            dics_processados.append(df)

        df_postal, df_operadora = dics_processados[0], dics_processados[1]

        v_tipo_pos = request.form.get('col_tipo_pos', '').strip().upper()
        v_esp_pos = request.form.get('col_esp_pos', '').strip().upper()
        v_tipo_op = request.form.get('col_tipo_op', '').strip().upper()
        v_esp_op = request.form.get('col_esp_op', '').strip().upper()

        col_cnpj_pos = cacador_de_colunas(df_postal, ['CNPJ', 'CPFCNPJ'])
        col_nome_pos = cacador_de_colunas(df_postal, ['NOME', 'RAZAO', 'PRESTADOR', 'FANTASIA'])
        col_mun_pos = cacador_de_colunas(df_postal, ['MUNICÍPIO', 'MUNICIPIO', 'CIDADE'])
        
        col_esp_pos = v_esp_pos if v_esp_pos in df_postal.columns else cacador_de_colunas(df_postal, ['ESPECIALIDADE', 'ESPECIALIDADES'])
        col_tipo_pos = v_tipo_pos if v_tipo_pos in df_postal.columns else cacador_de_colunas(df_postal, ['TIPOPRESTADOR', 'TIPO_PRESTADOR', 'TIPO PRESTADOR', 'TIPO'])
        col_esp_op = v_esp_op if v_esp_op in df_operadora.columns else cacador_de_colunas(df_operadora, ['ESPECIALIDADE', 'ESPECIALIDADES'])

        if not col_cnpj_pos: df_postal['CNPJ_TEMP'] = "S/CNPJ"; col_cnpj_pos = 'CNPJ_TEMP'
        if not col_nome_pos: df_postal['NOME_TEMP'] = "S/NOME"; col_nome_pos = 'NOME_TEMP'
        if not col_mun_pos: df_postal['MUN_TEMP'] = "S/MUNICIPIO"; col_mun_pos = 'MUN_TEMP'
        df_postal['CNPJ_LIMPO'] = df_postal[col_cnpj_pos].astype(str).str.replace(r'\D', '', regex=True)

        dict_agg = {}
        if col_esp_pos: dict_agg[col_esp_pos] = lambda s: ", ".join(sorted(list(extrair_especialidades(s))))
        if col_tipo_pos: dict_agg[col_tipo_pos] = lambda x: " | ".join(x.dropna().astype(str).unique())
        df_postal_agrupado = df_postal.groupby(['REGIAO_SAUDE', 'CNPJ_LIMPO', col_nome_pos, col_mun_pos], dropna=False).agg(dict_agg).reset_index()

        regioes_operadora = set(df_operadora['REGIAO_SAUDE'].dropna().unique())
        regioes_operadora.discard("REGIÃO NÃO IDENTIFICADA")
        df_postal_alvos = df_postal_agrupado[df_postal_agrupado['REGIAO_SAUDE'].isin(regioes_operadora)].copy()

        piscina_operadora = {}
        for regiao, grupo in df_operadora.groupby('REGIAO_SAUDE'):
            if regiao == "NÃO ENCONTRADO": continue
            piscina_operadora[regiao] = extrair_especialidades(grupo[col_esp_op]) if col_esp_op else set()

        res_reg = {}
        for _, row in df_postal_alvos.iterrows():
            regiao = row.get('REGIAO_SAUDE')
            if regiao not in res_reg: res_reg[regiao] = []
            espec_str = str(row.get(col_esp_pos, '')) if col_esp_pos else ""
            if not espec_str.strip() or espec_str.upper() == 'NAN': obs = "Vulnerabilidade: Especialidades não informadas na base da Postal."; espec_alvo = set()
            else: espec_alvo = set([e.strip() for e in espec_str.split(',') if e.strip()]); obs = ""
            
            if not obs:
                faltantes = espec_alvo - piscina_operadora.get(regiao, set())
                obs = "Vulnerabilidade: " + ", ".join(sorted(list(faltantes))) if faltantes else "Cobertura Total"
            res_reg[regiao].append({"CNPJ": str(row.get('CNPJ_LIMPO', 'N/A')), "Prestador": str(row.get(col_nome_pos, 'N/A')), "Município": str(row.get(col_mun_pos, 'N/A')), "OBSERVAÇÃO": obs})

        for regiao in regioes_operadora:
            if regiao not in res_reg: res_reg[regiao] = [{"CNPJ": "", "Prestador": "", "Município": "", "OBSERVAÇÃO": "Sem prestadores postal"}]

        output = BytesIO()
        wb = Workbook(); ws = wb.active; ws.title = "Analise Equivalencia"
        fill_regiao = PatternFill(start_color="005580", end_color="005580", fill_type="solid"); font_regiao = Font(color="FFFFFF", bold=True, size=11)
        fill_colunas = PatternFill(start_color="005580", end_color="005580", fill_type="solid"); font_colunas = Font(color="FFFFFF", bold=True)
        align_center = Alignment(horizontal="center", vertical="center"); align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        linha_atual = 1
        for regiao in sorted(res_reg.keys()):
            ws.merge_cells(start_row=linha_atual, start_column=1, end_row=linha_atual, end_column=4)
            c_reg = ws.cell(row=linha_atual, column=1, value=str(regiao).upper()); c_reg.fill = fill_regiao; c_reg.font = font_regiao; c_reg.alignment = align_center; linha_atual += 1
            for col_idx, texto in enumerate(["CNPJ", "Prestador", "Município", "OBSERVAÇÃO"], 1):
                cell = ws.cell(row=linha_atual, column=col_idx, value=texto); cell.fill = fill_colunas; cell.font = font_colunas; cell.alignment = align_left; cell.border = thin_border
            linha_atual += 1
            for item in res_reg[regiao]:
                c1 = ws.cell(row=linha_atual, column=1, value=item['CNPJ']); c2 = ws.cell(row=linha_atual, column=2, value=item['Prestador'])
                c3 = ws.cell(row=linha_atual, column=3, value=item['Município']); c4 = ws.cell(row=linha_atual, column=4, value=item['OBSERVAÇÃO'])
                for c in [c1, c2, c3, c4]: c.alignment = align_left; c.border = thin_border
                linha_atual += 1
            linha_atual += 1

        ws.column_dimensions['A'].width = 20; ws.column_dimensions['B'].width = 45; ws.column_dimensions['C'].width = 25; ws.column_dimensions['D'].width = 65
        wb.save(output); output.seek(0)
        return send_file(output, download_name=f"Relatorio_Regional_{datetime.now().strftime('%Hh%Mm')}.xlsx", as_attachment=True)
    except Exception as e: flash(f"Erro ao processar as planilhas: {e}", "danger"); return redirect(url_for('pentefino_regional'))

@app.route('/pentefino/individual', methods=['GET', 'POST'])
def pentefino_individual():
    if 'user' not in session: return redirect(url_for('login'))
    if request.method == 'GET': return render_template_string(TELA_PENTEFINO_INDIVIDUAL)
        
    try:
        f_postal = request.files.get('f_postal'); f_ibge = request.files.get('f_ibge')
        alvo_bruto = request.form.get('alvo', '').strip().upper()
        filtro_tipo = request.form.get('filtro_tipo', '').strip().upper()
        
        df_ind = ler_arquivo(f_postal); df_ibge = ler_arquivo(f_ibge)
        df_ind.columns = [str(c).strip().upper() for c in df_ind.columns]
        df_ibge.columns = [str(c).strip().upper() for c in df_ibge.columns]

        col_cnpj_ind = cacador_de_colunas(df_ind, ['CNPJ', 'CPFCNPJ'])
        col_nome_ind = cacador_de_colunas(df_ind, ['NOME', 'RAZAO', 'PRESTADOR', 'FANTASIA'])
        col_mun_ind = cacador_de_colunas(df_ind, ['MUNICÍPIO', 'MUNICIPIO', 'CIDADE'])
        col_uf_ind = cacador_de_colunas(df_ind, ['UF', 'ESTADO'])
        col_tipo_ind = cacador_de_colunas(df_ind, ['TIPOPRESTADOR', 'TIPO_PRESTADOR', 'TIPO PRESTADOR', 'TIPO', 'CATEGORIA', 'TIPO DE PRESTADOR'])
        col_esp_ind = cacador_de_colunas(df_ind, ['ESPECIALIDADE', 'ESPECIALIDADES'])
        col_ibge_ind = cacador_de_colunas(df_ind, ['IBGE', 'CÓDIGO IBGE'])
        
        if not col_cnpj_ind or not col_nome_ind: flash("Colunas básicas não encontradas.", "danger"); return redirect(url_for('pentefino_individual'))
            
        df_ind[col_cnpj_ind] = df_ind[col_cnpj_ind].astype(str).str.replace(r'\D', '', regex=True)
        
        alvo_numeros = re.sub(r'\D', '', alvo_bruto)
        mask = df_ind[col_nome_ind].fillna('').astype(str).str.upper().str.contains(alvo_bruto, na=False, regex=False)
        if len(alvo_numeros) > 3: mask = mask | (df_ind[col_cnpj_ind].str.contains(alvo_numeros, na=False, regex=False))
        
        linhas_alvo = df_ind[mask]
        if linhas_alvo.empty: flash(f"O Alvo '{alvo_bruto}' não foi encontrado na base.", "danger"); return redirect(url_for('pentefino_individual'))
            
        row_alvo = linhas_alvo.iloc[0]
        cnpj_alvo = str(row_alvo.get(col_cnpj_ind, '')).strip()
        nome_alvo_real = str(row_alvo.get(col_nome_ind, '')).strip()

        dict_agg = {}
        if col_esp_ind: dict_agg[col_esp_ind] = lambda s: ", ".join(sorted(list(extrair_especialidades(s))))
        if col_tipo_ind: dict_agg[col_tipo_ind] = lambda x: " | ".join(x.dropna().astype(str).unique())
        group_cols = [c for c in [col_cnpj_ind, col_nome_ind, col_mun_ind, col_uf_ind, col_ibge_ind] if c]
        df_ind = df_ind.groupby(group_cols, dropna=False).agg(dict_agg).reset_index()

        c_ibge_7 = cacador_de_colunas(df_ibge, ['CD_MUNICIPIO', 'COMPLETO', 'IBGE'])
        c_ibge_6 = cacador_de_colunas(df_ibge, ['AJUSTADO', 'CÓD. MUNIC'])
        c_regiao = cacador_de_colunas(df_ibge, ['NM_REGIAO_SAUDE', 'REGIÃO DE SAÚDE', 'REGIAO DE SAUDE', 'NOME DA REGIÃO DE SAÚDE', 'REGIÃO'])
        c_mun = cacador_de_colunas(df_ibge, ['NM_MUNICIPIO', 'MUNICÍPIO', 'MUNICIPIO', 'CIDADE'])
        c_uf = cacador_de_colunas(df_ibge, ['SG_UF', 'UF', 'ESTADO'])

        map_ibge_regiao = {}
        if c_ibge_7: map_ibge_regiao.update(dict(zip(df_ibge[c_ibge_7].apply(limpar_ibge), df_ibge[c_regiao].apply(lambda x: str(x).strip().upper() if pd.notna(x) else ""))))
        if c_ibge_6: map_ibge_regiao.update(dict(zip(df_ibge[c_ibge_6].apply(limpar_ibge), df_ibge[c_regiao].apply(lambda x: str(x).strip().upper() if pd.notna(x) else ""))))
        df_ibge['CHAVE_GEO'] = df_ibge[c_mun].apply(normalizar_texto) + "_" + df_ibge[c_uf].apply(normalizar_texto)
        map_geo_regiao = dict(zip(df_ibge['CHAVE_GEO'], df_ibge[c_regiao].apply(lambda x: str(x).strip().upper() if pd.notna(x) else "")))

        linhas_alvo_agrupadas = df_ind[df_ind[col_cnpj_ind] == cnpj_alvo]
        if not linhas_alvo_agrupadas.empty: row_alvo = linhas_alvo_agrupadas.iloc[0]

        mun_alvo = str(row_alvo.get(col_mun_ind, '')).strip().upper(); uf_alvo = str(row_alvo.get(col_uf_ind, '')).strip().upper()
        chave_geo_alvo = normalizar_texto(mun_alvo) + "_" + normalizar_texto(uf_alvo)
        tipo_alvo = str(row_alvo.get(col_tipo_ind, '')).strip().upper() if col_tipo_ind else ""
        
        ibge_alvo = str(row_alvo.get(col_ibge_ind, '')) if col_ibge_ind else None
        regiao_alvo = None
        if ibge_alvo and str(ibge_alvo).upper() not in ['NAN', 'NONE', 'NULL', '']: regiao_alvo = map_ibge_regiao.get(limpar_ibge(ibge_alvo))
        if not regiao_alvo: regiao_alvo = map_geo_regiao.get(chave_geo_alvo, "NÃO ENCONTRADA")
        
        esp_alvo_series = pd.Series([row_alvo.get(col_esp_ind, '')]) if col_esp_ind else pd.Series()
        esp_alvo_set = extrair_especialidades(esp_alvo_series)
        
        resultados = []; piscina_mun = set(); piscina_reg = set()
        
        for _, row in df_ind.iterrows():
            cnpj_cand = str(row.get(col_cnpj_ind, '')).strip()
            if cnpj_cand == cnpj_alvo or cnpj_cand == 'NAN' or not cnpj_cand: continue
            
            mun_cand = str(row.get(col_mun_ind, '')).strip().upper(); uf_cand = str(row.get(col_uf_ind, '')).strip().upper()
            chave_geo_cand = normalizar_texto(mun_cand) + "_" + normalizar_texto(uf_cand)
            
            ibge_cand = str(row.get(col_ibge_ind, '')) if col_ibge_ind else None
            regiao_cand = None
            if ibge_cand and str(ibge_cand).upper() not in ['NAN', 'NONE', 'NULL', '']: regiao_cand = map_ibge_regiao.get(limpar_ibge(ibge_cand))
            if not regiao_cand: regiao_cand = map_geo_regiao.get(chave_geo_cand, "NÃO ENCONTRADA")
            
            prox_str = ""; ordem_prox = 2
            if normalizar_texto(mun_cand) == normalizar_texto(mun_alvo) and normalizar_texto(uf_cand) == normalizar_texto(uf_alvo):
                prox_str = "1 - Mesmo Município"; ordem_prox = 0
            elif regiao_cand == regiao_alvo and regiao_alvo and regiao_alvo != "NÃO ENCONTRADA":
                prox_str = "2 - Mesma Região de Saúde"; ordem_prox = 1
            else:
                continue
            
            tipo_cand = str(row.get(col_tipo_ind, '')).strip().upper() if col_tipo_ind else "-"
            if filtro_tipo == "MESMO TIPO DO ALVO" and tipo_cand != tipo_alvo: continue
            elif filtro_tipo and filtro_tipo != "MESMO TIPO DO ALVO" and tipo_cand != filtro_tipo: continue
                
            esp_cand_series = pd.Series([row.get(col_esp_ind, '')]) if col_esp_ind else pd.Series()
            esp_cand_set = extrair_especialidades(esp_cand_series)
            
            if ordem_prox == 0: 
                piscina_mun.update(esp_cand_set)
                piscina_reg.update(esp_cand_set)
            elif ordem_prox == 1:
                piscina_reg.update(esp_cand_set)
            
            faltante = esp_alvo_set - esp_cand_set; nome_cand = str(row.get(col_nome_ind, '')).strip()
            if len(faltante) == 0: status_cob = "100% Coberto (Sem Vuln)"; txt_faltante = "Nenhuma"
            else: status_cob = f"Vulnerabilidade: Faltam {len(faltante)}"; txt_faltante = "\n".join([f"• {f}" for f in sorted(faltante)])
            resultados.append((ordem_prox, len(faltante), prox_str, cnpj_cand, nome_cand, tipo_cand, f"{mun_cand}/{uf_cand}", status_cob, txt_faltante))
            
        resultados.sort(key=lambda x: (x[0], x[1], x[4])) 
        
        faltante_mun = esp_alvo_set - piscina_mun
        if len(faltante_mun) == 0: status_mun = "100% Coberto pelo Município"; txt_faltante_mun = "A soma de todos os prestadores do município cobre as especialidades."
        else: status_mun = f"Vulnerabilidade no Município: Faltam {len(faltante_mun)}"; txt_faltante_mun = "\n".join([f"• {f}" for f in sorted(faltante_mun)])
        
        faltante_pool = esp_alvo_set - piscina_reg
        if len(faltante_pool) == 0: status_pool = "100% Coberto pela Rede da Região"; txt_faltante_pool = "A soma de todos os prestadores da região cobre as especialidades."
        else: status_pool = f"Vulnerabilidade na Região: Faltam {len(faltante_pool)}"; txt_faltante_pool = "\n".join([f"• {f}" for f in sorted(faltante_pool)])
            
        pool_mun_row = ("☆ POOL DO MUNICÍPIO ☆", "MÚLTIPLOS", "SOMA DOS PRESTADORES DO MUNICÍPIO", filtro_tipo, f"{mun_alvo}/{uf_alvo}", status_mun, txt_faltante_mun)
        pool_reg_row = ("☆ POOL DA REGIÃO ☆", "MÚLTIPLOS", "SOMA DOS PRESTADORES DA REGIÃO DE SAÚDE", filtro_tipo, regiao_alvo, status_pool, txt_faltante_pool)

        output = BytesIO()
        wb = Workbook(); ws = wb.active; ws.title = "Substituicao Individual"
        fill_header = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
        fill_cols = PatternFill(start_color="e0e0e0", end_color="e0e0e0", fill_type="solid")
        fill_100 = PatternFill(start_color="e2efda", end_color="e2efda", fill_type="solid"); fill_pool = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid") 
        font_white = Font(color="FFFFFF", bold=True); font_black = Font(color="000000", bold=False)
        align_c = Alignment(horizontal="center", vertical="center"); align_l = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin_border = Border(left=Side(style='thin', color='A0A0A0'), right=Side(style='thin', color='A0A0A0'), top=Side(style='thin', color='A0A0A0'), bottom=Side(style='thin', color='A0A0A0'))

        ws.merge_cells("A1:G1")
        c_title = ws.cell(row=1, column=1, value=f"RELATÓRIO DE SUBSTITUIÇÃO: {cnpj_alvo} - {nome_alvo_real} ({mun_alvo}/{uf_alvo})")
        c_title.fill = fill_header; c_title.font = font_white; c_title.alignment = align_c
        
        headers = ["Proximidade", "CNPJ", "Prestador Candidato", "Tipo Prestador", "Município/UF", "Status Cobertura", "OBSERVAÇÃO (O que falta)"]
        for col_idx, texto in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col_idx, value=texto); cell.fill = fill_cols; cell.font = Font(color="000000", bold=True); cell.alignment = align_c; cell.border = thin_border
            
        for i, val in enumerate(pool_mun_row, 1):
            cp = ws.cell(row=3, column=i, value=val); cp.fill = fill_pool; cp.font = Font(color="000000", bold=True); cp.alignment = align_l; cp.border = thin_border
            
        for i, val in enumerate(pool_reg_row, 1):
            cp = ws.cell(row=4, column=i, value=val); cp.fill = fill_pool; cp.font = Font(color="000000", bold=True); cp.alignment = align_l; cp.border = thin_border

        linha = 5
        for r in resultados:
            c1 = ws.cell(row=linha, column=1, value=r[2]); c2 = ws.cell(row=linha, column=2, value=r[3])
            c3 = ws.cell(row=linha, column=3, value=r[4]); c4 = ws.cell(row=linha, column=4, value=r[5])
            c5 = ws.cell(row=linha, column=5, value=r[6]); c6 = ws.cell(row=linha, column=6, value=r[7]); c7 = ws.cell(row=linha, column=7, value=r[8])
            for c in [c1, c2, c3, c4, c5, c6, c7]:
                c.alignment = align_l; c.border = thin_border
                if r[1] == 0: c.fill = fill_100
                else: c.font = font_black
            linha += 1

        ws.column_dimensions['A'].width = 25; ws.column_dimensions['B'].width = 20; ws.column_dimensions['C'].width = 50; ws.column_dimensions['D'].width = 20
        ws.column_dimensions['E'].width = 20; ws.column_dimensions['F'].width = 25; ws.column_dimensions['G'].width = 65

        wb.save(output); output.seek(0)
        
        res_100_mun = [r for r in resultados if r[1] == 0 and r[0] == 0]
        res_100_reg = [r for r in resultados if r[1] == 0 and r[0] == 1]
        res_vul_mun = [r for r in resultados if r[1] > 0 and r[0] == 0]
        res_vul_reg = [r for r in resultados if r[1] > 0 and r[0] == 1]
        
        b64_data = base64.b64encode(output.getvalue()).decode('utf-8')
        
        return render_template_string(TELA_PENTEFINO_INDIVIDUAL_RESULTADO, alvo_nome=nome_alvo_real, alvo_cnpj=cnpj_alvo, excel_b64=b64_data, 
                                      res_100_mun=res_100_mun, res_100_reg=res_100_reg,
                                      res_vul_mun=res_vul_mun, res_vul_reg=res_vul_reg, 
                                      regiao_alvo=regiao_alvo, mun_alvo=mun_alvo, uf_alvo=uf_alvo,
                                      pool_mun_status=status_mun, pool_mun_faltantes=txt_faltante_mun,
                                      pool_reg_status=status_pool, pool_reg_faltantes=txt_faltante_pool)
    except Exception as e: flash(f"Erro ao processar as planilhas: {e}", "danger"); return redirect(url_for('pentefino_individual'))

if __name__ == '__main__':
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port)
