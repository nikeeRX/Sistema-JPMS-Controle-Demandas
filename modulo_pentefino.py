import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import pandas as pd
import unicodedata
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

class PenteFinoFrame(tk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent, bg="#f0f4f8")
        self.controller = controller
        
        self.top_bar = tk.Frame(self, bg="#333", height=50)
        self.top_bar.pack(fill=tk.X)
        tk.Button(self.top_bar, text="⬅ Voltar ao Portal Principal", command=lambda: self.controller.show_frame("Portal"), bg="#777", fg="white", font=("Helvetica", 10, "bold"), relief=tk.FLAT).pack(side=tk.LEFT, padx=20, pady=10)
        self.lbl_user_info = tk.Label(self.top_bar, text="", bg="#333", fg="white", font=("Helvetica", 10, "bold"))
        self.lbl_user_info.pack(side=tk.RIGHT, padx=20, pady=10)
        
        self.path_postal, self.path_operadora, self.path_csv = tk.StringVar(), tk.StringVar(), tk.StringVar()
        self.var_tipo_postal, self.var_esp_postal, self.var_tipo_op, self.var_esp_op = tk.StringVar(), tk.StringVar(), tk.StringVar(), tk.StringVar()
        self.comboboxes = {}
        
        self.IGNORADAS = {"DIARIAS", "PACOTE", "TAXAS E GASES", "MATERIAIS E OPME", "MEDICAMENTOS", "FORNECEDOR DE MEDICAMENTO ONCOLOGICO", "INSTRUMENTADOR CIRURGICO", "HOME CARE - ATENDIMENTO DOMICILIAR", "REMOCAO", "MEDICINA LEGAL E PERICIA MEDICA", "MEDICO HIPERBARISTA"}
        self.DE_PARA = {"MEDICO HEMOTERAPEUTA": "HEMATOLOGIA E HEMOTERAPIA", "GENETICA MEDICA": "MEDICO GENETICISTA", "MEDICO CANCEROLOGISTA CIRURGICO": "CIRURGIA ONCOLOGICA", "MEDICO CANCEROLOGISTA PEDIATRICO": "CANCEROLOGIA", "MEDICO DE FAMILIA E COMUNIDADE": "MEDICINA DE FAMILIA E COMUNIDADE", "MEDICO GENERALISTA": "CLINICA MEDICA", "MEDICO PATOLOGISTA": "PATOLOGIA CLINICA/MEDICINA LABORATORIAL", "PATOLOGIA": "PATOLOGIA CLINICA/MEDICINA LABORATORIAL", "MEDICO NEUROFISIOLOGISTA": "NEUROFISIOLOGIA CLINICA", "NEUROPSICOLOGO": "PSICOLOGIA", "PSICOPEDAGOGO": "PSICOLOGIA", "PSICOMOTRICISTA": "FISIOTERAPIA", "MUSICOTERAPEUTA": "TERAPIA OCUPACIONAL", "ORTOPEDISTA": "ORTOPEDIA E TRAUMATOLOGIA", "BIOMEDICO": "PATOLOGIA CLINICA/MEDICINA LABORATORIAL", "FONOAUDIOLOGO EDUCACIONAL": "FONOAUDIOLOGIA", "FONOAUDIOLOGO EM AUDIOLOGIA": "FONOAUDIOLOGIA", "FONOAUDIOLOGO EM DISFAGIA": "FONOAUDIOLOGIA", "FONOAUDIOLOGO EM LINGUAGEM": "FONOAUDIOLOGIA", "FONOAUDIOLOGO EM MOTRICIDADE OROFACIAL": "FONOAUDIOLOGIA", "FONOAUDIOLOGO EM VOZ": "FONOAUDIOLOGIA", "FISIOTERAPEUTA NEUROFUNCIONAL": "FISIOTERAPIA", "FISIOTERAPEUTA OSTEOPATA": "FISIOTERAPIA", "FISIOTERAPEUTA RESPIRATORIA": "FISIOTERAPIA", "FISIOTERAPEUTA TRAUMATO-ORTOPEDICA FUNCIONAL": "FISIOTERAPIA", "ENFERMEIRO DA ESTRATEGIA DE SAUDE DA FAMILIA": "ENFERMEIRO", "CIRURGIAO DENTISTA - TRAUMATOLOGISTA BUCOMAXILOFAC": "CIRURGIA E TRAUMATOLOGIA BUCO-MAXILO-FACIAL", "CLINICA GERAL - ODONTOLOGIA": "ODONTOLOGIA", "CIRURGIAO DENTISTA - DENTISTICA": "ODONTOLOGIA", "CIRURGIAO DENTISTA - DISFUNCAO TEMPOROMANDIBULAR E": "ODONTOLOGIA", "DENTISTICA RESTAURADORA": "ODONTOLOGIA", "DISFUNCAO TEMPOROMANDIBULAR E DOR OROFACIAL": "ODONTOLOGIA", "ENDODONTIA": "ODONTOLOGIA", "ESTOMATOLOGIA": "ODONTOLOGIA", "ODONTOLOGIA - AUDITORIA INICIAL E FINAL": "ODONTOLOGIA", "ODONTOLOGIA DO TRABALHO": "ODONTOLOGIA", "ODONTOLOGIA P/ PACIENTE COM NECESSIDADE ESPECIAL": "ODONTOLOGIA", "ODONTOPEDIATRIA": "ODONTOLOGIA", "PERIODONTIA": "ODONTOLOGIA", "PROTESE DENTARIA": "ODONTOLOGIA", "RADIOLOGIA ODONTOLOGICA E IMAGINOLOGIA - RX ODONTO": "RADIOLOGIA E DIAGNOSTICO POR IMAGEM", "NEUROLIGIA": "NEUROLOGIA"}
        
        self.df_ind = None
        self.lista_alvos_completa = []
        self.path_postal_ind = tk.StringVar()
        self.combo_alvo_var = tk.StringVar()
        self.combo_filtro_tipo_var = tk.StringVar()
        
        self.setup_ui()

    def refresh_all(self):
        self.lbl_user_info.config(text=f"Usuário: {self.controller.current_user} | Nível: {self.controller.current_role.upper()}")

    def setup_ui(self):
        style = ttk.Style()
        style.theme_use('clam')
        style.configure("TButton", font=("Segoe UI", 10, "bold"), padding=6)
        style.configure("Header.TLabel", background="#f0f4f8", font=("Segoe UI", 16, "bold"), foreground="#004b87")

        header_frame = ttk.Frame(self, padding=20)
        header_frame.pack(fill=tk.X)
        ttk.Label(header_frame, text="Pente Fino RN 665 - Substituição de Prestadores", style="Header.TLabel").pack()
        
        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill=tk.BOTH, expand=True, padx=20, pady=10)
        
        self.tab_regional = ttk.Frame(self.notebook)
        self.tab_individual = ttk.Frame(self.notebook)
        
        self.notebook.add(self.tab_regional, text="1. Análise Regional de Cobertura (Massa)")
        self.notebook.add(self.tab_individual, text="2. Busca Individual (Substituição Interna na Postal)")
        
        self.setup_tab_regional()
        self.setup_tab_individual()

    # ------- ABA 1: REGIONAL (EM MASSA) -------
    def setup_tab_regional(self):
        main_frame = ttk.Frame(self.tab_regional, padding=20)
        main_frame.pack(fill=tk.BOTH, expand=True)
        self.create_upload_card_with_mapping(main_frame, "1. Base Post4l Saúde (Hospitais)", self.path_postal, self.load_postal, self.var_tipo_postal, self.var_esp_postal, "postal")
        self.create_upload_card_with_mapping(main_frame, "2. Base Operadora Intermediária", self.path_operadora, self.load_operadora, self.var_tipo_op, self.var_esp_op, "op")
        
        frame_csv = tk.Frame(main_frame, bg="white", highlightbackground="#cbd5e1", highlightthickness=1, bd=0)
        frame_csv.pack(fill=tk.X, pady=8, ipady=5)
        tk.Label(frame_csv, text="3. Base Geográfica IBGE (pop_ibge.csv)", bg="white", font=("Segoe UI", 10, "bold"), fg="#334155").pack(side=tk.TOP, anchor=tk.W, padx=10, pady=5)
        ttk.Button(frame_csv, text="Procurar Arquivo", command=self.load_csv).pack(side=tk.RIGHT, padx=10)
        tk.Label(frame_csv, textvariable=self.path_csv, bg="white", font=("Segoe UI", 9), fg="#64748b", width=65, anchor="w").pack(side=tk.LEFT, fill=tk.X, padx=10)

        btn_frame = tk.Frame(self.tab_regional, bg="#f0f4f8", pady=10)
        btn_frame.pack(fill=tk.X, padx=20)
        self.btn_processar = tk.Button(btn_frame, text="RODAR ANÁLISE COMPLETA E GERAR EXCEL", bg="#004b87", fg="white", font=("Segoe UI", 12, "bold"), relief="flat", command=self.processar_analise)
        self.btn_processar.pack(fill=tk.X, ipady=10)

        log_frame = ttk.Frame(self.tab_regional, padding=20)
        log_frame.pack(fill=tk.BOTH, expand=True)
        ttk.Label(log_frame, text="Log do Sistema:", background="#f0f4f8").pack(anchor=tk.W)
        self.txt_log = tk.Text(log_frame, height=8, bg="#1e293b", fg="#10b981", font=("Consolas", 10))
        self.txt_log.pack(fill=tk.BOTH, expand=True)
        self.log_message("Sistema Regional Online. Carregue as bases para mapear as colunas.")

    def create_upload_card_with_mapping(self, parent, title, var_path, command, var_tipo, var_esp, prefix):
        frame = tk.Frame(parent, bg="white", highlightbackground="#cbd5e1", highlightthickness=1, bd=0)
        frame.pack(fill=tk.X, pady=8, ipady=5)
        top_frame = tk.Frame(frame, bg="white")
        top_frame.pack(fill=tk.X)
        tk.Label(top_frame, text=title, bg="white", font=("Segoe UI", 10, "bold"), fg="#334155").pack(side=tk.TOP, anchor=tk.W, padx=10, pady=5)
        ttk.Button(top_frame, text="Carregar Arquivo", command=command).pack(side=tk.RIGHT, padx=10)
        tk.Label(top_frame, textvariable=var_path, bg="white", font=("Segoe UI", 9), fg="#64748b", anchor="w").pack(side=tk.LEFT, fill=tk.X, padx=10, expand=True)
        
        map_frame = tk.Frame(frame, bg="#f8fafc", pady=5)
        map_frame.pack(fill=tk.X, padx=10, pady=5)
        tk.Label(map_frame, text="Coluna TIPO PRESTADOR:", bg="#f8fafc", font=("Segoe UI", 9, "bold"), fg="#0f172a").grid(row=0, column=0, sticky=tk.W, padx=5)
        cb_tipo = ttk.Combobox(map_frame, textvariable=var_tipo, state="readonly", width=35)
        cb_tipo.grid(row=0, column=1, padx=5)
        tk.Label(map_frame, text="Coluna ESPECIALIDADE:", bg="#f8fafc", font=("Segoe UI", 9, "bold"), fg="#0f172a").grid(row=0, column=2, sticky=tk.W, padx=(20, 5))
        cb_esp = ttk.Combobox(map_frame, textvariable=var_esp, state="readonly", width=35)
        cb_esp.grid(row=0, column=3, padx=5)
        self.comboboxes[f"{prefix}_tipo"] = cb_tipo
        self.comboboxes[f"{prefix}_esp"] = cb_esp

    def carregar_colunas_no_combobox(self, filepath, prefix):
        try:
            if filepath.lower().endswith('.csv'):
                try: df_temp = pd.read_csv(filepath, sep=';', nrows=1, encoding='latin1')
                except: df_temp = pd.read_csv(filepath, sep=',', nrows=1, encoding='utf-8')
            else: df_temp = pd.read_excel(filepath, nrows=1)

            colunas = [str(c).strip().upper() for c in df_temp.columns]
            self.comboboxes[f"{prefix}_tipo"]['values'] = colunas
            self.comboboxes[f"{prefix}_esp"]['values'] = colunas
            self.log_message(f"Colunas lidas com sucesso! Mapeie o TIPO e a ESPECIALIDADE na caixa de '{prefix.upper()}'.")
        except Exception as e: 
            self.log_message(f"Erro ao ler cabeçalho: {str(e)}")

    def load_postal(self):
        f = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx;*.xls")])
        if f: 
            self.path_postal.set(f)
            self.carregar_colunas_no_combobox(f, "postal")

    def load_operadora(self):
        f = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx;*.xls")])
        if f: 
            self.path_operadora.set(f)
            self.carregar_colunas_no_combobox(f, "op")

    def load_csv(self):
        f = filedialog.askopenfilename(filetypes=[("CSV", "*.csv"), ("Excel", "*.xlsx;*.xls")])
        if f: self.path_csv.set(f)

    def log_message(self, message):
        self.txt_log.insert(tk.END, f"> {message}\n")
        self.txt_log.see(tk.END)
        self.update()

    def normalizar_texto(self, texto):
        if pd.isna(texto) or str(texto).lower() in ['nan', 'none', '']: return ""
        return ''.join(c for c in unicodedata.normalize('NFD', str(texto)) if unicodedata.category(c) != 'Mn').upper().strip()

    def cacador_de_colunas(self, df, palavras_chave):
        for k in palavras_chave:
            for col in df.columns:
                if k == str(col).strip().upper(): return col
        for k in palavras_chave:
            for col in df.columns:
                if k in str(col).strip().upper(): return col
        return None

    def limpar_ibge(self, val):
        if pd.isna(val): return ""
        v = ''.join(filter(str.isdigit, str(val).split('.')[0].strip()))
        return v[:6] if len(v) >= 6 else v

    def extrair_especialidades(self, series):
        specs = set()
        for val in series.dropna():
            for s in str(val).replace(';', ',').split(','):
                v = self.normalizar_texto(s)
                if v:
                    if v in self.IGNORADAS: continue
                    specs.add(self.DE_PARA.get(v, v))
        return specs

    def processar_analise(self):
        p_postal, p_operadora, p_csv = self.path_postal.get(), self.path_operadora.get(), self.path_csv.get()
        col_tipo_pos, col_esp_pos = self.var_tipo_postal.get(), self.var_esp_postal.get()
        col_tipo_op, col_esp_op = self.var_tipo_op.get(), self.var_esp_op.get()

        if not all([p_postal, p_operadora, p_csv]): return messagebox.showwarning("Aviso", "Importe as 3 bases.")
        if not all([col_tipo_pos, col_esp_pos, col_tipo_op, col_esp_op]): return messagebox.showwarning("Aviso", "Mapeie as colunas nas caixas suspensas!")

        self.btn_processar.config(state=tk.DISABLED, bg="#94a3b8")
        try:
            self.log_message("1. Extraindo as Bases do Excel...")
            df_postal = pd.read_excel(p_postal, dtype=str)
            df_operadora = pd.read_excel(p_operadora, dtype=str)
            df_postal.columns = [str(c).strip().upper() for c in df_postal.columns]
            df_operadora.columns = [str(c).strip().upper() for c in df_operadora.columns]

            self.log_message("2. Lendo Base IBGE e mapeando Regiões...")
            if p_csv.lower().endswith(('.xls', '.xlsx')): df_ibge = pd.read_excel(p_csv, dtype=str)
            else:
                try: df_ibge = pd.read_csv(p_csv, sep=';', encoding='latin1', dtype=str)
                except: df_ibge = pd.read_csv(p_csv, sep=',', encoding='utf-8', dtype=str)
            df_ibge.columns = [str(c).strip().upper() for c in df_ibge.columns]

            col_ibge_7 = self.cacador_de_colunas(df_ibge, ['COMPLETO'])
            col_ibge_6 = self.cacador_de_colunas(df_ibge, ['AJUSTADO', 'CÓD. MUNIC'])
            col_regiao_csv = self.cacador_de_colunas(df_ibge, ['REGIÃO', 'REGIAO DE SAUDE', 'NOME DA REGIÃO DE SAÚDE'])
            col_mun_csv = self.cacador_de_colunas(df_ibge, ['MUNICÍPIO', 'MUNICIPIO'])
            col_uf_csv = self.cacador_de_colunas(df_ibge, ['UF'])

            map_ibge_regiao = {}
            if col_ibge_7: map_ibge_regiao.update(dict(zip(df_ibge[col_ibge_7].apply(self.limpar_ibge), df_ibge[col_regiao_csv])))
            if col_ibge_6: map_ibge_regiao.update(dict(zip(df_ibge[col_ibge_6].apply(self.limpar_ibge), df_ibge[col_regiao_csv])))
            df_ibge['CHAVE_GEO'] = df_ibge[col_mun_csv].apply(self.normalizar_texto) + "_" + df_ibge[col_uf_csv].apply(self.normalizar_texto)
            map_geo_regiao = dict(zip(df_ibge['CHAVE_GEO'], df_ibge[col_regiao_csv]))

            self.log_message("3. Pente Fino Geográfico...")
            dics_processados = []
            for df in [df_postal, df_operadora]:
                col_ibge = self.cacador_de_colunas(df, ['IBGE', 'CÓDIGO IBGE'])
                col_mun = self.cacador_de_colunas(df, ['MUNICÍPIO', 'MUNICIPIO', 'CIDADE'])
                col_uf = self.cacador_de_colunas(df, ['UF', 'ESTADO'])
                
                if col_ibge: df['REG_IBGE'] = df[col_ibge].apply(self.limpar_ibge).map(map_ibge_regiao)
                else: df['REG_IBGE'] = None
                
                if col_mun and col_uf: df['REG_GEO'] = (df[col_mun].apply(self.normalizar_texto) + "_" + df[col_uf].apply(self.normalizar_texto)).map(map_geo_regiao)
                else: df['REG_GEO'] = None
                
                df['REGIAO_SAUDE'] = df['REG_IBGE'].fillna(df['REG_GEO']).fillna("REGIÃO NÃO IDENTIFICADA")
                dics_processados.append(df)

            df_postal, df_operadora = dics_processados[0], dics_processados[1]

            self.log_message("4. Traduzindo Especialidades e Agrupando a Postal...")
            col_cnpj_pos = self.cacador_de_colunas(df_postal, ['CNPJ', 'CPFCNPJ'])
            col_nome_pos = self.cacador_de_colunas(df_postal, ['NOME', 'RAZAO', 'PRESTADOR', 'FANTASIA'])
            col_mun_pos = self.cacador_de_colunas(df_postal, ['MUNICÍPIO', 'MUNICIPIO', 'CIDADE'])
            if not col_cnpj_pos: df_postal['CNPJ_TEMP'] = "S/CNPJ"; col_cnpj_pos = 'CNPJ_TEMP'
            if not col_nome_pos: df_postal['NOME_TEMP'] = "S/NOME"; col_nome_pos = 'NOME_TEMP'
            if not col_mun_pos: df_postal['MUN_TEMP'] = "S/MUNICIPIO"; col_mun_pos = 'MUN_TEMP'
            df_postal['CNPJ_LIMPO'] = df_postal[col_cnpj_pos].astype(str).str.replace(r'\D', '', regex=True)

            dict_agg = {col_esp_pos: lambda s: ", ".join(sorted(list(self.extrair_especialidades(s)))), col_tipo_pos: lambda x: " | ".join(x.dropna().astype(str).unique())}
            
            df_postal_agrupado = df_postal.groupby(['REGIAO_SAUDE', 'CNPJ_LIMPO', col_nome_pos, col_mun_pos], dropna=False).agg(dict_agg).reset_index()

            self.log_message("5. Mapeando a Piscina da Operadora por Região...")
            regioes_operadora = set(df_operadora['REGIAO_SAUDE'].dropna().unique())
            regioes_operadora.discard("REGIÃO NÃO IDENTIFICADA")
            df_postal_alvos = df_postal_agrupado[df_postal_agrupado['REGIAO_SAUDE'].isin(regioes_operadora)].copy()
            col_nome_op = self.cacador_de_colunas(df_operadora, ['NOME FANTASIA', 'RAZÃO SOCIAL', 'NOME', 'PRESTADOR'])

            piscina_operadora = {}
            for regiao, grupo in df_operadora.groupby('REGIAO_SAUDE'):
                if regiao == "NÃO ENCONTRADO": continue
                piscina_operadora[regiao] = self.extrair_especialidades(grupo[col_esp_op])

            self.log_message("6. Pente Fino Matemático (Alvos da Postal - Piscina Operadora)...")
            res_reg = {}
            for _, row in df_postal_alvos.iterrows():
                regiao = row.get('REGIAO_SAUDE')
                if regiao not in res_reg: res_reg[regiao] = []
                espec_str = str(row.get(col_esp_pos, ''))
                if not espec_str.strip() or espec_str.upper() == 'NAN': 
                    obs = "Vulnerabilidade: Especialidades não informadas na base da Postal."
                    espec_alvo = set()
                else: 
                    espec_alvo = set([e.strip() for e in espec_str.split(',') if e.strip()])
                    obs = ""
                
                if not obs:
                    faltantes = espec_alvo - piscina_operadora.get(regiao, set())
                    obs = "Vulnerabilidade: " + ", ".join(sorted(list(faltantes))) if faltantes else "Cobertura Total"
                        
                res_reg[regiao].append({"CNPJ": str(row.get('CNPJ_LIMPO', 'N/A')), "Prestador": str(row.get(col_nome_pos, 'N/A')), "Município": str(row.get(col_mun_pos, 'N/A')), "OBSERVAÇÃO": obs})

            self.log_message("7. Incluindo as Regiões da Operadora sem Hospitais da Postal...")
            for regiao in regioes_operadora:
                if regiao not in res_reg: res_reg[regiao] = [{"CNPJ": "", "Prestador": "", "Município": "", "OBSERVAÇÃO": "Sem prestadores postal"}]

            self.log_message("8. Desenhando a planilha final...")
            nome_arquivo = f"Relatorio_RN665_{datetime.now().strftime('%Hh%Mm%Ss')}.xlsx"
            output_path = os.path.join(os.path.dirname(p_postal), nome_arquivo)
            self.exportar_excel_formatado(res_reg, output_path)
            
            self.log_message(f"✅ Sucesso Absoluto! Relatório gerado em: {output_path}")
            messagebox.showinfo("Motor Concluído", f"Cruzamento Regional Realizado com Sucesso!\nRelatório salvo:\n{nome_arquivo}")

        except Exception as e:
            self.log_message(f"ERRO CRÍTICO: {str(e)}")
            messagebox.showerror("Erro de Processamento", f"Ocorreu um erro no motor:\n{str(e)}")
        finally: 
            self.btn_processar.config(state=tk.NORMAL, bg="#004b87")

    def exportar_excel_formatado(self, dicionario_resultados, output_path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Analise Equivalencia"
        
        fill_regiao = PatternFill(start_color="005580", end_color="005580", fill_type="solid")
        font_regiao = Font(color="FFFFFF", bold=True, size=11)
        fill_colunas = PatternFill(start_color="005580", end_color="005580", fill_type="solid")
        font_colunas = Font(color="FFFFFF", bold=True)
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        linha_atual = 1
        for regiao in sorted(dicionario_resultados.keys()):
            ws.merge_cells(start_row=linha_atual, start_column=1, end_row=linha_atual, end_column=4)
            c_reg = ws.cell(row=linha_atual, column=1, value=str(regiao).upper())
            c_reg.fill = fill_regiao
            c_reg.font = font_regiao
            c_reg.alignment = align_center
            linha_atual += 1
            
            for col_idx, texto in enumerate(["CNPJ", "Prestador", "Município", "OBSERVAÇÃO"], 1):
                cell = ws.cell(row=linha_atual, column=col_idx, value=texto)
                cell.fill = fill_colunas
                cell.font = font_colunas
                cell.alignment = align_left
                cell.border = thin_border
            linha_atual += 1
            
            for item in dicionario_resultados[regiao]:
                c1 = ws.cell(row=linha_atual, column=1, value=item['CNPJ'])
                c2 = ws.cell(row=linha_atual, column=2, value=item['Prestador'])
                c3 = ws.cell(row=linha_atual, column=3, value=item['Município'])
                c4 = ws.cell(row=linha_atual, column=4, value=item['OBSERVAÇÃO'])
                for c in [c1, c2, c3, c4]: 
                    c.alignment = align_left
                    c.border = thin_border
                linha_atual += 1
            linha_atual += 1

        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 45
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 65
        wb.save(output_path)

    # ------- ABA 2: BUSCA INDIVIDUAL (SNIPER) - TELA COM VULNERABILIDADE E POOL DA REGIÃO -------
    def setup_tab_individual(self):
        main_frame = ttk.Frame(self.tab_individual, padding=20)
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # 1. Carregar Bases
        f_top = tk.Frame(main_frame, bg="white", highlightbackground="#cbd5e1", highlightthickness=1, bd=0)
        f_top.pack(fill=tk.X, pady=5, ipady=5)
        tk.Label(f_top, text="1. Carregar Bases de Dados", bg="white", font=("Segoe UI", 10, "bold"), fg="#334155").pack(side=tk.TOP, anchor=tk.W, padx=10, pady=5)
        
        f_p = tk.Frame(f_top, bg="white")
        f_p.pack(fill=tk.X, pady=2)
        tk.Label(f_p, text="Postal Saúde:", bg="white", width=15, anchor=tk.W, font=("Segoe UI", 9, "bold")).pack(side=tk.LEFT, padx=10)
        ttk.Button(f_p, text="Carregar", command=self.load_postal_ind).pack(side=tk.RIGHT, padx=10)
        tk.Label(f_p, textvariable=self.path_postal_ind, bg="white", font=("Segoe UI", 8), fg="#64748b").pack(side=tk.LEFT, fill=tk.X, expand=True)

        f_i = tk.Frame(f_top, bg="white")
        f_i.pack(fill=tk.X, pady=2)
        tk.Label(f_i, text="IBGE (Geográfica):", bg="white", width=15, anchor=tk.W, font=("Segoe UI", 9, "bold")).pack(side=tk.LEFT, padx=10)
        ttk.Button(f_i, text="Carregar", command=self.load_csv).pack(side=tk.RIGHT, padx=10)
        tk.Label(f_i, textvariable=self.path_csv, bg="white", font=("Segoe UI", 8), fg="#64748b").pack(side=tk.LEFT, fill=tk.X, expand=True)
        
        # 2. Critérios
        f_mid = tk.Frame(main_frame, bg="#f8fafc", pady=10, highlightbackground="#cbd5e1", highlightthickness=1, bd=0)
        f_mid.pack(fill=tk.X, pady=10)
        tk.Label(f_mid, text="2. Selecionar Alvo para Substituição", bg="#f8fafc", font=("Segoe UI", 10, "bold"), fg="#334155").grid(row=0, column=0, columnspan=4, sticky=tk.W, padx=10, pady=(0,10))
        
        tk.Label(f_mid, text="Prestador a ser substituído (Pesquise CNPJ/Nome):", bg="#f8fafc", font=("Segoe UI", 9, "bold")).grid(row=1, column=0, sticky=tk.W, padx=10)
        self.combo_alvo = ttk.Combobox(f_mid, textvariable=self.combo_alvo_var, width=55)
        self.combo_alvo.grid(row=1, column=1, padx=5, sticky=tk.W)
        self.combo_alvo.bind("<KeyRelease>", self.filtrar_alvos_tempo_real) 
        
        tk.Label(f_mid, text="Filtro de Tipo:", bg="#f8fafc", font=("Segoe UI", 9, "bold")).grid(row=1, column=2, sticky=tk.W, padx=(20,5))
        self.combo_filtro_tipo = ttk.Combobox(f_mid, textvariable=self.combo_filtro_tipo_var, state="readonly", width=30)
        self.combo_filtro_tipo.grid(row=1, column=3, padx=5, sticky=tk.W)
        
        # BOTOES DE AÇÃO
        f_botoes = tk.Frame(f_mid, bg="#f8fafc")
        f_botoes.grid(row=2, column=0, columnspan=4, pady=15)
        ttk.Button(f_botoes, text="Buscar Substitutos na Tela", command=self.buscar_substitutos).pack(side=tk.LEFT, padx=10)
        ttk.Button(f_botoes, text="Exportar Relatório Excel (Geral + Vulnerabilidades + Pool)", command=self.exportar_relatorio_individual).pack(side=tk.LEFT, padx=10)
        
        # 3. Resultados Limpos
        f_bot = tk.Frame(main_frame)
        f_bot.pack(fill=tk.BOTH, expand=True)
        tk.Label(f_bot, text="3. Prestadores Aptos Encontrados", font=("Segoe UI", 10, "bold"), fg="#334155").pack(anchor=tk.W, pady=5)
        
        cols = ("Proximidade", "CNPJ", "Prestador", "Tipo Prestador", "Município/UF", "Status Cobertura")
        sy = ttk.Scrollbar(f_bot, orient="vertical")
        sy.pack(side=tk.RIGHT, fill=tk.Y)
        sx = ttk.Scrollbar(f_bot, orient="horizontal")
        sx.pack(side=tk.BOTTOM, fill=tk.X)
        
        self.tree_subs = ttk.Treeview(f_bot, columns=cols, show="headings", yscrollcommand=sy.set, xscrollcommand=sx.set)
        for c in cols:
            self.tree_subs.heading(c, text=c)
            w = 500 if c == "Prestador" else 150 
            self.tree_subs.column(c, width=w, stretch=False, anchor=tk.W if c != "CNPJ" else tk.CENTER)
            
        self.tree_subs.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        sy.config(command=self.tree_subs.yview)
        sx.config(command=self.tree_subs.xview)

    def load_postal_ind(self):
        f = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx;*.xls")])
        if not f: return
        self.path_postal_ind.set(f)
        try:
            self.df_ind = pd.read_excel(f, dtype=str)
            self.df_ind.columns = [str(c).strip().upper() for c in self.df_ind.columns]
            
            self.col_cnpj_ind = self.cacador_de_colunas(self.df_ind, ['CNPJ', 'CPFCNPJ'])
            self.col_nome_ind = self.cacador_de_colunas(self.df_ind, ['NOME', 'RAZAO', 'PRESTADOR', 'FANTASIA'])
            self.col_mun_ind = self.cacador_de_colunas(self.df_ind, ['MUNICÍPIO', 'MUNICIPIO', 'CIDADE'])
            self.col_uf_ind = self.cacador_de_colunas(self.df_ind, ['UF', 'ESTADO'])
            self.col_tipo_ind = self.cacador_de_colunas(self.df_ind, ['TIPO PRESTADOR', 'TIPO_PRESTADOR', 'TIPOPRESTADOR', 'TIPO'])
            self.col_esp_ind = self.cacador_de_colunas(self.df_ind, ['ESPECIALIDADE', 'ESPECIALIDADES'])
            self.col_ibge_ind = self.cacador_de_colunas(self.df_ind, ['IBGE', 'CÓDIGO IBGE'])
            
            if not all([self.col_cnpj_ind, self.col_nome_ind, self.col_mun_ind, self.col_uf_ind]):
                return messagebox.showerror("Erro", "Colunas básicas não encontradas na Postal.")
                
            self.df_ind[self.col_cnpj_ind] = self.df_ind[self.col_cnpj_ind].astype(str).str.replace(r'\D', '', regex=True)
            
            # --- AGRUPAMENTO OBRIGATÓRIO PARA REMOVER DUPLICADAS (CLONES) DA TELA E DO EXCEL ---
            dict_agg = {}
            if self.col_esp_ind: dict_agg[self.col_esp_ind] = lambda s: ", ".join(sorted(list(self.extrair_especialidades(s))))
            if self.col_tipo_ind: dict_agg[self.col_tipo_ind] = lambda x: " | ".join(x.dropna().astype(str).unique())
            
            group_cols = [c for c in [self.col_cnpj_ind, self.col_nome_ind, self.col_mun_ind, self.col_uf_ind, self.col_ibge_ind] if c]
            self.df_ind = self.df_ind.groupby(group_cols, dropna=False).agg(dict_agg).reset_index()
            # -----------------------------------------------------------------------------------
            
            self.lista_alvos_completa = []
            for _, row in self.df_ind.iterrows():
                cnpj = str(row.get(self.col_cnpj_ind, '')).strip()
                nome = str(row.get(self.col_nome_ind, '')).strip()
                if cnpj and cnpj != 'NAN' and nome and nome != 'NAN': 
                    self.lista_alvos_completa.append(f"{cnpj} - {nome}")
                    
            self.lista_alvos_completa = sorted(list(set(self.lista_alvos_completa)))
            self.combo_alvo['values'] = self.lista_alvos_completa
            
            if self.col_tipo_ind:
                tipos = sorted([str(t).strip().upper() for t in self.df_ind[self.col_tipo_ind].dropna().unique()])
                self.combo_filtro_tipo['values'] = ["Todos (Sem Filtro)", "Mesmo Tipo do Alvo"] + tipos
            else: 
                self.combo_filtro_tipo['values'] = ["Todos (Sem Filtro)"]
                
            self.combo_filtro_tipo.current(0)
            
            messagebox.showinfo("Sucesso", "Base Postal carregada! Lembre-se de carregar também o IBGE.")
        except Exception as e: 
            messagebox.showerror("Erro", f"Erro: {e}")

    def filtrar_alvos_tempo_real(self, event):
        if event.keysym in ('Up', 'Down', 'Left', 'Right', 'Return', 'Tab', 'Shift_L', 'Shift_R', 'Caps_Lock'): return
        pesquisa = self.combo_alvo_var.get().upper()
        if pesquisa == "":
            self.combo_alvo['values'] = self.lista_alvos_completa
        else:
            filtrados = [alvo for alvo in self.lista_alvos_completa if pesquisa in alvo.upper()]
            self.combo_alvo['values'] = filtrados

    def buscar_substitutos(self):
        alvo_str = self.combo_alvo.get()
        if not alvo_str or alvo_str not in self.lista_alvos_completa: return messagebox.showwarning("Aviso", "Selecione um Prestador Alvo válido na lista.")
        p_csv = self.path_csv.get()
        if not p_csv: return messagebox.showwarning("Aviso", "Carregue a Base Geográfica IBGE para mapear as Regiões de Saúde.")
        
        try:
            if p_csv.lower().endswith(('.xls', '.xlsx')): df_ibge = pd.read_excel(p_csv, dtype=str)
            else:
                try: df_ibge = pd.read_csv(p_csv, sep=';', encoding='latin1', dtype=str)
                except: df_ibge = pd.read_csv(p_csv, sep=',', encoding='utf-8', dtype=str)
            df_ibge.columns = [str(c).strip().upper() for c in df_ibge.columns]

            c_ibge_7 = self.cacador_de_colunas(df_ibge, ['COMPLETO'])
            c_ibge_6 = self.cacador_de_colunas(df_ibge, ['AJUSTADO', 'CÓD. MUNIC'])
            c_regiao = self.cacador_de_colunas(df_ibge, ['REGIÃO', 'REGIAO DE SAUDE', 'NOME DA REGIÃO DE SAÚDE'])
            c_mun = self.cacador_de_colunas(df_ibge, ['MUNICÍPIO', 'MUNICIPIO'])
            c_uf = self.cacador_de_colunas(df_ibge, ['UF'])

            map_ibge_regiao = {}
            if c_ibge_7: map_ibge_regiao.update(dict(zip(df_ibge[c_ibge_7].apply(self.limpar_ibge), df_ibge[c_regiao])))
            if c_ibge_6: map_ibge_regiao.update(dict(zip(df_ibge[c_ibge_6].apply(self.limpar_ibge), df_ibge[c_regiao])))
            df_ibge['CHAVE_GEO'] = df_ibge[c_mun].apply(self.normalizar_texto) + "_" + df_ibge[c_uf].apply(self.normalizar_texto)
            map_geo_regiao = dict(zip(df_ibge['CHAVE_GEO'], df_ibge[c_regiao]))
        except Exception as e:
            return messagebox.showerror("Erro", f"Erro ao ler IBGE: {e}")
        
        cnpj_alvo = alvo_str.split(' - ')[0]
        linhas_alvo = self.df_ind[self.df_ind[self.col_cnpj_ind] == cnpj_alvo]
        if linhas_alvo.empty: return
        
        row_alvo = linhas_alvo.iloc[0]
        mun_alvo = str(row_alvo.get(self.col_mun_ind, '')).strip().upper()
        uf_alvo = str(row_alvo.get(self.col_uf_ind, '')).strip().upper()
        chave_geo_alvo = self.normalizar_texto(mun_alvo) + "_" + self.normalizar_texto(uf_alvo)
        tipo_alvo = str(row_alvo.get(self.col_tipo_ind, '')).strip().upper() if self.col_tipo_ind else ""
        
        ibge_alvo = str(row_alvo.get(self.col_ibge_ind, '')) if hasattr(self, 'col_ibge_ind') and self.col_ibge_ind else None
        regiao_alvo = None
        if ibge_alvo and ibge_alvo != 'NAN': regiao_alvo = map_ibge_regiao.get(self.limpar_ibge(ibge_alvo))
        if not regiao_alvo: regiao_alvo = map_geo_regiao.get(chave_geo_alvo, "NÃO ENCONTRADA")
        
        esp_alvo_series = pd.Series([row_alvo.get(self.col_esp_ind, '')]) if self.col_esp_ind else pd.Series()
        esp_alvo_set = self.extrair_especialidades(esp_alvo_series)
        
        if not esp_alvo_set: return messagebox.showwarning("Aviso", "O prestador selecionado não possui especialidades cadastradas para comparar.")
            
        filtro_tipo = self.combo_filtro_tipo.get()
        for i in self.tree_subs.get_children(): self.tree_subs.delete(i)
        
        resultados = []
        for _, row in self.df_ind.iterrows():
            cnpj_cand = str(row.get(self.col_cnpj_ind, '')).strip()
            if cnpj_cand == cnpj_alvo or cnpj_cand == 'NAN' or not cnpj_cand: continue
            
            mun_cand = str(row.get(self.col_mun_ind, '')).strip().upper()
            uf_cand = str(row.get(self.col_uf_ind, '')).strip().upper()
            chave_geo_cand = self.normalizar_texto(mun_cand) + "_" + self.normalizar_texto(uf_cand)
            
            ibge_cand = str(row.get(self.col_ibge_ind, '')) if hasattr(self, 'col_ibge_ind') and self.col_ibge_ind else None
            regiao_cand = None
            if ibge_cand and ibge_cand != 'NAN': regiao_cand = map_ibge_regiao.get(self.limpar_ibge(ibge_cand))
            if not regiao_cand: regiao_cand = map_geo_regiao.get(chave_geo_cand, "NÃO ENCONTRADA")
            
            prox_str = ""
            ordem_prox = 2
            if mun_cand == mun_alvo and uf_cand == uf_alvo: prox_str = "1 - Mesmo Município"; ordem_prox = 0
            elif regiao_cand == regiao_alvo and regiao_alvo != "NÃO ENCONTRADA": prox_str = "2 - Mesma Região de Saúde"; ordem_prox = 1
            else: continue
            
            tipo_cand = str(row.get(self.col_tipo_ind, '')).strip().upper() if self.col_tipo_ind else "-"
            if filtro_tipo == "Mesmo Tipo do Alvo" and tipo_cand != tipo_alvo: continue
            elif filtro_tipo not in ["Todos (Sem Filtro)", "Mesmo Tipo do Alvo"] and tipo_cand != filtro_tipo.upper(): continue
                
            esp_cand_series = pd.Series([row.get(self.col_esp_ind, '')]) if self.col_esp_ind else pd.Series()
            esp_cand_set = self.extrair_especialidades(esp_cand_series)
            
            faltante = esp_alvo_set - esp_cand_set
            
            # Mostra todos na tela, mas identifica a Cobertura
            nome_cand = str(row.get(self.col_nome_ind, '')).strip()
            if len(faltante) == 0:
                status_cob = "100% Coberto (Sem Vuln)"
            else:
                status_cob = f"Vulnerabilidade: Faltam {len(faltante)}"
                
            resultados.append((ordem_prox, len(faltante), prox_str, cnpj_cand, nome_cand, tipo_cand, f"{mun_cand}/{uf_cand}", status_cob))
            
        # Ordenação: Município primeiro -> Os com menos faltas -> Nome
        resultados.sort(key=lambda x: (x[0], x[1], x[4]))
        
        for r in resultados: self.tree_subs.insert("", tk.END, values=(r[2], r[3], r[4], r[5], r[6], r[7]))
        if not resultados: messagebox.showinfo("Busca Concluída", "Nenhum prestador encontrado com os critérios selecionados.")

    # ------ EXPORTAÇÃO EXCEL DO ALVO (AGRUPADO E COM POOL DA REGIÃO) ------
    def exportar_relatorio_individual(self):
        alvo_str = self.combo_alvo.get()
        if not alvo_str or alvo_str not in self.lista_alvos_completa: return messagebox.showwarning("Aviso", "Selecione um Prestador Alvo válido na lista.")
        p_csv = self.path_csv.get()
        if not p_csv: return messagebox.showwarning("Aviso", "Carregue a Base Geográfica IBGE para mapear as Regiões de Saúde.")
        
        try:
            if p_csv.lower().endswith(('.xls', '.xlsx')): df_ibge = pd.read_excel(p_csv, dtype=str)
            else:
                try: df_ibge = pd.read_csv(p_csv, sep=';', encoding='latin1', dtype=str)
                except: df_ibge = pd.read_csv(p_csv, sep=',', encoding='utf-8', dtype=str)
            df_ibge.columns = [str(c).strip().upper() for c in df_ibge.columns]

            c_ibge_7 = self.cacador_de_colunas(df_ibge, ['COMPLETO']); c_ibge_6 = self.cacador_de_colunas(df_ibge, ['AJUSTADO', 'CÓD. MUNIC'])
            c_regiao = self.cacador_de_colunas(df_ibge, ['REGIÃO', 'REGIAO DE SAUDE', 'NOME DA REGIÃO DE SAÚDE'])
            c_mun = self.cacador_de_colunas(df_ibge, ['MUNICÍPIO', 'MUNICIPIO']); c_uf = self.cacador_de_colunas(df_ibge, ['UF'])

            map_ibge_regiao = {}
            if c_ibge_7: map_ibge_regiao.update(dict(zip(df_ibge[c_ibge_7].apply(self.limpar_ibge), df_ibge[c_regiao])))
            if c_ibge_6: map_ibge_regiao.update(dict(zip(df_ibge[c_ibge_6].apply(self.limpar_ibge), df_ibge[c_regiao])))
            df_ibge['CHAVE_GEO'] = df_ibge[c_mun].apply(self.normalizar_texto) + "_" + df_ibge[c_uf].apply(self.normalizar_texto)
            map_geo_regiao = dict(zip(df_ibge['CHAVE_GEO'], df_ibge[c_regiao]))
        except Exception as e: return messagebox.showerror("Erro", f"Erro ao ler IBGE: {e}")
            
        cnpj_alvo = alvo_str.split(' - ')[0]
        linhas_alvo = self.df_ind[self.df_ind[self.col_cnpj_ind] == cnpj_alvo]
        if linhas_alvo.empty: return
        
        row_alvo = linhas_alvo.iloc[0]
        mun_alvo = str(row_alvo.get(self.col_mun_ind, '')).strip().upper()
        uf_alvo = str(row_alvo.get(self.col_uf_ind, '')).strip().upper()
        chave_geo_alvo = self.normalizar_texto(mun_alvo) + "_" + self.normalizar_texto(uf_alvo)
        tipo_alvo = str(row_alvo.get(self.col_tipo_ind, '')).strip().upper() if self.col_tipo_ind else ""
        nome_alvo_real = str(row_alvo.get(self.col_nome_ind, '')).strip()
        
        ibge_alvo = str(row_alvo.get(self.col_ibge_ind, '')) if hasattr(self, 'col_ibge_ind') and self.col_ibge_ind else None
        regiao_alvo = None
        if ibge_alvo and ibge_alvo != 'NAN': regiao_alvo = map_ibge_regiao.get(self.limpar_ibge(ibge_alvo))
        if not regiao_alvo: regiao_alvo = map_geo_regiao.get(chave_geo_alvo, "NÃO ENCONTRADA")
        
        esp_alvo_series = pd.Series([row_alvo.get(self.col_esp_ind, '')]) if self.col_esp_ind else pd.Series()
        esp_alvo_set = self.extrair_especialidades(esp_alvo_series)
        
        if not esp_alvo_set: return messagebox.showwarning("Aviso", "O prestador não possui especialidades.")
            
        filtro_tipo = self.combo_filtro_tipo.get()
        resultados = []
        piscina_regiao = set() 
        
        for _, row in self.df_ind.iterrows():
            cnpj_cand = str(row.get(self.col_cnpj_ind, '')).strip()
            if cnpj_cand == cnpj_alvo or cnpj_cand == 'NAN' or not cnpj_cand: continue
            
            mun_cand = str(row.get(self.col_mun_ind, '')).strip().upper()
            uf_cand = str(row.get(self.col_uf_ind, '')).strip().upper()
            chave_geo_cand = self.normalizar_texto(mun_cand) + "_" + self.normalizar_texto(uf_cand)
            
            ibge_cand = str(row.get(self.col_ibge_ind, '')) if hasattr(self, 'col_ibge_ind') and self.col_ibge_ind else None
            regiao_cand = None
            if ibge_cand and ibge_cand != 'NAN': regiao_cand = map_ibge_regiao.get(self.limpar_ibge(ibge_cand))
            if not regiao_cand: regiao_cand = map_geo_regiao.get(chave_geo_cand, "NÃO ENCONTRADA")
            
            prox_str = ""
            ordem_prox = 2
            if mun_cand == mun_alvo and uf_cand == uf_alvo: prox_str = "1 - Mesmo Município"; ordem_prox = 0
            elif regiao_cand == regiao_alvo and regiao_alvo != "NÃO ENCONTRADA": prox_str = "2 - Mesma Região de Saúde"; ordem_prox = 1
            else: continue
            
            tipo_cand = str(row.get(self.col_tipo_ind, '')).strip().upper() if self.col_tipo_ind else "-"
            if filtro_tipo == "Mesmo Tipo do Alvo" and tipo_cand != tipo_alvo: continue
            elif filtro_tipo not in ["Todos (Sem Filtro)", "Mesmo Tipo do Alvo"] and tipo_cand != filtro_tipo.upper(): continue
                
            esp_cand_series = pd.Series([row.get(self.col_esp_ind, '')]) if self.col_esp_ind else pd.Series()
            esp_cand_set = self.extrair_especialidades(esp_cand_series)
            
            piscina_regiao.update(esp_cand_set)
            
            faltante = esp_alvo_set - esp_cand_set
            nome_cand = str(row.get(self.col_nome_ind, '')).strip()
            
            if len(faltante) == 0:
                status_cob = "100% Coberto (Sem Vuln)"
                txt_faltante = "Nenhuma"
            else:
                status_cob = f"Vulnerabilidade: Faltam {len(faltante)}"
                txt_faltante = "\n".join([f"• {f}" for f in sorted(faltante)])
            
            resultados.append((ordem_prox, len(faltante), prox_str, cnpj_cand, nome_cand, tipo_cand, f"{mun_cand}/{uf_cand}", status_cob, txt_faltante))
            
        resultados.sort(key=lambda x: (x[0], x[1], x[4])) 
        
        faltante_pool = esp_alvo_set - piscina_regiao
        if len(faltante_pool) == 0:
            status_pool = "100% Coberto pela Rede"
            txt_faltante_pool = "A soma de todos os prestadores da região cobre as especialidades."
        else:
            status_pool = f"Vulnerabilidade na Rede: Faltam {len(faltante_pool)}"
            txt_faltante_pool = "\n".join([f"• {f}" for f in sorted(faltante_pool)])
            
        pool_row = ("☆ POOL DA REGIÃO ☆", "MÚLTIPLOS", "SOMA DOS PRESTADORES VÁLIDOS DA REGIÃO", filtro_tipo, regiao_alvo, status_pool, txt_faltante_pool)

        nome_arquivo = f"Relatorio_Substituicao_{cnpj_alvo}_{datetime.now().strftime('%Hh%Mm%Ss')}.xlsx"
        output_path = os.path.join(os.path.dirname(self.path_postal_ind.get()), nome_arquivo)
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Substituicao Individual"
        
        fill_header = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
        fill_cols = PatternFill(start_color="e0e0e0", end_color="e0e0e0", fill_type="solid")
        fill_100 = PatternFill(start_color="e2efda", end_color="e2efda", fill_type="solid") 
        fill_pool = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid") 
        font_white = Font(color="FFFFFF", bold=True)
        font_black = Font(color="000000", bold=False)
        align_c = Alignment(horizontal="center", vertical="center")
        align_l = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin_border = Border(left=Side(style='thin', color='A0A0A0'), right=Side(style='thin', color='A0A0A0'), top=Side(style='thin', color='A0A0A0'), bottom=Side(style='thin', color='A0A0A0'))

        ws.merge_cells("A1:G1")
        c_title = ws.cell(row=1, column=1, value=f"RELATÓRIO DE SUBSTITUIÇÃO: {cnpj_alvo} - {nome_alvo_real} ({mun_alvo}/{uf_alvo})")
        c_title.fill = fill_header; c_title.font = font_white; c_title.alignment = align_c
        
        headers = ["Proximidade", "CNPJ", "Prestador Candidato", "Tipo Prestador", "Município/UF", "Status Cobertura", "OBSERVAÇÃO (O que falta)"]
        for col_idx, texto in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col_idx, value=texto)
            cell.fill = fill_cols; cell.font = Font(color="000000", bold=True); cell.alignment = align_c; cell.border = thin_border
            
        for i, val in enumerate(pool_row, 1):
            cp = ws.cell(row=3, column=i, value=val)
            cp.fill = fill_pool; cp.font = Font(color="000000", bold=True); cp.alignment = align_l; cp.border = thin_border

        linha = 4
        for r in resultados:
            c1 = ws.cell(row=linha, column=1, value=r[2]); c2 = ws.cell(row=linha, column=2, value=r[3])
            c3 = ws.cell(row=linha, column=3, value=r[4]); c4 = ws.cell(row=linha, column=4, value=r[5])
            c5 = ws.cell(row=linha, column=5, value=r[6]); c6 = ws.cell(row=linha, column=6, value=r[7])
            c7 = ws.cell(row=linha, column=7, value=r[8])
            
            for c in [c1, c2, c3, c4, c5, c6, c7]:
                c.alignment = align_l; c.border = thin_border
                if r[1] == 0: 
                    c.fill = fill_100
                else:
                    c.font = font_black
            linha += 1

        ws.column_dimensions['A'].width = 25; ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 50; ws.column_dimensions['D'].width = 20
        ws.column_dimensions['E'].width = 20; ws.column_dimensions['F'].width = 25
        ws.column_dimensions['G'].width = 65

        wb.save(output_path)
        messagebox.showinfo("Sucesso", f"Relatório Detalhado gerado!\nSalvo em: {nome_arquivo}")
        if platform.system() == 'Windows': os.startfile(output_path)
        else: subprocess.call(('open', output_path))
